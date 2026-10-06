"""Conclusion-first brief, using saved experiment evidence and existing figures."""
import csv
import hashlib
import json
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs/overnight_3d_ideas"
TARGET = ROOT / "publication/report"
TMP = ROOT / "tmp/report_redesign"
FONT = "Microsoft YaHei"
INK = "202A35"
BLUE = "234A68"
LIGHT = "EAF1F6"
GRAY = "647382"
BORDER = "D9D9D9"
BREAK_NEXT = False

rows = {}
for rd in ("v5", "v6", "v7", "ideas", "dctrl"):
    with (OUT / rd / "comparison.csv").open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            rows[row["group"]] = row
verification = json.loads((OUT / "verification/final_verification.json").read_text(encoding="utf-8"))
for vr in verification["rows"]:
    row = rows.setdefault(vr["group"], {})
    for k in ("final_collision_pairs", "accepted_moves", "candidate_evaluations",
              "relocations", "returns", "stage_length_delta_mm"):
        if k in vr:
            if k in row and row[k] not in ("", None):
                old = float(row[k])
                assert abs(old - float(vr[k])) < 1e-7, (vr["group"], k, old, vr[k])
            row[k] = vr[k]

def r(group):
    return rows[group]

def num(row, key):
    return float(row[key])

def ni(row, key):
    return int(float(row[key]))

def fmt(row, key):
    return f"{ni(row,key):,}"

def length(row):
    return f"{num(row,'stage_length_delta_mm'):.2f}"

def font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    pr = run._element.get_or_add_rPr()
    fonts = pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        pr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), FONT)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def p(doc, text, *, size=11, color=INK, bold=False, after=8, before=0, keep=False):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.23
    pf.keep_with_next = keep
    pf.widow_control = True
    font(para.add_run(text), size, bold, color)
    return para

def heading(doc, text, *, level=1):
    global BREAK_NEXT
    para = doc.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    if BREAK_NEXT:
        para.paragraph_format.page_break_before = True
        BREAK_NEXT = False
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(9)
    para.paragraph_format.keep_with_next = True
    font(para.add_run(text), 17 if level == 1 else 12.5, True, "000000")
    return para

def label(doc, text):
    return p(doc, text, size=10, bold=True, color=BLUE, after=5, before=4, keep=True)

def note(doc, text):
    return p(doc, text, size=9, color=GRAY, after=7)

def setup_section(sec, landscape=False):
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width = Inches(11 if landscape else 8.5)
    sec.page_height = Inches(8.5 if landscape else 11)
    sec.top_margin = sec.bottom_margin = Inches(.45 if landscape else .65)
    sec.left_margin = sec.right_margin = Inches(.7)
    sec.header_distance = Inches(.25)
    sec.footer_distance = Inches(.28)

def next_page(doc, landscape=False):
    global BREAK_NEXT
    current = doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE
    if current == landscape:
        BREAK_NEXT = True
    else:
        setup_section(doc.add_section(WD_SECTION_START.NEW_PAGE), landscape)

def table(doc, headers, data, widths, *, body_size=10, numeric=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top","left","bottom","right","insideH","insideV"):
        edge = OxmlElement("w:"+side)
        edge.set(qn("w:val"),"single")
        edge.set(qn("w:sz"),"4")
        edge.set(qn("w:color"),BORDER)
        borders.append(edge)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, size in (("top",66),("bottom",66),("left",100),("right",100)):
        child=OxmlElement("w:"+side)
        child.set(qn("w:w"),str(size))
        child.set(qn("w:type"),"dxa")
        margins.append(child)
    props.append(margins)
    for j, text in enumerate(headers):
        cell=t.rows[0].cells[j]
        cell.width=Inches(widths[j])
        shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),BLUE)
        cell._tc.get_or_add_tcPr().append(shade)
        para=cell.paragraphs[0]
        para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after=Pt(0)
        para.paragraph_format.keep_with_next=True
        para.paragraph_format.line_spacing=1.12
        font(para.add_run(text),body_size,True,"FFFFFF")
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    trpr=t.rows[0]._tr.get_or_add_trPr()
    repeat=OxmlElement("w:tblHeader"); trpr.append(repeat)
    for i, data_row in enumerate(data):
        row=t.add_row()
        for j, text in enumerate(data_row):
            cell=row.cells[j]; cell.width=Inches(widths[j])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i%2==0:
                shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),LIGHT)
                cell._tc.get_or_add_tcPr().append(shade)
            para=cell.paragraphs[0]
            para.alignment=WD_ALIGN_PARAGRAPH.CENTER if numeric is None or j in numeric else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_after=Pt(0)
            para.paragraph_format.line_spacing=1.15
            para.paragraph_format.widow_control=True
            font(para.add_run(str(text)),body_size,False,INK)
        no_split=OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
    return t

def picture(doc, name, caption, width=9.25):
    path=OUT/"figures"/name
    para=doc.add_paragraph()
    para.paragraph_format.line_spacing=1.0
    para.paragraph_format.space_after=Pt(3)
    para.paragraph_format.keep_with_next=True
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path),width=Inches(width))
    cp=p(doc,caption,size=9,color=GRAY,after=5)
    cp.alignment=WD_ALIGN_PARAGRAPH.LEFT
    return path

def footer(sec):
    para=sec.footer.paragraphs[0]
    para.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(para.add_run("三维布线通宵实验结论  |  "),8.5,color=GRAY)
    for field in ("PAGE","NUMPAGES"):
        if field=="NUMPAGES": font(para.add_run(" / "),8.5,color=GRAY)
        run=para.add_run(); font(run,8.5,color=GRAY)
        el=OxmlElement("w:fldSimple"); el.set(qn("w:instr"),field)
        text=OxmlElement("w:r"); vt=OxmlElement("w:t"); vt.text="1"; text.append(vt); el.append(text)
        run._r.addnext(el)

doc=Document()
setup_section(doc.sections[0])
for s in ("Normal","Title","Heading 1","Heading 2"):
    st=doc.styles[s]
    st.font.name=FONT
    st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(11 if s=="Normal" else 22 if s=="Title" else 17 if s=="Heading 1" else 12.5)
    st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),FONT)
doc.styles["Normal"].paragraph_format.line_spacing=1.23
for border in doc.styles.element.xpath(".//w:pBdr"):
    border.getparent().remove(border)
doc.core_properties.title="三维布线通宵实验结论简报"
doc.core_properties.subject="v5 v6 v7及八项机制的效果 成本与限制"
doc.core_properties.author=""

# Page 1
title=doc.add_paragraph(style="Title")
title.paragraph_format.space_after=Pt(7)
font(title.add_run("三维布线通宵实验结论简报"),22,True,"000000")
p(doc,"v5 v6 v7及八项机制复盘",size=12,color=GRAY,after=5)
p(doc,"2026年10月6日",size=9.5,color=GRAY,after=15)
p(doc,"本轮最明确的改进是每次目标尝试，两侧合计最多评价16个候选。相同的2880次追加评价下，N组终态近距对由30,999降至18,444，减少40.5%；R组降至19,233，减少38.0%。单纯改变候选评价顺序没有改变结果。",size=11.5,after=11)
p(doc,"收益伴随长度和结构代价。N组阶段长度增加从23.78 mm升至57.50 mm，终态升降过渡数从154增至370。现有结果支持提高评价预算的利用率，尚不能说明在相同长度预算下仍有同样优势。",after=11)
label(doc,"表1  2880档主试验的关键结果")
base=r("V6_E0_N2880"); en=r("V6_E2_N2880"); er=r("V6_E2_R2880")
an=r("IDEA_A_FAMILY_N2880"); ar=r("IDEA_A_FAMILY_R2880")
table(doc,["配置","终态近距对","阶段长度\n增加 mm","实际评价\n次数"],[
 ["旧配置 N与R",fmt(base,"final_collision_pairs"),length(base),fmt(base,"candidate_evaluations")],
 ["K16 N",fmt(en,"final_collision_pairs"),length(en),fmt(en,"candidate_evaluations")],
 ["K16 R",fmt(er,"final_collision_pairs"),length(er),fmt(er,"candidate_evaluations")],
 ["K16加家族轮询 N",fmt(an,"final_collision_pairs"),length(an),fmt(an,"candidate_evaluations")],
 ["K16加家族轮询 R",fmt(ar,"final_collision_pairs"),length(ar),fmt(ar,"candidate_evaluations")],
],[2.40,1.12,1.40,1.48],body_size=10.5,numeric={1,2,3})
note(doc,"主起点为42,909对、3对未决、24条已抬升路线。阶段长度相对该起点计算。A家族轮询在2880上限下未用足评价预算，须与预算用满的结果区别陈述。")
heading(doc,"当前可以保留的结论",level=2)
p(doc,"保留原目标排序与K16评价配额。家族轮询值得继续比较长度约束和实际评价对齐后的效果。分层调度让重定位获得了真实执行机会，但不足以支持整体采用；覆盖遍历也没有稳定优势。",after=8)
note(doc,"N只允许首次抬升；R还允许已抬升路线重定位。近距对越少越好；长度与升降结构用于说明代价。")
note(doc,"2880档K16动作数由旧配置53增至N161、R154；评价/动作由54.34降至17.89、18.70。")

# Page 2
next_page(doc,True)
heading(doc,"评价配额是主要有效改动")
p(doc,"E0与先排序的E1逐位一致；E2将每次目标尝试的两侧评价合计限制到16个，同预算完成更多动作。",size=11,after=5)
fig1=picture(doc,"v6_f1_pairs_vs_budget.png","图1  原有v6预算对照图。左侧为主起点，右侧为固定挑战起点；两种起点分开比较。",width=8.4)
label(doc,"表2  主试验三档预算对照")
table(doc,["追加预算","旧配置 N与R","K16 N","K16 R","N相对旧配置少"],[
 [f"{b:,}",fmt(r(f"V6_E0_N{b}"),"final_collision_pairs"),fmt(r(f"V6_E2_N{b}"),"final_collision_pairs"),
  fmt(r(f"V6_E2_R{b}"),"final_collision_pairs"),
  f"{ni(r(f'V6_E0_N{b}'),'final_collision_pairs')-ni(r(f'V6_E2_N{b}'),'final_collision_pairs'):,} 对"]
 for b in (720,1440,2880)
],[1.20,2.0,1.8,1.8,2.4],body_size=10.5)

# Page 3
next_page(doc,False)
heading(doc,"更少近距对伴随更多长度代价")
p(doc,"K16提高了单位评价的全局修复量，同时使用了更多首次抬升。家族轮询使用相同的完整候选池，为16个评价名额分配更多窗口家族，继续降低了终态近距对，但增加长度。",after=10)
label(doc,"表3  2880档主试验的几何成本")
table(doc,["配置","阶段长度\n增加 mm","终态过渡\n段数","已抬升\n路线数"],[
 ["旧配置 N与R",length(base),fmt(base,"final_transition_count"),fmt(base,"final_elevated_route_count")],
 ["K16 N",length(en),fmt(en,"final_transition_count"),fmt(en,"final_elevated_route_count")],
 ["K16 R",length(er),fmt(er,"final_transition_count"),fmt(er,"final_elevated_route_count")],
 ["家族轮询 N",length(an),fmt(an,"final_transition_count"),fmt(an,"final_elevated_route_count")],
 ["家族轮询 R",length(ar),fmt(ar,"final_transition_count"),fmt(ar,"final_elevated_route_count")],
],[2.4,1.4,1.3,1.3],body_size=10.5,numeric={1,2,3})
note(doc,"过渡数是当前终态的升段与降段数量，不是累计生成或累计执行数量。重定位替换原结构，因此动作数与终态结构数不能互相替代。")
label(doc,"表4  家族轮询相对K16的进一步变化")
a_data=[]
for b in (720,1440,2880):
    for mode in (("N",) if b<2880 else ("N","R")):
        aa=r(f"IDEA_A_FAMILY_{mode}{b}"); ee=r(f"V6_E2_{mode}{b}")
        a_data.append([f"{b:,}  "+("N与R" if b<2880 else mode),
          f"{ni(ee,'final_collision_pairs')-ni(aa,'final_collision_pairs'):,}",
          f"{num(aa,'stage_length_delta_mm')-num(ee,'stage_length_delta_mm'):.2f}",
          f"{ni(aa,'candidate_evaluations'):,} / {b:,}"])
table(doc,["预算与模式","额外少的\n近距对","额外长度\n增加 mm","实际评价\n相对上限"],a_data,[1.35,1.35,1.45,2.25],body_size=10.5)
p(doc,"A在2880档触及200次目标尝试上限，实际评价未用满。不能把相同预算上限表述成相同实际评价次数。",after=8)
heading(doc,"需要补的公平对照",level=2)
p(doc,"下一步应同时对齐评价次数和阶段长度上限，判断家族轮询的收益是否值得增加长度。",after=6)
note(doc,"K16也会漏掉更好候选：固定诊断集里75个侧的前缀胜者差于完整枚举，合计收益少14.7%。更多有效动作使全布局结果仍得到改善。")

# Page 4
next_page(doc,False)
heading(doc,"重定位获得检验但贡献有限")
p(doc,"分层调度按当前抬升状态轮询目标，使原先排在很后面的重定位目标进入评价与执行。小预算下整体结果改善主要来自首次抬升转向了新的目标；重定位直接贡献很小。",after=10)
label(doc,"表5  分层调度的主试验结果")
table(doc,["预算","旧配置 N与R","分层 N","分层 R"],[
 [f"{b:,}",fmt(r(f"V5_LEGACY_N{b}"),"final_collision_pairs"),fmt(r(f"V5_STRATIFIED_N{b}"),"final_collision_pairs"),fmt(r(f"V5_STRATIFIED_R{b}"),"final_collision_pairs")]
 for b in (720,1440,2880)
],[1.0,1.9,1.75,1.75],body_size=10.5)
note(doc,"分层N2880的30,425来自完整组产物与最终复核，补足原Markdown表的遗漏。各格使用相同对应预算。")
label(doc,"表6  分层R组重定位的独立贡献")
rel=[]
for b in (720,1440,2880):
    row=r(f"V5_STRATIFIED_R{b}")
    rr=ni(row,"relocation_net_reduction"); net=ni(row,"net_collision_reduction")
    rel.append([f"{b:,}",fmt(row,"relocations"),f"{rr:,}",f"{rr/net*100:.2f}%"])
table(doc,["预算","执行重定位","净减少对数","占本组净收益"],rel,[1.0,1.9,1.75,1.75],body_size=10.5)
p(doc,"重定位净收益是11、27、38对，三档占比均不到0.5%。2880档分层R反而比旧配置多2,344对；预声明采纳规则因此保留旧调度。分层N在该档少574对，方向不能用R结果代替。",after=10)
p(doc,"挑战起点上也有模式差异：分层N较旧配置少930对，分层R多4,212对。这表明等权分配调度机会并未稳定提高全布局净收益。",after=10)
heading(doc,"模式权限与执行收益要区分",level=2)
p(doc,"主试验K16的N和R都没有执行重定位，结果差异来自候选池与K16前缀变化。挑战K16 R执行3次、净减少67对；A主R2880执行1次、净减少21对。不能把所有N/R差异归因于已执行的重定位。")

# Page 5
next_page(doc,True)
heading(doc,"提高目标上限只在受限起点释放预算")
p(doc,"主起点把目标尝试上限从200提高到2000，六个结果逐位一致。挑战起点原配置提前触及200上限，提高上限后可用完剩余评价预算；覆盖遍历在对齐评价数后更差。",size=11,after=5)
fig2=picture(doc,"v7_f2_same_budget_N_vs_R.png","图2  原有v7主起点对照图。T0上限200、T1上限2000的结果相同；T2覆盖遍历的终态近距对更高。",width=9.2)
label(doc,"表7  挑战起点上限变化")
table(doc,["模式","上限200终态","上限2000终态","实际评价 200到2000","少的近距对"],[
 ["N","15,837","14,583","2,353 → 2,880","1,254"],
 ["R","16,537","16,285","2,481 → 2,880","252"],
],[.7,2.0,2.0,2.8,1.7],body_size=10.5)
note(doc,"挑战起点为30,999对。该改善包含释放原未用评价预算的作用，不能视为同实际评价次数下的独立效率提升。")

# Page 6
next_page(doc,False)
heading(doc,"八项机制的证据与当前决定")
p(doc,"单项消融的结论按实际触发条件与正确对照陈述。零终态变化、没有执行、计算调用减少分别回答不同问题，不能统一写成机制无效。",after=10)
label(doc,"表8  八项机制复盘")
table(doc,["机制","本轮观察","当前判断"],[
 ["A 家族轮询","三档终态更少；长度增加；2880实际评价未用满。","有改善潜力\n补成本对照"],
 ["B 动作配额","终态与基线相同；当前设置未体现预留配额的增益。","触发条件受限\n不能外推无效"],
 ["C 等待年龄","六个主终态与K16基线逐位一致。","本设置无改善"],
 ["D 路线覆盖","对专门dctrl：N较差、R较好；2880实际评价不齐。","结果混合\n不普遍采纳"],
 ["E 多锚点","720略好；三档均在668次评价和200目标处停止。","高预算退化\n需分开查上限"],
 ["F 几何去重","没有完全重复几何，也没有与现态相同的候选。","没有可去重项"],
 ["G 判定缓存","终态逐位一致，复用了邻居判定；未做受控串行计时。","计算量有证据\n时间收益未证"],
 ["H 撤回平面","R2880评价7个撤回候选，0次执行。","额外收益未测到"],
],[1.45,3.45,1.50],body_size=10,numeric=set())
note(doc,"D只与STRATIFIED加K16且不启用覆盖的dctrl比较；直接拿D与LEGACY加K16比较会同时改变两个因素。")
heading(doc,"未完成的核心工程",level=2)
p(doc,"v8沿路径弧长的窗口扩展尚未实现、运行或验收。旧审计中33/120个目标侧无合法窗口，本轮没有测量它们是否因扩展生成域而变得可行。",after=10)
p(doc,"固定512条合成路线上的几何结果不能直接外推真实损耗、串扰或工艺合规。实验结论限定于所列起点、模式、策略和评价预算。")

# Page 7
next_page(doc,False)
heading(doc,"证据可信度与下一步")
label(doc,"表9  本轮验证与尚缺证据")
table(doc,["项目","已完成证据或状态"],[
 ["全量回归","778 passed，4 skipped；4个跳过为模块级scipy条件。"],
 ["引擎针对性测试","21 passed；对应本轮实现，历史测试数未充作本轮结果。"],
 ["正式产物复核","126组：104组运行PASS，22组引用完整性校验通过，0 FAIL。"],
 ["终态复核","每个正式终态保存重载后全扫130,816对，核对集合、几何与长度。"],
 ["受控串行计时","尚未进行；旧并行墙钟只作运行描述。"],
 ["v8窗口扩展","未实现、未运行、未验收，不能报告无窗比例改善。"],
],[1.65,4.75],body_size=10.5,numeric=set())
heading(doc,"建议按这个顺序补证据",level=2)
p(doc,"先对齐A的实际评价预算，再在固定阶段长度上限下比较K16与A；这样可以判断减少的近距对是否值得增加长度。之后完成v8几何支持链和正式对照，检验当前无窗侧的适用范围。",after=10)
p(doc,"缓存G采用相同轨迹的开关对照做受控串行计时。继续保留严格验收，不因负结果改判据；分层调度与覆盖遍历应按模式和起点讨论，避免把局部改善写成普遍优势。",after=10)
heading(doc,"数据来源与口径更正",level=2)
note(doc,"数值来源：outputs/overnight_3d_ideas 下v5、v6、v7、ideas、dctrl的comparison.csv，以及verification/final_verification.json。漏列的组以正式组产物和最终复核补齐。图1和图2直接复用原有图件。")
note(doc,"原Markdown把旧配置1440和2880档的终态过渡数写为140和212；当前CSV及路线数对应96和154。本简报采用当前终态口径。原Markdown遗漏分层N2880，本简报补入30,425。")
note(doc,"v7采用决定以adoption_note.json的有效配置为准：LEGACY加E2_ORDERED_K16，K为16，目标上限200。原始实验与旧版报告保留。")
note(doc,"图件选择：v7原预算图的挑战N格仍标注未完成，因CSV漏列；本简报复用无该标注的主起点图，挑战结果以表7中的最终复核为准。")
note(doc,"工程规则、组级明细、完整图集和恢复命令仍见原32页报告与原始产物；本简报用于读取主要结论和证据边界。")

footer(doc.sections[0])
# Later sections inherit the footer automatically.
settings=doc.settings.element
upd=OxmlElement("w:updateFields"); upd.set(qn("w:val"),"true"); settings.append(upd)
TARGET.mkdir(parents=True,exist_ok=True)
TMP.mkdir(parents=True,exist_ok=True)
output=TARGET/"三维布线通宵实验结论简报.docx"
doc.save(output)
evidence={"docx":str(output),"source_verification":str(OUT/"verification/final_verification.json"),
 "groups_total":verification["groups_total"],"new_pass":verification["groups_pass"],
 "referenced":verification["groups_reused_reference"],
 "figures_reused":[{"path":str(x),"sha256":hashlib.sha256(x.read_bytes()).hexdigest()} for x in (fig1,fig2)],
 "corrected_transition_counts":{"1440":96,"2880":154},
 "key_groups":{name:rows[name] for name in ("V6_E0_N2880","V6_E2_N2880","V6_E2_R2880","IDEA_A_FAMILY_N2880","IDEA_A_FAMILY_R2880")},
 "source_files":[{"path":str(OUT/rd/"comparison.csv"),"sha256":hashlib.sha256((OUT/rd/"comparison.csv").read_bytes()).hexdigest()} for rd in ("v5","v6","v7","ideas","dctrl")]}
(TMP/"brief_evidence.json").write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({"docx":str(output),"figures_reused":2,"tables":len(doc.tables),"sections":len(doc.sections)},ensure_ascii=False))

