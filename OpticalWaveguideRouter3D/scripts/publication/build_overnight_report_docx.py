"""Build the overnight round DOCX from the markdown report and the real figures.

Every number in the document comes from docs/reports/step_18_overnight_3d_v5_v8.md,
which was written from the saved artifacts; the figures are the PNGs actually
produced by scripts/overnight_figures.py for rounds v5/v6/v7/ideas.

Usage: <bundled python> scripts/publication/build_overnight_report_docx.py PROJECT
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

CJK = "SimSun"
CJK_HEAD = "SimHei"


def style_run(run, size=None, bold=None, head=False):
    run.font.name = CJK_HEAD if head else CJK
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_HEAD if head else CJK)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def add_body(document, text, size=10.5):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    for piece in re.split(r"(\*\*[^*]+\*\*)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            style_run(paragraph.add_run(piece[2:-2]), size=size, bold=True)
        else:
            style_run(paragraph.add_run(piece), size=size)
    return paragraph


def _repeat_header(row):
    """Mark a table row as a header row so Word repeats it on every page."""
    properties = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    properties.append(element)


def _keep_row_together(row):
    """Forbid Word from splitting one table row across two pages."""
    properties = row._tr.get_or_add_trPr()
    element = OxmlElement("w:cantSplit")
    element.set(qn("w:val"), "true")
    properties.append(element)


def _estimate_widths(header, body, font_pt):
    """Width needed by each column, in points, from its widest cell.

    ASCII runs are measured at about half the font size and CJK at the full
    size, plus cell padding. The estimate is what the column must get so that a
    long identifier such as P4_A_FAMILY_CAP24_R2880 is never broken across two
    lines (which is how a trailing digit ends up alone on the next line)."""
    def width(value):
        text_value = str(value or "")
        total = 0.0
        for ch in text_value:
            total += font_pt if ord(ch) > 127 else font_pt * 0.55
        return total + 8.0
    widths = []
    for index in range(len(header)):
        values = [header[index]] + [row[index] if index < len(row) else "" for row in body]
        widths.append(max(28.0, max(width(v) for v in values)))
    return widths


def _fixed_layout(table, widths_pt):
    properties = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.append(layout)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("left", 40), ("right", 40), ("top", 0), ("bottom", 0)):
        element = OxmlElement("w:" + side)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    properties.append(margins)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for column, width in zip(grid.findall(qn("w:gridCol")), widths_pt):
            column.set(qn("w:w"), str(int(width * 20)))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if index < len(widths_pt):
                cell.width = Pt(widths_pt[index])


def add_table(document, rows, available_pt=458.0, font_pt=9.0):
    header, body = rows[0], rows[2:]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.autofit = False
    estimates = _estimate_widths(header, body, font_pt)
    total = sum(estimates)
    factor = min(1.0, available_pt / total) if total else 1.0
    size = max(6.5, font_pt * factor)
    widths = [estimate * available_pt / total for estimate in estimates] if total else estimates
    for index, cell in enumerate(header):
        paragraph = table.rows[0].cells[index].paragraphs[0]
        style_run(paragraph.add_run(cell), size=size, bold=True, head=True)
    _repeat_header(table.rows[0])
    _keep_row_together(table.rows[0])
    for row in body:
        cells = table.add_row().cells
        for index, value in enumerate(row[:len(header)]):
            paragraph = cells[index].paragraphs[0]
            style_run(paragraph.add_run(value.replace("**", "")), size=size)
        _keep_row_together(table.rows[-1])
    _fixed_layout(table, widths)
    return table


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build(markdown, figures, out_path):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = CJK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            block = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index]); index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            style_run(paragraph.add_run("\n".join(block)), size=8.5)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and set(lines[index + 1]) <= set("|-: "):
            rows = [split_row(line)]
            index += 1
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(split_row(lines[index])); index += 1
            add_table(document, rows)
            continue
        if line.startswith("### "):
            paragraph = document.add_heading(level=3)
            style_run(paragraph.add_run(line[4:]), size=12, bold=True, head=True)
        elif line.startswith("## "):
            paragraph = document.add_heading(level=2)
            style_run(paragraph.add_run(line[3:]), size=14, bold=True, head=True)
        elif line.startswith("# "):
            paragraph = document.add_heading(level=1)
            style_run(paragraph.add_run(line[2:]), size=17, bold=True, head=True)
        elif line.strip().startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            base = line.strip()[2:]
            for piece in re.split(r"(\*\*[^*]+\*\*)", base):
                if not piece:
                    continue
                if piece.startswith("**") and piece.endswith("**"):
                    style_run(paragraph.add_run(piece[2:-2]), size=10.5, bold=True)
                else:
                    style_run(paragraph.add_run(piece), size=10.5)
        elif re.match(r"^\d+\. ", line.strip()):
            add_body(document, line.strip())
        elif line.strip():
            add_body(document, line)
        index += 1
    if figures:
        paragraph = document.add_heading(level=2)
        style_run(paragraph.add_run("附：本轮生成的图（每张同时提供 PDF 与 manifest）"), size=14, bold=True, head=True)
        for caption, path in figures:
            document.add_picture(str(path), width=Inches(6.1))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = document.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_run(cap.add_run(caption + "  (" + Path(path).name + ")"), size=9)
    document.save(str(out_path))
    return out_path


def main():
    root = Path(sys.argv[1]).resolve()
    out = root / "outputs/overnight_3d_ideas"
    markdown = (root / "docs/reports/step_18_overnight_3d_v5_v8.md").read_text(encoding="utf-8")
    figures = []
    for prefix, title in (("v5", "v5 目标调度"), ("v6", "v6 候选评价"), ("v7", "v7 目标覆盖"), ("ideas", "八项机制消融")):
        for path in sorted((out / "figures").glob(prefix + "_f*.png")):
            figures.append((title + " " + path.stem, path))
    target = root / "publication/report/brief_build"
    target.mkdir(parents=True, exist_ok=True)
    docx = target / "三维布线通宵轮次v5v8实验报告.docx"
    build(markdown, figures, docx)
    print(str(docx))
    print("figures", len(figures))


if __name__ == "__main__":
    main()
