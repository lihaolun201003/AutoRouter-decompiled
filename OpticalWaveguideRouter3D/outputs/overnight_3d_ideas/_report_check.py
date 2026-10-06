import sys, json, hashlib
from pathlib import Path
from docx import Document
root = Path(sys.argv[1]).resolve()
docx = root / "publication/report/brief_build/三维布线通宵轮次v5v8实验报告.docx"
d = Document(str(docx))
text = "\n".join(p.text for p in d.paragraphs)
table_text = "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
required = ["v5 目标调度", "v6 候选评价调度", "v7 目标覆盖与遍历", "八项机制单项消融",
            "E2_ORDERED_K16", "30,999", "33,343", "18,444", "19,233", "778 passed",
            "v8 沿路径弧长的窗口扩展完全未实现", "21 passed"]
missing = [r for r in required if r not in text and r not in table_text]
report = dict(docx=str(docx), paragraphs=len(d.paragraphs), tables=len(d.tables),
              inline_images=len(d.inline_shapes), characters=len(text),
              required_present=len(required) - len(missing), required_total=len(required),
              missing_required=missing)
eng = root / "scripts/publication/build_overnight_report_docx.py"
report["sha256"] = {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in [docx, root / "docs/reports/step_18_overnight_3d_v5_v8.md",
                              root / "src/overnight_engine_3d.py", eng]}
(root / "outputs/overnight_3d_ideas/report_check.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
print(json.dumps({k: v for k, v in report.items() if k != "sha256"}, indent=2, ensure_ascii=False))
print("sha256 keys:", len(report["sha256"]))
