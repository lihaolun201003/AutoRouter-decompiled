"""使用临时工作簿验证前代连接解析，不依赖真实数据路径。"""

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from openpyxl import Workbook
from src.io import load_legacy_fiberboard

raises = TestCase().assertRaises


def write_sample(path: Path, rows: list) -> None:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()


def test_connections_identity_direction_and_unknown_state():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1", "Port2"], [1, 2], [1, 2], [2, 1]])
        before = path.read_bytes()
        first = load_legacy_fiberboard(path)
        second = load_legacy_fiberboard(path)
        assert path.read_bytes() == before
        assert len(first) == 3
        assert [w.id for w in first] == [0, 1, 2]
        assert first == second
        assert [(w.start_port.pmt_id, w.end_port.pmt_id) for w in first] == [(1, 2), (1, 2), (2, 1)]
        ports = [p for w in first for p in (w.start_port, w.end_port)]
        assert [p.id for p in ports] == list(range(6))
        assert len({id(p) for p in ports}) == 6
        assert all(p.local_id is None and p.position is None for p in ports)


def test_blank_rows_and_integral_numbers():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1", "Port2"], [None, None], [1.0, 2.0], [None, None], [3, 4]])
        result = load_legacy_fiberboard(path)
        assert [w.id for w in result] == [0, 1]
        assert result[0].start_port.pmt_id == 1
        assert type(result[0].start_port.pmt_id) is int


def test_missing_file():
    with TemporaryDirectory() as directory:
        with raises(FileNotFoundError):
            load_legacy_fiberboard(Path(directory) / "absent.xlsx")


def test_missing_port1_header():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Other","Port2"],[1,2]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_missing_port2_header():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1","Other"],[1,2]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_missing_headers():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [[1,2],[3,4]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_missing_port1():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1","Port2"],[None,2]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_missing_port2():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1","Port2"],[1,None]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_fraction():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1","Port2"],[1.5,2]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_string():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1","Port2"],["bad",2]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_numeric_string():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1","Port2"],["1",2]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_boolean():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1","Port2"],[True,2]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_formula():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1","Port2"],["=1+1",2]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_duplicate_header():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1","Port2","Port1"],[1,2,3]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_header_only():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1", "Port2"]])
        assert load_legacy_fiberboard(path) == []


def test_multiple_sheets_rejected():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        book = Workbook()
        book.active.append(["Port1", "Port2"])
        book.create_sheet("Other")
        book.save(path)
        book.close()
        with raises(ValueError):
            load_legacy_fiberboard(path)


def test_both_formulas_without_cache_rejected():
    with TemporaryDirectory() as directory:
        path = Path(directory) / "sample.xlsx"
        write_sample(path, [["Port1", "Port2"], ["=1+1", "=2+1"]])
        with raises(ValueError):
            load_legacy_fiberboard(path)


SNAPSHOT_HEADERS = [None, "Port1", "Port2", "index1", "index2", "sy", "ly", "dz", "sx", "lx", "dx"]


def test_snapshot_restores_direction_positions_and_ids():
    from src.io import load_legacy_512_snapshot
    from src.models import Point2D
    with TemporaryDirectory() as directory:
        original, snapshot = Path(directory) / "input.xlsx", Path(directory) / "snapshot.xlsx"
        write_sample(original, [["Port1", "Port2"], [1, 2], [1, 2]])
        write_sample(snapshot, [SNAPSHOT_HEADERS,
            [1, 2, 1, 20, 10, 0, 150, 0, 8, 3, 5],
            [0, 1, 2, 10, 20, 150, 0, 0, 2, 7, 5]])
        before = [original.read_bytes(), snapshot.read_bytes()]
        result = load_legacy_512_snapshot(original, snapshot)
        assert result == load_legacy_512_snapshot(original, snapshot)
        assert [w.id for w in result] == [0, 1]
        assert [(w.start_port.pmt_id, w.end_port.pmt_id) for w in result] == [(1, 2), (1, 2)]
        assert result[0].start_port.position == Point2D(2, 150)
        assert result[0].end_port.position == Point2D(7, 0)
        assert result[1].start_port.position == Point2D(3, 150)
        assert result[1].end_port.position == Point2D(8, 0)
        ports = [p for w in result for p in (w.start_port, w.end_port)]
        assert [p.id for p in ports] == [0, 1, 2, 3]
        assert all(p.local_id is None and isinstance(p.position, Point2D) for p in ports)
        assert before == [original.read_bytes(), snapshot.read_bytes()]


def snapshot_rejects(rows: list) -> None:
    from src.io import load_legacy_512_snapshot
    with TemporaryDirectory() as directory:
        original, snapshot = Path(directory) / "input.xlsx", Path(directory) / "snapshot.xlsx"
        write_sample(original, [["Port1", "Port2"], [1, 2]])
        write_sample(snapshot, [SNAPSHOT_HEADERS, *rows])
        with raises(ValueError):
            load_legacy_512_snapshot(original, snapshot)


def test_snapshot_duplicate():
    snapshot_rejects([[0,1,2,10,20,150,0,0,2,7,5],[0,1,2,10,20,150,0,0,2,7,5]])


def test_snapshot_missing():
    snapshot_rejects([])


def test_snapshot_out_of_range():
    snapshot_rejects([[1,1,2,10,20,150,0,0,2,7,5]])


def test_snapshot_negative_index():
    snapshot_rejects([[-1,1,2,10,20,150,0,0,2,7,5]])


def test_snapshot_fractional_index():
    snapshot_rejects([[0.5,1,2,10,20,150,0,0,2,7,5]])


def test_snapshot_pair_mismatch():
    snapshot_rejects([[0,1,3,10,20,150,0,0,2,7,5]])


def test_snapshot_bad_coordinate():
    snapshot_rejects([[0,1,2,10,20,150,0,0,"bad",7,5]])


def test_snapshot_missing_coordinate():
    snapshot_rejects([[0,1,2,10,20,150,0,0,None,7,5]])


def test_snapshot_bad_dx():
    snapshot_rejects([[0,1,2,10,20,150,0,0,2,7,6]])


def test_snapshot_nonzero_dz():
    snapshot_rejects([[0,1,2,10,20,150,0,1,2,7,5]])
