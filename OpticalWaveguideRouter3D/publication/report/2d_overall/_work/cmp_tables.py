
import sys, pathlib, re, json
sys.stdout.reconfigure(encoding="utf-8")
md = pathlib.Path(r"C:\Users\lihao\Desktop\Graduation Project\OpticalWaveguideRouter3D\publication\report\2d_overall\二维光波导自动布线总体实验报告.md")
lines = md.read_text(encoding="utf-8").split("\n")
seps = [i for i,l in enumerate(lines) if re.match(r"^\|\s*-{3,}", l)]
print("md tables:", len(seps))
from collections import Counter
c = Counter()
for i in seps:
    n = len([x for x in lines[i-1].strip().strip("|").split("|")])
    c[n] += 1
print("md table column histogram:", dict(sorted(c.items())))
for i in seps:
    n = len([x for x in lines[i-1].strip().strip("|").split("|")])
    if n >= 12:
        print("  ", n, lines[i-1][:120])
from docx import Document
d = Document(r"C:\Users\lihao\Desktop\Graduation Project\OpticalWaveguideRouter3D\publication\report\2d_overall\二维光波导自动布线总体实验报告.docx")
print("docx tables:", len(d.tables))
cc = Counter((len(t.rows), len(t.columns)) for t in d.tables)
print("docx (rows,cols) histogram:", dict(sorted(cc.items(), key=lambda kv: -kv[0][1])))
for t in d.tables:
    if len(t.columns) >= 12:
        print("   DOCX", len(t.rows), "x", len(t.columns), "|", " / ".join(c.text[:12] for c in t.rows[0].cells)[:150])
