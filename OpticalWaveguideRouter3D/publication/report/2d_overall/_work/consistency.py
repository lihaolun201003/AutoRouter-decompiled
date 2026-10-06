
import sys, pathlib, re
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
d = Document(r"C:\Users\lihao\Desktop\Graduation Project\OpticalWaveguideRouter3D\publication\report\2d_overall\二维光波导自动布线总体实验报告.docx")
parts = []
for p in d.paragraphs: parts.append(p.text)
for t in d.tables:
    for row in t.rows:
        for c in row.cells: parts.append(c.text)
docx_text = "\n".join(parts)
md_text = (pathlib.Path(r"C:\Users\lihao\Desktop\Graduation Project\OpticalWaveguideRouter3D\publication\report\2d_overall") / "二维光波导自动布线总体实验报告.md").read_text(encoding="utf-8")
keys = ["5.2786","4.335995","5.516158","4.5582","2.970298","3.294685","5.078298","5.499483",
        "4.654524","4.900387","−2.3089","−2.2220","152.007142675","10.466","0.0924","0.9504",
        "−0.7749","8.4996","8.4333","2.9986","3.3492","+0.3885","+0.0092","+0.0238","0.0",
        "5.7264","5.6340","5.9657","6.1946","1.9037","2.3826","5.5147","6.5761","9.8096"]
missing_docx = [k for k in keys if k not in docx_text]
missing_md  = [k for k in keys if k not in md_text]
print("key tokens checked:", len(keys))
print("missing in DOCX:", missing_docx)
print("missing in MD:", missing_md)
print("DOCX paragraphs:", len(d.paragraphs), "tables:", len(d.tables))
print("docx text chars:", len(docx_text), "md chars:", len(md_text))
