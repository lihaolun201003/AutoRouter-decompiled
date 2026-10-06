# -*- coding: utf-8 -*-
"""Build the DOCX deliverable from the master Markdown.

Master source:  二维光波导自动布线总体实验报告.md   (single source of truth for numbers)
Output:         二维光波导自动布线总体实验报告.docx
PDF is produced from the DOCX with the bundled LibreOffice kit.

Markdown dialect understood here:
  # / ## / ###          headings (levels 1-3)
  - item                bullet
  **bold**              inline bold (kept in a Song face, never switched to a hei face)
  | a | b |             pipe table, preceded by a paragraph "表N 标题"
  ![caption](path)      centred figure, the alt text becomes the caption below the figure
  > text                pull-quote paragraph (used for 结论边界 notes)
  plain paragraph       body text
"""
from pathlib import Path
import re, sys, json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
MASTER = HERE / '二维光波导自动布线总体实验报告.md'
OUT_DOCX = HERE / '二维光波导自动布线总体实验报告.docx'
EAST = '宋体'
LATIN = 'Times New Roman'
FIG_WIDTH = json.loads((HERE / '_work' / 'fig_widths.json').read_text(encoding='utf-8'))
TABLE_COLS = json.loads((HERE / '_work' / 'table_cols.json').read_text(encoding='utf-8'))

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.2)
sec.left_margin = sec.right_margin = Cm(2.3)
sec.header_distance = sec.footer_distance = Cm(1.1)

NAMES = ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption',
         'Header', 'Footer', 'List Bullet']
for name in NAMES:
    try:
        s = doc.styles[name]
    except KeyError:
        continue
    s.font.name = LATIN
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), EAST)
    s.paragraph_format.widow_control = True
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.5
doc.styles['Normal'].paragraph_format.space_after = Pt(0)
doc.styles['Normal'].paragraph_format.first_line_indent = Pt(21)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.first_line_indent = Pt(0)
doc.styles['Title'].paragraph_format.space_after = Pt(6)
for name, size in [('Heading 1', 15), ('Heading 2', 12), ('Heading 3', 10.5)]:
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.line_spacing = 1.3
    s.paragraph_format.space_before = Pt(12 if name == 'Heading 1' else 9)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.first_line_indent = Pt(0)
    s.paragraph_format.keep_with_next = True
doc.styles['Caption'].font.size = Pt(9.5)
doc.styles['Caption'].font.bold = False
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].paragraph_format.line_spacing = 1.2
doc.styles['Caption'].paragraph_format.first_line_indent = Pt(0)
doc.styles['Caption'].paragraph_format.space_before = Pt(6)
doc.styles['Caption'].paragraph_format.space_after = Pt(3)
doc.styles['List Bullet'].font.size = Pt(10.5)
doc.styles['List Bullet'].paragraph_format.line_spacing = 1.5
doc.styles['List Bullet'].paragraph_format.first_line_indent = Pt(0)

h = sec.header.paragraphs[0]
h.text = '二维光波导自动布线总体实验报告'
h.style = doc.styles['Header']
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in h.runs:
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
f = sec.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.CENTER
f.add_run('第 ')
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
f._p.append(fld)
f.add_run(' 页 / 共 ')
fld2 = OxmlElement('w:fldSimple')
fld2.set(qn('w:instr'), 'NUMPAGES')
f._p.append(fld2)
f.add_run(' 页')
for r in f.runs:
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

STATS = {'tables': 0, 'figures': 0, 'images': [], 'headings': [], 'table_titles': []}


def set_run(run, size=None, bold=None):
    run.font.name = LATIN
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), EAST)
    fonts.set(qn('w:ascii'), LATIN)
    fonts.set(qn('w:hAnsi'), LATIN)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    return run


def add_rich(par, text, size=None, base_bold=False):
    """Render **bold** spans; bold CJK keeps the Song face (no hei substitution)."""
    text = text.replace('`', '')
    for i, chunk in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if chunk == '':
            continue
        run = par.add_run(chunk)
        set_run(run, size=size, bold=(base_bold or (i % 2 == 1)))
    return par


PENDING_BREAK = {'on': False}


def apply_break(par):
    """Turn a pending explicit page break into page_break_before on this paragraph.

    Using page_break_before instead of an empty break paragraph avoids producing a
    blank page when the previous content already ended at a page boundary.
    """
    if PENDING_BREAK['on']:
        par.paragraph_format.page_break_before = True
        PENDING_BREAK['on'] = False
    return par


def body(text, style=None, indent=True, size=None, space_after=0):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    apply_break(p)
    if not indent:
        p.paragraph_format.first_line_indent = Pt(0)
    if space_after:
        p.paragraph_format.space_after = Pt(space_after)
    add_rich(p, text, size=size)
    return p


def border(parent, side, val='nil', size=0):
    el = OxmlElement('w:' + side)
    el.set(qn('w:val'), val)
    if val != 'nil':
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:color'), '000000')
    parent.append(el)


def add_table(headers, rows, caption=None, widths=None, notes=(), text_cols=(0,)):
    STATS['tables'] += 1
    number = STATS['tables']
    if caption:
        cp = apply_break(doc.add_paragraph(caption, style='Caption'))
        cp.paragraph_format.keep_with_next = True
        cp.paragraph_format.space_after = Pt(4)
        for r in cp.runs:
            r.font.bold = False
        STATS['table_titles'].append(caption)
    ncol = len(headers)
    key = caption if caption else ''
    if widths is None:
        widths = TABLE_COLS.get(key[:6])
    if widths is None or len(widths) != ncol:
        widths = [16.4 / ncol] * ncol
    scale = 16.4 / sum(widths)
    widths = [w * scale for w in widths]
    t = doc.add_table(rows=1, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    tw = pr.find(qn('w:tblW'))
    tw.set(qn('w:type'), 'dxa')
    tw.set(qn('w:w'), str(round(sum(widths) * 1440 / 2.54)))
    for old in pr.findall(qn('w:tblBorders')):
        pr.remove(old)
    b = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border(b, side, 'single' if side in ('top', 'bottom') else 'nil', 8)
    pr.append(b)
    for c, w in zip(t.columns, widths):
        c.width = Cm(w)
    hdr = t.rows[0]
    for cell, text in zip(hdr.cells, headers):
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.first_line_indent = Pt(0)
        add_rich(p, text, size=9.5, base_bold=True)
    for values in rows:
        row = t.add_row()
        for cell, text, ci in zip(row.cells, values, range(ncol)):
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            add_rich(p, str(text), size=9)
    for ri, row in enumerate(t.rows):
        rp = row._tr.get_or_add_trPr()
        rp.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            rp.append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            for tag in ['tcBorders', 'tcMar']:
                for old in cp.findall(qn('w:' + tag)):
                    cp.remove(old)
            mar = OxmlElement('w:tcMar')
            for side, val in [('top', 50), ('bottom', 50), ('left', 70), ('right', 70)]:
                e = OxmlElement('w:' + side)
                e.set(qn('w:w'), str(val))
                e.set(qn('w:type'), 'dxa')
                mar.append(e)
            cp.append(mar)
            cb = OxmlElement('w:tcBorders')
            border(cb, 'top', 'single' if ri == 0 else 'nil', 8)
            border(cb, 'bottom', 'single' if ri in (0, len(t.rows) - 1) else 'nil',
                   8 if ri == len(t.rows) - 1 else 5)
            border(cb, 'left', 'nil')
            border(cb, 'right', 'nil')
            cp.append(cb)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                p.paragraph_format.keep_with_next = (ri == 0 or ri == len(t.rows) - 1)
                p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if ci in text_cols
                               else WD_ALIGN_PARAGRAPH.CENTER)
    for item in notes:
        body(item, indent=False, size=9, space_after=3)
    return t


def clean_caption(path, caption):
    """Turn a leading source-file name into a trailing provenance note."""
    name = Path(path).name
    cap = caption.strip()
    for prefix in ('`' + name + '`：', name + '：', '`' + name + '`', name):
        if cap.startswith(prefix):
            cap = cap[len(prefix):].strip()
            break
    if cap and not cap.endswith(('。', '）', ')')):
        cap += '。'
    return cap + '（原始图：' + name + '）'


def add_figure(path, caption):
    STATS['figures'] += 1
    number = STATS['figures']
    p = doc.add_paragraph()
    apply_break(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    width = FIG_WIDTH.get(Path(path).name, 15.5)
    caption = clean_caption(path, caption)
    run = p.add_run()
    shape = run.add_picture(str(HERE / path), width=Cm(width))
    shape._inline.docPr.set('descr', caption)
    cp = doc.add_paragraph('图' + str(number) + ' ' + caption, style='Caption')
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(4)
    STATS['images'].append({'figure': number, 'path': path,
                            'sha256': hashlib.sha256((HERE / path).read_bytes()).hexdigest()})


# ----------------------------------------------------------------- markdown parse
lines = MASTER.read_text(encoding='utf-8').split('\n')
i = 0
pending_caption = None
first_h1 = True


def flush_caption_as_paragraph():
    global pending_caption
    if pending_caption:
        body(pending_caption, indent=False)
        pending_caption = None
while i < len(lines):
    line = lines[i].rstrip()
    if line.strip() == '<!-- pagebreak -->':
        PENDING_BREAK['on'] = True
        i += 1
        continue
    if line.lstrip().startswith('|'):
        block = []
        while i < len(lines) and lines[i].lstrip().startswith('|'):
            block.append(lines[i].strip())
            i += 1
        cells = [[c.strip() for c in r.strip('|').split('|')] for r in block]
        cells = [c for c in cells if not all(set(x) <= set('-: ') for x in c)]
        add_table(cells[0], cells[1:], caption=pending_caption)
        pending_caption = None
        continue
    if line.startswith('!['):
        flush_caption_as_paragraph()
        m = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        add_figure(m.group(2), m.group(1))
        i += 1
        continue
    if line.strip() == '':
        i += 1
        continue
    if line.startswith('#'):
        flush_caption_as_paragraph()
        level = len(line) - len(line.lstrip('#'))
        text = line.lstrip('#').strip()
        if level == 1 and first_h1:
            apply_break(doc.add_paragraph(text, style='Title'))
            first_h1 = False
        else:
            apply_break(doc.add_heading(text, level=min(level, 3) if level > 1 else 1))
            STATS['headings'].append(text)
        i += 1
        continue
    if line.startswith('- '):
        flush_caption_as_paragraph()
        body(line[2:], style='List Bullet', size=10.5)
        i += 1
        continue
    if line.startswith('> '):
        flush_caption_as_paragraph()
        p = body(line[2:], indent=False, size=10)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.right_indent = Cm(0.6)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        i += 1
        continue
    stripped = line.strip()
    if re.match(r'^表\d+\s', stripped):
        flush_caption_as_paragraph()
        pending_caption = stripped
        i += 1
        continue
    if re.match(r'^(注：|数据来源：|来源：|说明：|图\d+ )', stripped):
        body(stripped, indent=False, size=9, space_after=3)
        i += 1
        continue
    flush_caption_as_paragraph()
    body(line)
    i += 1

doc.core_properties.title = '二维光波导自动布线总体实验报告'
doc.core_properties.author = '李昊伦'
doc.core_properties.subject = '二维复现、半径与路径结构优化、几何修复、冻结对照与保护性优化'

for style in doc.styles:
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), EAST)
    fonts.attrib.pop(qn('w:eastAsiaTheme'), None)
defaults = doc.styles.element.find(qn('w:docDefaults'))
if defaults is not None:
    for fonts in defaults.iter(qn('w:rFonts')):
        fonts.set(qn('w:eastAsia'), EAST)
        fonts.attrib.pop(qn('w:eastAsiaTheme'), None)
parts = [doc.part]
for section in doc.sections:
    parts.extend([section.header.part, section.footer.part])
for part in parts:
    for run in part.element.iter(qn('w:r')):
        rpr = run.find(qn('w:rPr'))
        if rpr is None:
            continue
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            continue
        fonts.set(qn('w:eastAsia'), EAST)
        fonts.attrib.pop(qn('w:eastAsiaTheme'), None)
doc.save(OUT_DOCX)
(HERE / '_work' / 'docx_stats.json').write_text(
    json.dumps(STATS, ensure_ascii=False, indent=1), encoding='utf-8')
print(json.dumps({'docx': str(OUT_DOCX), 'tables': STATS['tables'],
                  'figures': STATS['figures'], 'headings': len(STATS['headings'])},
                 ensure_ascii=False))