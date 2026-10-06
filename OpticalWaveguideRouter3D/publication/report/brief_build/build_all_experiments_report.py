from pathlib import Path
import csv, json, hashlib, zipfile, re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[4]
P3, P2 = ROOT/'OpticalWaveguideRouter3D', ROOT/'OpticalWaveguideRouter2D'
OUT = Path(__file__).resolve().parent
NAME = '光波导布线全部实验报告'
IMAGES, TABLES, COVERAGE = [], [], {}
SECTION_NUMBER=0
PENDING_PAGE=False

def csvread(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))
def jsread(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))
def tabdata(name): return csvread(P3/'publication/tables'/name)
def raw(folder,n,name): return csvread(P3/f'outputs/{folder}/{n}/{name}')
def find(rows, scheme): return next(r for r in rows if r['scheme']==scheme)
def fmt(v,d=4):
    if v is None: return '缺失'
    if isinstance(v,str) and v in ('不适用','未完成','缺失'): return v
    return f'{float(v):.{d}f}'
def integer(v): return str(int(float(v)))

doc=Document(); sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.7)
sec.left_margin=sec.right_margin=Cm(2)
sec.header_distance=sec.footer_distance=Cm(0.7)
for name in ['Normal','Title','Heading 1','Heading 2','Caption','Header','Footer']:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    s.paragraph_format.widow_control=True
    for b in s.element.xpath('./w:pPr/w:pBdr'): b.getparent().remove(b)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.5
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(20); doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(10)
for name,size in [('Heading 1',14),('Heading 2',11.5)]:
    s=doc.styles[name]; s.font.size=Pt(size); s.font.bold=True
    s.paragraph_format.line_spacing=1.2
    s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(6)
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].font.italic=False; doc.styles['Caption'].font.bold=False
doc.styles['Caption'].paragraph_format.line_spacing=1.25
doc.styles['Caption'].paragraph_format.space_after=Pt(6)
h=sec.header.paragraphs[0]; h.text='光波导布线实验记录'; h.style='Header'; h.runs[0].font.size=Pt(9)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
f.add_run('第 '); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field); f.add_run(' 页')
for r in f.runs: r.font.size=Pt(9)

def para(text,label=None,small=False):
    p=doc.add_paragraph()
    if label: p.add_run(label+'：').bold=True
    p.add_run(text)
    if small:
        p.paragraph_format.line_spacing=1.3
        p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(5)
        for r in p.runs: r.font.size=Pt(9)
    return p
def heading(text,ids=(),level=1):
    global SECTION_NUMBER,PENDING_PAGE
    if level==1:
        SECTION_NUMBER+=1
        text=f'{SECTION_NUMBER} '+re.sub(r'^\d+\s+','',text)
    for code in ids:
        if code in COVERAGE: raise ValueError('Duplicate coverage '+code)
        COVERAGE[code]=text
    p=doc.add_heading(text,level)
    if (level==1 and SECTION_NUMBER>1) or PENDING_PAGE:
        p.paragraph_format.page_break_before=True; PENDING_PAGE=False
    return p
def page():
    global PENDING_PAGE
    PENDING_PAGE=True
def continuation(text):
    p=heading(f'实验 {SECTION_NUMBER} 续 {text}',level=2)
    p.paragraph_format.page_break_before=True
    return p
def note(text): return para('注：'+text,small=True)
def source(text):
    # Keep a source with the preceding explanation instead of leaving it alone.
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
    return para('数据来源：'+text,small=True)
def border(parent,side,val='nil',size=0):
    el=OxmlElement('w:'+side); el.set(qn('w:val'),val)
    if val!='nil': el.set(qn('w:sz'),str(size)); el.set(qn('w:color'),'000000')
    parent.append(el)

def table(title,headers,rows,widths,text_cols=(0,),groups=None,best=(),notes=None,row_groups=()):
    """Editable three-line table with values centered under their column headings."""
    assert len(headers)==len(widths) and sum(widths)<=17.001
    assert all(len(row)==len(headers) for row in rows)
    if not row_groups:
        if len(rows)<=5 or headers[0]=='路线ID':
            row_groups=[(0,len(rows))]
        elif headers[0] in ('规模','规模与ID'):
            starts=[0]+[i for i in range(1,len(rows)) if rows[i][0]!=rows[i-1][0]]+[len(rows)]
            row_groups=list(zip(starts,starts[1:]))
    number=len(TABLES)+1
    p=doc.add_paragraph(f'表{number} {title}',style='Caption')
    p.paragraph_format.keep_with_next=True; p.paragraph_format.space_after=Pt(6)
    hr=2 if groups else 1
    t=doc.add_table(rows=hr,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    pr=t._tbl.tblPr
    tw=pr.find(qn('w:tblW')); tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'),str(round(sum(widths)*1440/2.54)))
    for old in pr.findall(qn('w:tblBorders')): pr.remove(old)
    b=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        border(b,side,'single' if side in ('top','bottom') else 'nil',8)
    pr.append(b)
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for row in t.rows:
        for ci,cell in enumerate(row.cells): cell.width=Cm(widths[ci])
    for ci,(c,titletext) in enumerate(zip(t.rows[hr-1].cells,headers)): c.text=titletext
    if groups:
        covered=set()
        for start,end,titletext in groups:
            cell=t.cell(0,start).merge(t.cell(0,end)) if end>start else t.cell(0,start)
            cell.text=titletext; covered.update(range(start,end+1))
        for ci in range(len(headers)):
            if ci not in covered:
                cell=t.cell(0,ci).merge(t.cell(1,ci)); cell.text=headers[ci]
    for values in rows:
        for c,v in zip(t.add_row().cells,values): c.text=str(v)
    best=set(best)
    positions={}
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            positions.setdefault(cell._tc,[]).append((ri,ci))
    seen=set()
    for ri,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr(); rp.append(OxmlElement('w:cantSplit'))
        if ri<hr: rp.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            if cell._tc in seen: continue
            seen.add(cell._tc); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri>=hr or not groups: cell.width=Cm(widths[ci])
            cp=cell._tc.get_or_add_tcPr()
            for tag in ['tcBorders','shd','tcMar']:
                for old in cp.findall(qn('w:'+tag)): cp.remove(old)
            margins=OxmlElement('w:tcMar')
            for side,val in [('top',80),('bottom',80),('left',85),('right',85)]:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); margins.append(e)
            cp.append(margins)
            cb=OxmlElement('w:tcBorders')
            first=min(r for r,c in positions[cell._tc]); last=max(r for r,c in positions[cell._tc])
            for side in ['left','right','top','bottom']:
                is_top=side=='top' and first==0
                is_middle=side=='bottom' and last==hr-1
                is_bottom=side=='bottom' and last==len(t.rows)-1
                border(cb,side,'single' if is_top or is_middle or is_bottom else 'nil',8 if is_top or is_bottom else 5)
            cp.append(cb)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.2
                p.paragraph_format.keep_with_next=ri<hr
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci in text_cols else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size=Pt(10.5); r.font.bold=(ri<hr or (ri-hr,ci) in best)
    # Physical continuation cells in vertically merged headers also need the separator.
    for tc in t.rows[hr-1]._tr.findall(qn('w:tc')):
        cp=tc.find(qn('w:tcPr'))
        cb=cp.find(qn('w:tcBorders'))
        if cb is None: cb=OxmlElement('w:tcBorders'); cp.append(cb)
        for old in cb.findall(qn('w:bottom')): cb.remove(old)
        border(cb,'bottom','single',5)
    for tc in t.rows[hr]._tr.findall(qn('w:tc')):
        cp=tc.find(qn('w:tcPr')); cb=cp.find(qn('w:tcBorders'))
        for old in cb.findall(qn('w:top')): cb.remove(old)
        border(cb,'top','single',5)
    # Include the empty continuation cells of vertically merged headers.
    for row in t.rows[:hr]:
        for tc in row._tr.findall(qn('w:tc')):
            for element in tc.findall(qn('w:p')):
                pp=element.get_or_add_pPr()
                keep=pp.find(qn('w:keepNext'))
                if keep is None: keep=OxmlElement('w:keepNext'); pp.append(keep)
                keep.set(qn('w:val'),'1')
    # Keep each comparable group together while allowing breaks between groups.
    for start,end in row_groups:
        assert 0<=start<end<=len(rows)
        for row in t.rows[hr+start:hr+end-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next=True
    if notes:
        for cell in t.rows[-1].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next=True
    TABLES.append({'number':number,'title':title,'rows':len(rows),'width_cm':sum(widths),'best_cells':sorted(best)})
    if notes: note(notes)
    else:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=Pt(2)
    return t

def picture(paths,widths,caption):
    global FIG_NUMBER
    FIG_NUMBER+=1
    if isinstance(paths,(str,Path)): paths=[paths]; widths=[widths]
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(5); p.paragraph_format.keep_with_next=True
    for i,(path,w) in enumerate(zip(paths,widths)):
        path=Path(path); assert path.is_file()
        if i: p.add_run('  ')
        shape=p.add_run().add_picture(str(path),width=Cm(w)); shape._inline.docPr.set('descr',caption)
        IMAGES.append({'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    p=doc.add_paragraph(f'图{FIG_NUMBER} {caption}',style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
FIG_NUMBER=0

repro=tabdata('t32_reproduction_compare_data.csv')
radius=tabdata('t33_radius_sweep_512_data.csv')
s12={n:raw('opt2d',n,'comparison.csv') for n in (256,512)}
for rows in s12.values():
    for r in rows: r['scheme']=r['label'].split('-')[0]
s13={n:raw('opt2d_step13_fix',n,'comparison.csv') for n in (256,512)}
s14={n:raw('opt2d_step14',n,'comparison.csv') for n in (256,512)}

# Practical introduction and the user's actual reproduction images.
doc.add_paragraph('光波导布线全部实验报告',style='Title')
para('我们完成了二维布线复现、轨道与路径优化、交叉和间距诊断，以及三维抬层实验。下面逐项说明做了什么、结果怎么样、得出了什么结论，并保留未布通、被拒绝和中止的尝试。图与数据均来自自己的项目结果。')
heading('R5 和 R6 是什么',level=2)
para('R 是圆弧转弯半径，R5 = 5 mm，R6 = 6 mm。R6更缓，也占更多空间；模型的一次90°转弯损耗约2.3826/1.9037 dB。F5是R5自由弯角，F56是允许R5/R6的自由弯角。清单编号R05、R06分别指512条的2 mm、3 mm实验，与半径名称不同。')
heading('1 原二维波导布线复现',('R02','R03','R04'))
para('恢复原程序的输入、轨道分配与圆弧构造，在150 × 150 mm区域布置256、512条连接；分别计算直线、弯曲和交叉损耗。', '做了什么')
picture([P2/'results/fiberBoard256bend.png',P2/'results/fiberBoard512bend.png'],[8.15,8.15],'自己的二维复现布局 左为256条 右为512条 红线为波导')
rows=[]
for case,label in [('256','256条 R5'),('512','512条 R5'),('512_R4','512条 R4')]:
    a=next(r for r in repro if r['case']==case and r['metric']=='平均损耗'); b=next(r for r in repro if r['case']==case and r['metric']=='最大损耗')
    rows.append([label,fmt(a['repro_db']),fmt(b['repro_db']),fmt(a['thesis_db'],1),fmt(b['thesis_db'],1)])
table('复现损耗与原报告数值',['实验','平均','最大','平均','最大'],rows,[4.4,3.15,3.15,3.15,3.15],groups=[(1,2,'复现损耗（dB）'),(3,4,'原报告损耗（dB）')],notes='全部连接完成。原报告只保留1位小数；复现结果保留4位小数。不同规模或半径不作复现优劣排名。')
para('主要损耗数值接近原报告，256、512两组均全部布通。这里的损耗是模型计算值；交叉表来自原图数字化近似，尚无器件实测结果。', '结果和结论')
source('2D/results/fiberBoard{256,512}_loss_summary.json、fiberBoard512_loss_R4_summary.json；三线表取自t32。2D/3D目录缩写与t表编号见末页。')

page(); heading('2 弯曲模型与半径扫描',('R01','R05','R06'))
para('先用原程序的弯曲表重算2、3、4、5、6 mm半径的一次90°转弯损耗，再让512条连接分别以2、3、4、5 mm半径重新布线。', '做了什么')
bend=tabdata('t31_bend_model_thesis_vs_repro_data.csv')
table('一次90°转弯的模型核对',['半径（mm）','原印刷值（dB）','复算值（dB）','复算减印刷（dB）'],[[fmt(r['radius_mm'],0),fmt(r['thesis_db'],2),fmt(r['repro_db']),fmt(float(r['repro_db'])-float(r['thesis_db']))] for r in bend],[3.0,4.0,4.8,5.2],text_cols=(),notes='印刷值保留2位小数。差值有正有负，列名明确表示有符号差，不沿用旧CSV中误写的“abs”列名。')
table('512条连接的完整半径扫描',['半径（mm）','平均损耗（dB）','P95损耗（dB）','最大损耗（dB）'],[[integer(r['radius_mm']),fmt(r['mean_loss_db']),fmt(r['p95_loss_db']),fmt(r['max_loss_db'])] for r in radius],[3.1,4.6,4.6,4.7],text_cols=(),best=[(3,1),(3,2),(3,3)],notes='各半径独立构造几何，均完整布通。粗体是本组完整结果的最低损耗，不能解释为固定几何下只改模型参数。P95表示一次布局内95%的路线损耗不超过该值。')
continuation('损耗分布与结论')
picture(P2/'results/loss_vs_radius.png',16.5,'自己的半径扫描结果 左为损耗汇总 右为逐条路线分布')
para('平均损耗由16.5873降至5.5147 dB。R5中弯曲损耗占82.66%，直线占11.52%，交叉占5.81%。在这批输入和当前模型下，转弯损耗是主要来源，扩大半径是有效方向，但仍要留出足够空间。', '结果和结论')
source('t31、t33；2D/loss_model.py；2D/results/fiberBoard512_loss_R{2,3,4}_summary.json。')

heading('3 交叉表敏感性与原程序逐路核对',('R07','R09'))
para('原始交叉损耗数据表没有完整恢复，所以在同一512条R5几何上分别采用全90°估算、图线数字化并按90°文字值校准、图线原值。另在原Python 3.8字节码上实跑256条R5及512条R2/R3/R4/R5，逐条核对直线和弯曲项。', '做了什么')
cross=tabdata('t34_crossing_model_sensitivity_data.csv')
labels={'A':'全部按90°估算','B':'图线数字化加90°校准','C':'图线数字化原值'}
table('固定几何的交叉模型复算',['交叉表口径','平均损耗（dB）','最大损耗（dB）'],[[labels[r['scope']],fmt(r['mean_loss_db']),fmt(r['max_loss_db'])] for r in cross],[7.8,4.6,4.6],notes='这三行是不同模型假设，无算法最优排名。主报告采用有文字锚点的校准口径。')
para('三种口径的平均跨度0.0986 dB，最大跨度0.1436 dB，主要半径趋势不变。字节码对照的五组逐路直线加弯曲损耗最大绝对差均为0。原字节码运行用的是全0交叉占位表，因此它只能验证非交叉项。', '结果')
para('二维几何和非交叉损耗的复现有直接逐路证据；交叉项仍含近似，不能把近似表说成恢复了原始测量数据。损耗分布、与长度及交叉次数的散点图是同一批复现结果的分析，不是新的输入实验。', '结论')
para('端到端输出另进行了256/512两组各36项验收，GDS重读分别得到1个cell和256/512条路径。512条与原字节码逐阶段对照，端点、轨道、全部弯曲参数以及GDS点阵和宽度均512/512一致，只有8个时间戳字节不同。', '输出与几何保真核对')
para('还对随附的历史图片进行了差异诊断：只有266/512条与当前恢复结果一致，235条相差一根轨道。把sn_calc条件由大于改为大于等于没有影响；去掉sn_calc项后仅43/512一致。这些负假设不能解释图差异，证据表明历史图片与恢复EXE来自不同构建。', '历史图差异与负假设')
source('t34；2D/data/crossing_loss_from_thesis_fig3_12.csv；2D/scratch/legacy_loss/legacy256_R5.json及legacy512_R{2,3,4,5}.json。')
source('2D/docs/exact_legacy_fidelity_report.md与migration_report.md。')

# Early completed verification and diagnostic experiments.
page(); heading('4 早期输入与二维骨架验证')
para('检查256/512连接输入，恢复512条连接的快照（共1024个端点）与分类，再执行独占轨道分配和直角骨架构造。轨道策略文档中的B、C只是设计，不当作已完成实验。', '做了什么')
table('真实输入和快照恢复检查',['输入规模','PMT数量','端点每PMT','自连接','额外重复行'],[['256','32','16','0','134'],['512','64','16','0','268']],[3.1,3.3,3.7,3.0,3.9],text_cols=(),notes='重复行是多重连接，保留原输入，不去重。512后半工作簿数据使用公式缓存。')
para('512快照关联512行，371行方向相同、141行恢复方向；64个PMT各16个唯一端点，间距0.175 mm。分类得到上侧U112、下侧U112、Z288，恢复输入方向后两类Z各144，非法端点0。', '输入与分类结果')
table('早期512条轨道分配',['配置','候选轨道（条）','普通路线（条）','未支持特殊Z（条）'],[['独占轨道 A G1 S1','800','454','58']],[6.2,3.6,3.6,3.6],notes='只完成454/512条，不能与后续完整512条结果直接排名。特殊Z横向跨度不足2R，没有出现普通轨道耗尽。')
table('454条零宽直角骨架的全局检查',['已检路线对','段交叉事件','交叉路线对','接触','重合','自交'],[['102831','47292','41147','0','0','0']],[3.3,3.3,3.3,2.4,2.4,2.3],text_cols=(),notes='454条中224条U形、230条Z形。端点错误、斜段、重复点、轨道不符均为0；通过骨架检查不等于平滑版图完成。')
para('普通骨架与逐对检测能够运行，但58条跨度不足2R的特殊连接需要专门路径。后续补齐普通圆角与特殊Z解析几何，才建立完整512条工程基线。', '结果和结论')
source('3D/docs/reports/step_8_legacy_512_track_assignment_v01.md、step_8_legacy_512_routes_v01.md、step_8_legacy_512_collision_validation_v01.md。')
source('输入和快照：step_8_legacy_input_validation.md、step_8_legacy_512_snapshot.md、step_8_legacy_512_route_preparation.md，均在3D/docs/reports/。')

page()
heading('5 精确求交与多交叉诊断',('R08','G3D-10'))
para('先把454条普通骨架全部圆弧平滑，得到1362条直段和908个四分之一圆弧；再独立构造58条特殊Z的双弧几何，得到116个圆弧和116条直段。特殊Z的新几何不是原EXE特殊路径的忠实恢复，原分配器的58条未支持状态仍保留。', '平滑和特殊Z实验')
para('解析求交新增直与直、直与弧、弧与弧及自交测试54项，当时全测试281项通过。完整512条共130816对，原段级统计交叉55932、交叉对49499、touch12、overlap0、自交0；12个段级touch来自3对特殊弧接点，归并后成为3个物理cross，物理touch为0。', '求交基础与事件归并')
para('对这套512条工程解析路线按直线与圆弧精确求交，合并同位置的段级交点，分别统计物理交叉、双交叉和相邻端口之间的多交叉。', '做了什么')
para('完整平滑基线得到55935个物理交叉事件、49502对交叉路线、最小交叉角约2.84°。这与二维原版逐路整数角统计属于不同口径；原报告的14°不能直接拿来判断精确求交是否正确。', '结果')
para('诊断发现6433对路线出现双交叉，318个端部多交叉结构全部涉及圆弧。多交叉的954个交点按段组合归因为：弧与水平303、弧与弧31、弧与竖直331、水平与竖直289。求交口径与端口拓扑会影响诊断数量。', '结论')
table('多交叉和双交叉的分类审计',['诊断项','类别','计数'],[['多交叉','纯普通路线','222'],['多交叉','含特殊Z','96'],['多交叉','阈值边界待定','2'],['双交叉对','普通与普通','6092'],['双交叉对','普通与特殊','341'],['双交叉对','U与U','3735'],['双交叉对','Z与Z','2698']],[4.4,8.2,4.4],text_cols=(0,1),notes='多交叉使用0.125 mm间距和1e-9容差。两类双交叉分解是不同分类视角，不可把全部行求和。阈值外841477项不能直接解释为安全。')
para('普通路线的6092个双交叉对全部在直角骨架中已经存在，本组双交叉并非由圆角才创造。这个早期基线的非交叉损耗平均5.2646、最大6.1874 dB，采用整段传播与印刷弯曲表，未加入交叉损耗；不能和后续总损耗直接相减。', '损耗与拓扑核对')
source('3D/outputs/step_8_5_legacy_512_{physical,multi_crossing,double_cross,loss}_summary.json；step_8_5_m1_5_multi_attribution_summary.json。')

page(); heading('6 排序反转与硬约束负结果')
para('把112条上侧U形路线的主排序反转，保持其他路线不变；再测试遇到端部多交叉就拒绝候选的硬约束，并核验拒绝原因。', '做了什么')
table('上侧U形排序的A B对照',['指标','原排序A','反转排序B'],[['双交叉路线对','6433','4530'],['上侧U与上侧U的双交叉对','1903','0'],['物理交叉事件','55935','52129'],['小于20°交叉','1508','984'],['端部多交叉结构','318','345']],[8.2,4.4,4.4],notes='每行只比较同一指标。反转改善部分交叉指标，却使多交叉结构增加，所以不标总体最优。')
para('多交叉审计显示204个稳定、114个移除、141个新增，净增27。新增中的87个阈值余量大于0.05 mm，不能把增加全部归因于数值临界误差。', '结果')
para('硬约束的初版全量诊断被中止。快速版运行180.018 s，只提交217条普通路线，21条出现真实轨道耗尽，累计14696次候选拒绝；最终全板阶段没有完成。独立核验的14384条耗尽拒绝均有有效见证。', '结果')
table('硬约束的局部速度基准',['指标','原局部核验','快速局部核验'],[['30次重复的中位耗时（s）',fmt(0.00390345,6),fmt(0.00052845,6)]],[7.8,4.6,4.6],notes='只针对固定40条路线的局部核验，约7.39倍加速，不能外推为整个512条布线加速。全局损耗和最终多交叉数为不适用，因为全板没有完成。')
para('改变排序有局部收益；当前贪心硬约束无法同时保持完整布通。加速求交不会自动解决轨道耗尽，未完成的方案没有可比较的全板损耗。', '结论')
source('3D/outputs/step_8_5_top_u_order_ab_summary.json、top_u_multi_crossing_delta_summary.json、exact_multi_guard_fast_routing.json、exact_multi_guard_exhaustion_summary.json。以上文件名均有step_8_5_前缀。')

page(); heading('7 局部单线恢复沙箱与归因诊断')
para('先实现Dynamic D端点视图，用4类合成样本和4条真实路线检查每侧最近端点与连续长度，458项测试通过。它只是工程诊断，危险侧尚未确认，没有整数K或新的布线改善结果。', '端点长度诊断')
para('对318个多交叉进一步按局部结构归因：204个类似原文结构、23个有歧义、91个不类似；统计最低候选路线时241个目标唯一、77个并列，最低仅普通270、仅特殊29、普通与特殊混合19，没有全部不可移动目标。统计可以帮助选候选，但尚未证明某个顺序最优。', '结构归因与候选统计')
para('从同一个318多交叉的基线分别选择四类局部案例，每例尝试347个候选，并独立验证；这些候选保留在沙箱中，没有写回全板。', '做了什么')
table('四类局部恢复的独立结果',['结构案例类型','尝试候选','严格候选','多交叉前','多交叉后'],[['A 类似原文结构 纯普通','347','0','318','318'],['B 跨侧 不类似原文结构','347','331','318','317'],['C 弧与弧案例','347','342','318','316'],['D 含特殊路径载体','347','344','318','316']],[7.0,2.5,2.5,2.5,2.5],notes='四行独立从同一基线出发，不是连续优化，收益不能相加；普通受试线均复用同一R5构造器。全板端点、损耗与布通尚未联合验收。')
para('部分局部轨道替换能减少1或2个多交叉。本次类似原文结构的纯普通案例没有找到严格候选；这些结果尚未合并成全板恢复，不能宣称多交叉已消除。', '结果和结论')
para('四个原位负对照均保持318，目标仍在且差分为空。其他候选也出现过目标消失却总数不降、甚至总数增加，因此只消除目标不足以接受改动。', '负对照')
source('3D/outputs/step_8_5_m1_6_probe_summary.json与step_8_5_m1_6_independent_validation.json。')
source('3D/docs/reports/step_8_5_m1_dynamic_d_diagnostic.md；outputs/step_8_5_m1_5_multi_attribution_summary.json、step_8_5_m1_6_victim_statistics.json。')

# Step 12 complete main results and all saved ablations.
page(); heading('8 候选轨道与自适应半径主实验',('S12-01',))
para('在256和512条输入上比较A原版R5、B统一R6、C固定R5候选轨道搜索、D允许R5/R6选择。这个早期阶段允许同端口槽位调整；后续实验才逐路固定端点。损耗用解析求交统计，与原版复现的小数略有差别。', '做了什么')
rows=[]; best=[]
for n in (256,512):
    eligible=[r for r in s12[n] if int(r['unplaced_count'])==0]
    minima={key:min(float(r[key]) for r in eligible) for key in ['mean_loss_db','p95_loss_db','max_loss_db']}
    for r in s12[n]:
        ri=len(rows); rows.append([str(n),r['scheme'],integer(r['unplaced_count']),fmt(r['mean_loss_db']),fmt(r['p95_loss_db']),fmt(r['max_loss_db']),fmt(r['runtime_s'],2)])
        if r in eligible:
            for col,key in [(3,'mean_loss_db'),(4,'p95_loss_db'),(5,'max_loss_db')]:
                if abs(float(r[key])-minima[key])<1e-10: best.append((ri,col))
table('两种规模的A B C D主对照',['规模','方案','未布通','平均','P95','最大','耗时（s）'],rows,[1.5,1.5,1.7,3.0,3.0,3.0,3.3],text_cols=(1,),groups=[(3,5,'损耗（dB）')],best=best,notes='只在同规模且完整布通的方案中标粗最低损耗。512条B缺2条，损耗仅统计510条，排除排名；运行时间是保存的一次运行，不是重复计时。')
para('256条统一R6完整布通，平均4.3360 dB；512条统一R6缺2条，自适应半径D则全部布通，平均4.5593 dB。固定R5的轨道搜索只有约0.0014/0.0015 dB平均改善。', '结果')
para('半径策略的改善远大于当前候选轨道搜索；512条需要小半径回退才能保证完整连接。这个阶段改变了部分槽位，结论属于早期探索。', '结论')
source('3D/outputs/opt2d/{256,512}/comparison.csv；S12-01。')

page(); heading('9 轨道优化的全部消融',('S12-02','S12-03','S12-04','S12-05','S12-06'))
para('依次改变候选数量、布线顺序、拆线重布、位置惩罚和半径策略，其余采用保存的默认配置。每种设置分别在256和512条输入上执行。', '做了什么')
abl=tabdata('t52_step12_ablation_data.csv')
names={'cand_k1':'候选 K = 1','cand_k4':'候选 K = 4','cand_k8':'候选 K = 8','cand_k16':'候选 K = 16','order_span':'顺序 按跨度','order_congestion':'顺序 按拥挤度','refine_on':'拆线重布 开','penalty_0':'位置惩罚 0','penalty_0.001':'位置惩罚 0.001','penalty_0.01':'位置惩罚 0.01','penalty_0.05':'位置惩罚 0.05','radius_fixed_R6':'半径 固定R6','radius_adaptive':'半径 自适应'}
rows=[[names[r['config']],fmt(r['mean_256']),integer(r['unplaced_256']),fmt(r['runtime_256'],2),fmt(r['mean_512']),integer(r['unplaced_512']),fmt(r['runtime_512'],2)] for r in abl]
best=[(3,1),(1,4),(2,4),(3,4),(8,1),(9,4),(11,1),(12,1)]
table('全部保存的轨道消融结果',['配置','平均（dB）','未布通','时间（s）','平均（dB）','未布通','时间（s）'],rows,[3.8,2.4,1.65,2.3,2.4,1.65,2.8],groups=[(1,3,'256条连接'),(4,6,'512条连接')],best=best,notes='粗体只在候选数量、位置惩罚或半径策略的各自组内、同规模且完整布通的结果中比较；跨配置组不排名。拆线重布关闭对应K = 8默认结果。')
para('候选1到16的平均收益不到0.002 dB，时间增加。跨度顺序缺76/144条，拥挤度顺序缺33/99条；这两种顺序都没有保持完整。拆线重布没有进一步改善。位置惩罚为0时缺59/113条；256的0.001和512的0.01保留完整且损耗较低。512固定R6消融缺2条，自适应完整。', '结果')
para('更多候选、改变顺序或拆线重布都不保证改善。在当前候选预筛、惩罚和预算下，搜索收益较小；未布通配置不作为成功优化结果。', '结论')
source('t52；3D/outputs/opt2d/{256,512}/sensitivity.json与sensitivity/*/summary.json。')

# Step 13 all six corrected schemes.
page(); heading('10 固定端点与自由弯角六方案',('S13-01','S13-02','S13-07'))
para('逐路固定起终点，比较原版A、固定R5的U形搜索R5U、优先R6且失败回退R5的D0、自适应U形D56、固定R5自由弯角F5、自适应自由弯角F56。跨侧路线允许圆弧加斜线的S形，同侧保留U形。', '做了什么')
rows=[]
for n in (256,512):
    for r in s13[n]: rows.append([str(n),r['scheme'],fmt(r['mean_loss_db']),fmt(r['p95_loss_db']),fmt(r['max_loss_db']),fmt(r['runtime_s'],2)])
table('补修后固定端点的全部六方案',['规模','方案','平均','P95','最大','耗时（s）'],rows,[1.6,2.4,3.3,3.3,3.3,3.1],text_cols=(1,),groups=[(2,4,'损耗（dB）')],best=[(5,2),(5,3),(5,4),(11,2),(11,3),(11,4)],notes='全部256/512条均完整连接、端点零变化。粗体按规模标识本组整体方案最低损耗；方案间半径策略不同，不把总改善都归因于路径形状。')
para('F56的平均损耗为2.9703/3.2947 dB，最大为5.0783/5.4995 dB。D0到D56只有约0.0421/0.0140 dB平均收益，搜索时间明显增加。F5降低平均损耗，但512条最大6.4459 dB，只比原版低约0.1314 dB。', '结果')
para('逐路核对D56与F56的半径完全相同：256条全部R6；512条510条R6、2条R5。相同半径下，U形到自由弯角使平均再降1.3690/1.2661 dB，支持减少总转弯角度的作用。', '结论')
para('历史256条F56曾把路线7/31重合152.007 mm误判成交叉。补修后48条路线结果变化，重合归零，平均由2.9689变为2.9703 dB；512条补修前后逐项不变。上表采用补修数据，历史非法结果不参与排名。', '补修记录')
source('3D/outputs/opt2d_step13_fix/{256,512}/comparison.csv、radius_compare.csv；256/per_route_delta.csv；历史记录在opt2d_step13/。')

page(); heading('11 自由弯角的损耗分解与几何代价')
rows=[]
for n in (256,512):
    for r in s13[n]: rows.append([str(n),r['scheme'],fmt(r['mean_straight_loss_db']),fmt(r['mean_bend_loss_db']),fmt(r['mean_crossing_loss_db']),fmt(r['mean_bend_total_deg'],2)])
table('六方案平均损耗分解',['规模','方案','直线','弯曲','交叉','平均总转角（°）'],rows,[1.6,2.0,3.2,3.2,3.2,3.8],text_cols=(1,),groups=[(2,4,'平均损耗分量（dB）')],notes='不同分量分别描述代价，不按单项最低值给方案作总体排名。')
rows=[]
for n in (256,512):
    for r in s13[n]: rows.append([str(n),r['scheme'],integer(r['unique_crossing_events']),integer(r['crossings_under_10deg']),integer(r['spacing_violation_count']),integer(r['contact_overlap_count'])])
continuation('交叉与间距代价')
table('六方案交叉与间距结果',['规模','方案','交叉事件','小于10°','间距不足对','重合'],rows,[1.6,2.0,3.5,3.0,4.0,2.9],text_cols=(1,),notes='交叉和小角列均为唯一交叉事件数，间距列为路线对。旧段级touch数不作为物理接触数。重合为零仍可有间距不足。')
para('512条A到F56平均弯曲损耗4.5587→2.1384 dB；交叉损耗0.3219→0.5650 dB，间距不足692→1231对。自由弯角降低弯曲损耗，也增加交叉与间距代价。', '结果和结论')
source('3D/outputs/opt2d_step13_fix/{256,512}/comparison.csv；Step14对段级接触重新归并为物理事件。')

heading('12 最差路线的同一编号跟踪',('S13-06',))
picture([P3/'outputs/opt2d_step13/512/figures/worst_route_A.png',P3/'outputs/opt2d_step13/512/figures/worst_route_F56.png'],[8.15,8.15],'自己的最差路线局部图 左为A路线287 右为F56路线34')
para('两张保存图分别选择各方案最差路线，编号不同；灰色为周围路线，红色为局部关注段。下面另固定原版最差路线编号，比较同一连接的损耗变化。', '做了什么')
rows=[]
for n,rid in [(256,'44'),(512,'287')]:
    for name in ['A','R5U','D0','D56','F5','F56']:
        r=next(r for r in raw('opt2d_step13_fix',n,f'{name}/per_route.csv') if r['route_id']==rid)
        rows.append([f'{n} / {rid}',name,fmt(r['total_loss_db']),fmt(r['bend_loss_db']),fmt(r['crossing_loss_db']),fmt(r['bend_total_deg'],0)])
continuation('同一连接的逐路结果')
table('固定原版最差连接的逐路结果',['规模与ID','方案','总损耗','弯曲损耗','交叉损耗','总转角（°）'],rows,[2.7,1.8,3.2,3.2,3.2,2.9],text_cols=(0,1),groups=[(2,4,'逐路损耗（dB）')],best=[(5,2),(11,2)],notes='粗体只表示同一个连接在六方案中的最低总损耗。图中的F56最差ID 34与表中的原版最差ID 44/287不同。',row_groups=[(0,6),(6,12)])
para('原版最差连接44/287从6.3333/6.5773降至约3.7512/4.3308 dB，总转角180°→90°。但F56的新最差连接换成34，最大损耗为5.0783/5.4995 dB，必须同时看固定ID和全局最大。', '结果和结论')
source('3D/outputs/opt2d_step13_fix/{256,512}/{scheme}/per_route.csv；原跟踪文件worst_route_tracking.csv。')

heading('13 自适应半径条件下的约束与惩罚扫描',('S13-03','S13-05'))
para('在base上分别加入间距惩罚spacing、小角惩罚small及二者both，重新选择路径；另对F5、F56扫描间距惩罚0、0.02、0.05 dB。此阶段半径仍自适应，不能当作纯固定半径的约束对照。', '做了什么')
con=tabdata('t63_step13_constrained_data.csv')
table('未冻结半径的四配置结果',['规模','配置','平均（dB）','最大（dB）','小于20°','间距不足对'],[[r['channels'],r['config'],fmt(r['mean_loss_db']),fmt(r['max_loss_db']),integer(r['crossings_under_20_deg']),integer(r['spacing_violation_count'])] for r in con],[1.6,2.2,3.3,3.3,3.1,3.5],text_cols=(1,),notes='重新布线并自适应半径，不是固定几何只改分数。作为早期权衡实验，不按后续严格保护口径评为最终方案。')
rows=[]
for n in (256,512):
    for name in ['F5','F56']:
        r=find(s13[n],name)
        rows.append([str(n),name,'0.00',fmt(r['mean_loss_db']),fmt(r['max_loss_db']),integer(r['spacing_violation_count'])])
        for r in raw('opt2d_step13_fix',n,f'penalty_ablation_{name}.csv'):
            rows.append([str(n),name,fmt(r['spacing_penalty_db'],2),fmt(r['mean_loss_db']),fmt(r['max_loss_db']),integer(r['spacing_violation_count'])])
continuation('间距惩罚扫描')
table('两个路径族的全部间距惩罚扫描',['规模','方案','惩罚（dB）','平均（dB）','最大（dB）','间距不足对'],rows,[1.6,2.0,2.9,3.5,3.5,3.5],text_cols=(1,),notes='每个规模、路径族独立扫描；损耗和间距存在权衡，故不定义单一最优。0.00行取补修后的默认结果。',row_groups=[(0,3),(3,6),(6,9),(9,12)])
para('256条F5间距不足576→436→401对，F56为415→324→288对；损耗小幅变化，并非全部增加。512条惩罚也减少间距不足，但不能保证最大损耗和关键路线同步改善。后续冻结半径实验进一步检验纯约束作用。', '结果和结论')
source('t63；3D/outputs/opt2d_step13_constrained/{256,512}/constrained_comparison.csv；opt2d_step13_fix/{256,512}/penalty_ablation_F{5,56}.csv。')

heading('14 六路径方案的小角交叉压力实验',('S13-04',))
para('不重新布线，在各方案的固定几何上，把小于20°的交叉损耗分别乘以1、2、5、10，再复算最大损耗。倍率是人为压力假设，用来检查交叉表不确定性的影响。', '做了什么')
sens=tabdata('t64_step13_sensitivity_data.csv')
table('补修后六方案的最大损耗压力结果',['规模','方案','倍率1','倍率2','倍率5','倍率10'],[[r['channels'],r['scheme'],*[fmt(r[k]) for k in ['max_x1','max_x2','max_x5','max_x10']]] for r in sens],[1.6,2.2,3.3,3.3,3.3,3.3],text_cols=(1,),groups=[(2,5,'各倍率情景的最大损耗（dB）')],notes='只在相同规模、相同倍率下有比较意义。此处展现敏感性而不评定统一最优；倍率5/10不代表已验证的真实误差范围。')
para('512条F56最大损耗从5.4995增到8.4333 dB，倍率10时超过原版A的6.5773 dB；256条F56从5.0783增到5.9941 dB。自由弯角的标称收益会受极小角交叉模型影响，不能从标称改善直接推断所有压力情景都改善。', '结果和结论')
source('t64；3D/outputs/opt2d_step13_fix/{256,512}/sensitivity.csv。')

# Step 14 all configurations, protected fixed IDs, margins, probes and pressure tests.
heading('15 冻结半径的六配置与严格验收',('S14-01','S14-02','S14-08','S14-09'))
para('固定每条路线的端点和F56半径，比较base、spacing、small、touch、both及保护性局部优化opt_base。要求平均增加不超过0.05 dB、标称全局最大不恶化、固定11条保护路线逐条不恶化，并保持几何合法和无重合。', '做了什么')
acc={n:raw('opt2d_step14',n,'acceptance.csv') for n in (256,512)}
rows=[]; best=[]
for n in (256,512):
    for r in s14[n]:
        a=find(acc[n],r['scheme']); ri=len(rows)
        rows.append([str(n),r['scheme'],fmt(r['mean_loss_db']),fmt(r['p95_loss_db']),fmt(r['max_loss_db']),integer(r['spacing_violation_count']),'通过' if a['ok']=='True' else '未通过'])
        if (n==512 and r['scheme']=='opt_base') or (n==256 and r['scheme'] in ['base','touch','opt_base']):
            best.extend([(ri,2),(ri,4)])
table('固定半径后的全部配置及验收',['规模','配置','平均','P95','最大','间距不足对','验收'],rows,[1.5,2.5,2.8,2.8,2.8,2.5,2.0],text_cols=(1,6),groups=[(2,4,'损耗（dB）')],best=best,notes='粗体只比较同规模、通过严格验收的配置；256条三行结果相同。P95独立描述尾部，512条opt_base的P95略有增加。未通过配置的单项较低值不评为最终最优。')
para('spacing、small、both在两种规模均未通过。256条的拒绝主要来自个别保护路线轻微恶化；512条三者平均增加分别0.0677、0.0545、0.1081 dB，都超过0.05 dB预算，且有保护对象恶化。touch与base相同，因为两组物理接触为0。', '结果')
para('冻结半径核对为256/256、512/512一致；base和未冻结F56参照的数值一致。D56参照平均4.3393/4.5607、最大5.3462/5.6217 dB；它保留U形，因此属于路径族参照，不是本组六个约束配置之一。只有512条保护优化产生了通过全部规则的非零改善。', '结论')
source('3D/outputs/opt2d_step14/{256,512}/comparison.csv、acceptance.csv、radius_freeze_check.csv、references.csv。')

for n in (256,512):
    heading(f'16 {n}条固定保护对象的逐路结果',('S14-04',) if n==256 else ())
    para('固定同一组11条ID，包含base最差、D56最差及被穿越的U形连接。每个配置都评价这些原ID，不能逐方案重选最差十条。表中列出相对base的损耗差，正数为恶化，负数为改善。', '做了什么')
    by={name:raw('opt2d_step14',n,f'{name}/protection_routes.csv') for name in ['base','spacing','small','both','opt_base']}
    rows=[]
    for base in by['base']:
        rid=base['route_id']; values=[]
        for name in ['spacing','small','both','opt_base']:
            r=next(r for r in by[name] if r['route_id']==rid); values.append(fmt(r['delta_vs_base_db']))
        rows.append([rid,fmt(base['total_loss_db']),*values])
    table(f'{n}条实验固定11条连接的损耗变化',['路线ID','base损耗（dB）','spacing','small','both','opt_base'],rows,[2.0,3.0,3.0,3.0,3.0,3.0],text_cols=(),groups=[(2,5,'相对base的差值（dB）')],notes='差值保留4位小数；显示0.0000可能是小于舍入精度的数值，不表示用0替代缺失。touch差值全部为0，与base一致，省去重复列。不对不同路线ID作损耗排名。')
    if n==256:
        para('spacing让185恶化约0.0072 dB，small让186、185、177恶化，both让52恶化约0.0021 dB。虽然多数保护对象改善，每种约束仍有至少一条恶化，所以零余量验收拒绝。opt_base没有接受候选，全部保持原值。', '结果和结论')
    else:
        para('spacing和both对311分别恶化约0.3885、0.3376 dB；small对152恶化约0.0238 dB。opt_base改善7条、其余不恶化，满足逐路保护。集合均值或全局最大改善，不能代替逐路验收。', '结果和结论')
    source(f'3D/outputs/opt2d_step14/{n}/protection_set.json与各配置/protection_routes.csv。')
    if n==256:
        heading('保护余量诊断',('S14-03',),level=2)
        para('对固定保护集合扫描允许恶化0、0.001、0.002、0.005、0.01、0.02 dB，不改变已经生成的几何。256条both在0.005 dB时达到逐路余量要求，spacing和small在0.01 dB时达到；512条small、both、spacing的最大恶化分别0.0238、0.3376、0.3885 dB，在本次扫描的全部余量下仍未通过逐路保护。')
        para('这是对保护条件的诊断，任务仍采用0 dB余量；“通过余量”只描述保护集合，并不自动满足平均预算等其他条件。', '结论')
        source('3D/outputs/opt2d_step14/{256,512}/protection_margins.csv。')

heading('17 保护优化过程与放宽条件探针',('S14-05','S14-06'))
para('每个目标路线最多试4个候选，目标包含固定保护集合和压力情景的高损耗路线。逐候选评价平均预算、全局最大、11条保护路线和倍率5情景最大；失败则完整回滚。另改变候选排序强度及保护余量，检查拒绝限制。', '做了什么')
rows=[]
for n in (256,512):
    rec=jsread(P3/f'outputs/opt2d_step14/{n}/optimize_record.json')
    # Main counts and runtime are recorded by the actual optimizer, not comparison runtime.
    attempts=raw('opt2d_step14',n,'optimize_attempts_base.csv')
    runtime=233.1024483999936 if n==256 else 762.2401490000193
    accepted=sum(r.get('accepted')=='True' for r in attempts)
    rows.append([str(n),'严格主实验','0.00',str(len(attempts)),str(accepted),fmt(runtime,2),fmt(find(s14[n],'opt_base')['scenario5_max_loss_db'])])
    for r in raw('opt2d_step14',n,'probe/probe_summary.csv'):
        label=f"排序{fmt(r['rank_small_penalty_db'],1)} 余量{fmt(r['protected_slack_db'],2)}"
        rows.append([str(n),label,fmt(r['protected_slack_db'],2),integer(r['candidates_tried']),integer(r['accepted']),fmt(r['runtime_s'],2),fmt(r['scenario5_max_loss_db'])])
table('实际优化和全部保存的诊断探针',['规模','变体','余量（dB）','尝试','接受','耗时（s）','倍率5最大（dB）'],rows,[1.5,4.2,2.0,1.2,1.2,2.9,4.0],text_cols=(1,),notes='诊断预算与候选排序有差别，不按接受数或耗时跨行排优。256的strict_slack005实际参数是0.05 dB，名称不等于0.005。全部回滚指纹不一致次数均为0。')
para('严格主实验256条接受0/68，512条接受6/74，512平均3.2947→3.2936、最大5.4995→5.4654、倍率5最大5.7264→5.6340 dB。256条主要拒绝原因包括情景无改善62次、最大恶化32次及保护对象恶化；一个候选可同时触发多条拒绝，计数不能相加成总候选数。', '结果')
para('256条把排序强度调到0.5或2.0仍0接受；保护余量0.01/0.05时各接受2个候选，属于放宽条件诊断。512条严格排序0.5和余量0.01探针均接受5/33和5/28，情景最大同为约5.6392 dB，不能把“严格排序0接受”的256结论推广到512。', '结论')
source('3D/outputs/opt2d_step14/{256,512}/optimize_attempts_base.csv、optimize_record.json、probe/probe_summary.csv。')

heading('18 冻结半径后全部压力情景',('S14-07',))
para('保留六配置已经生成的几何，分别把小于20°交叉损耗乘以1、2、5、10，重新计算最大损耗；不重新优化各情景。', '做了什么')
sens=tabdata('t74_step14_sensitivity_data.csv')
table('六配置的全倍率最大损耗',['规模','配置','倍率1','倍率2','倍率5','倍率10'],[[r['channels'],r['scheme'],*[fmt(r[k]) for k in ['max_x1','max_x2','max_x5','max_x10']]] for r in sens],[1.6,2.2,3.3,3.3,3.3,3.3],text_cols=(1,),groups=[(2,5,'各倍率情景最大损耗（dB）')],best=[(11,2),(11,3),(11,4),(6,5),(9,5)],notes='仅在同规模、同倍率且已通过严格验收的base/touch/opt_base之间标粗有改善或更低的结果；256条三者相同，不标重复最优。spacing/small/both仍保留未验收状态。')
para('512条opt_base在倍率1、2、5时改善，但倍率10最大从8.4333变为8.4996 dB，略有恶化。small在倍率10的最大较低，但仍违反平均预算和固定保护条件。256条opt_base没有改动，全部压力结果与base一致。', '结果和结论')
heading('小角计数与间距权衡',level=2)
para('冻结半径后，256条small把小于20°交叉657→315，both降至371；512条small从2939→1125，both降至1209。spacing在512条虽然把间距不足1231→821，却使小于20°交叉增到3263。不能用一项改善代表所有风险改善。')
source('t74；3D/outputs/opt2d_step14/{256,512}/sensitivity.csv、comparison.csv。')

# Three-dimensional work, including all foundation steps and all executed scale experiments.
heading('19 三维几何与距离检测验证',('G3D-01','G3D-02','G3D-03'))
para('建立三维路线容器和局部抬层几何；用余弦过渡连接不同层，检查端点与切向连续；对直线、圆弧和过渡段的最小距离进行样本验证。', '做了什么')
table('三维基础验证记录',['验证项','真实样本路线（条）','独立差分点（个）','全测试通过（项）'],[['9A 路线容器与抬层','6','不适用','545'],['9B 曲率解析与数值','不适用','7','579'],['9C clearance分类','不适用','不适用','623']],[6.0,3.8,3.6,3.6],notes='9C使用6对路线和150对几何段。基础验证支持几何实现，不是全板路由结果。不同阶段测试数包含当时已有测试，不相加。')
para('9A将普通0/1/2、特殊13/14/15抬到2.75 mm，XY误差0、平面段长度不变。9B的Lxy = 10 mm、高度差1 mm过渡，真实最小半径20.2642 mm、有效参数25.25 mm；7点差分最大曲率误差6.356 × 10⁻¹² mm⁻¹，579项测试通过。')
para('9C在0.01 mm阈值下，同层0/24、0/25、0/26为碰撞；0/1、0/2、0/3间距0.175、0.350、0.525 mm为分离，另覆盖异层、阈值和过渡内部最近点，623项测试通过。未决不当作分离。', '距离验证')
para('几何容器、切向连接和距离检测能用于后续抬层；有效半径不等于真实最小半径，切向连续也不保证曲率连续。此验证没有计算三维器件传播损耗。', '结论')
source('3D/outputs/step_9_a_tests.json、step_9_b_curvature_join_validation.json、step_9_c_clearance_validation.json。')
heading('20 单路线抬层探针',('G3D-04',))
para('选3组真实相交路线对，逐次抬高其中一条；每次再检查该路线与其余511条邻居，测试能否消除目标近距对且不引入新对。阈值0.1 mm。', '做了什么')
table('三个真实路线对的独立抬层结果',['目标路线对','移动ID','碰撞邻线前','碰撞邻线后','新增邻线','增长（mm）'],[['0 / 24','24','230','217','0','0.243892'],['0 / 25','0','95','52','0','0.243892'],['0 / 26','26','230','217','0','0.243892']],[3.0,2.0,3.2,3.2,2.8,2.8],text_cols=(0,),notes='每组各45候选，20个通过目标检查、25个因自身判定有歧义拒绝。三个案例从各自原状态出发，不能同时叠加。')
para('3组均校验511个邻居，零新增碰撞，总耗时约68.6 s。合成十字探针从0→1→0 mm后，目标间距1 mm、额外长度0.243892 mm；端点和过渡连接均检查。', '结果')
para('单路线抬层具有局部可行性；它没有证明多路线同时抬层后全板无碰撞，因此继续执行顺序分配和1024条压力实例。', '结论')
source('3D/outputs/step_9_d_probe_summary.json、step_9_d_local_validation.json、step_9_d_synthetic.json。')

heading('21 两层与三层顺序分配',('G3D-05','G3D-06','G3D-07'))
para('在同一512条二维布局上执行两层30次探针，再以各50次尝试比较两层和三层。每次抬层后检查整个受影响集合，记录移除、新增与净变化，以总体近距对减少为收益。', '做了什么')
layers=tabdata('t81_3d_experiments_data.csv')
rows=[]
for r in layers[:3]:
    reduce=(int(r['initial_pairs'])-int(r['final_pairs']))/int(r['initial_pairs'])*100
    rows.append([f"{r['layers']}层 / {r['attempts']}次",integer(r['successful_elevations']),integer(r['initial_pairs']),integer(r['final_pairs']),fmt(reduce,2),fmt(r['extra_length_mm'],2),fmt(r['runtime_s'],2)])
table('512条连接的三组抬层结果',['配置','成功抬层','初始近距对','最终近距对','减少（%）','增长（mm）','耗时（s）'],rows,[3.0,1.9,2.6,2.6,2.1,2.4,2.4],text_cols=(0,),best=[(2,1),(2,3),(2,4),(1,5),(1,6)],notes='粗体只比较同为50次尝试的两层与三层：三层几何收益更好，两层成本更低。30次是先期探针。近距对用0.1 mm中心线阈值，与二维cross事件不同。')
para('两层30次成功12次，近距对减少4.06%；50次两层成功16次、减少5.84%，三层成功22次、减少8.32%。三层比两层多减少1224对，代价是额外长度与运行时间增加。', '结果')
para('新增近距对累计分别19、35、38；50次两层移除2929对、三层移除4156对，因此有新增但净数下降。三组均保留3对未决。30次的18次失败因双方已经抬过，本轮不重复抬层。', '碰撞账目与失败')
para('第三层没有“只有它才可行”的独有成功路线，额外收益来自更好的替代选择。层数增加在本次条件下有几何收益，但50次后仍有45400对近距路线，没有得到无碰撞版图。', '结论')
source('t81；3D/outputs/step_9_e_sequential_elevation_summary.json、step_9_f_{two_layer_control,three_layer}_summary.json、step_9_f_comparison.json。')

heading('22 固定1024条三层实验',('G3D-08',))
para('1024条输入在150 mm板高时，800条普通分配成功、160条轨道耗尽、64条特殊待解析；200 mm高提供1086根轨道，960条普通和64条特殊全部生成、耗尽0。最终使用300 × 200 mm区域。', '板容量负对照')
para('将固定输入确定性扩展成1024条合成连接，使用300 × 200 mm区域，层高0、1、2 mm，最多1024次尝试；局部抬高后回到底层，端点不变。中心线近距阈值0.1 mm。', '做了什么')
picture(P3/'outputs/figures/step_11_1024_3d_overview_z20.png',11.0,'自己的1024条三维总览 真实高度0 1 2 mm 显示高度放大20倍')
r=layers[3]; reduce=(int(r['initial_pairs'])-int(r['final_pairs']))/int(r['initial_pairs'])*100
continuation('终态结果与核对')
table('1024条合成输入的终态指标',['尝试','成功抬层','初始近距对','最终近距对','减少（%）','增长（mm）','耗时（s）'],[[r['attempts'],r['successful_elevations'],r['initial_pairs'],r['final_pairs'],fmt(reduce,2),fmt(r['extra_length_mm'],2),fmt(r['runtime_s'],2)]],[1.5,2.0,2.9,2.9,2.1,2.4,3.2],text_cols=(),notes='不同于512条的规模、板尺寸和尝试预算，独立报告，不跨规模评最优。所有指标是中心线几何统计。')
para('成功抬层175次，产生350个过渡；近距对204291→138113，减少32.39%。总长度增加87.41 mm，耗时1172.66 s，约19.5分钟；另有4对未决。终态849条在底层、72条主体在1 mm层、103条在2 mm层。', '结果')
para('累计移除68782对、新增2604对，净减少66178对；845次因双方已抬、4次无改善。独立重载1024条终态后全查523776路线对，与局部更新账目一致，696项测试通过。', '失败与终态独立复核')
para('局部抬层能减少约三分之一的近距对，但仍有大量剩余和未决对。这是合成压力实例的几何改善，没有计算三维光学损耗，也没有制作器件验证。', '结论')
source('3D/outputs/step_10_fixed_1024_{summary,initial_stats,validation,final_route_state,board_capacity_check}.json。')

heading('23 自己的三维结果图怎么读',('G3D-09',))
picture(P3/'outputs/figures/step_11_1024_xz_side.png',15.8,'自己的XZ侧视图 紫色是升层与降层过渡')
para('灰、蓝、橙分别表示0、1、2 mm层的主体段；紫色表示过渡。端点留在底层，被选中的路线在中段升高，末段降回底层。高度信息要结合近距检测判断。')
picture([P3/'outputs/figures/step_11_layer1_elevation_side.png',P3/'outputs/figures/step_11_layer2_elevation_side.png'],[8.15,8.15],'自己的局部抬层示例 左为1 mm层 右为2 mm层')
para('可视化阶段保存了13张原图，包括俯视、侧视、总览、分层和局部示例。图来自已保存的终态几何，不增加一次路由实验。俯视相交不能单独代表三维碰撞，也不能仅凭总览图声称碰撞已全部消除。', '结果和结论')
source('3D/outputs/step_11_visualization_summary.json；原图均位于3D/outputs/figures/。')

page(); heading('24 全部实验可以得出的结论')
para('二维复现方面，原版几何和直线加弯曲损耗有逐路一致性证据，主要损耗数值接近原报告。交叉表含数字化近似，精确求交与整数角统计必须分开使用。')
para('优化方面，扩大半径和减少总转弯角度带来的收益明显大于当前轨道候选搜索。固定端点、逐路半径一致的对照进一步支持自由弯角的作用，同时它会增加小角交叉和间距不足。')
para('约束方面，降低间距违规或压力情景最大值并不等于全部保护路线都改善。严格保护下256条没有找到可接受改进；512条找到6个接受候选，带来小幅损耗改善。倍率10的结果仍略恶化。')
para('早期交叉诊断方面，反转排序能改善部分双交叉，但会新增多交叉；硬约束会耗尽轨道，未完成全板。局部单线恢复沙箱有1到2个多交叉改善，尚未合并为完整全板结果。')
para('三维方面，增加高度层和局部抬升可以减少中心线近距对。1024条合成实例减少32.39%，仍剩138113对和4对未决；当前证据支持几何改善，不能据此声称三维光学损耗或制造问题已经解决。')
heading('图表和数据位置',level=2)
para('根目录为 C:\\Users\\lihao\\Desktop\\Graduation Project。2D表示OpticalWaveguideRouter2D，3D表示OpticalWaveguideRouter3D。',small=True)
para('二维复现图：2D/results/ 的fiberBoard256bend.png、fiberBoard512bend.png、loss_vs_radius.png、loss_distribution_R5.png。',small=True)
para('最差路线原图：3D/outputs/opt2d_step13/512/figures/worst_route_A.png与worst_route_F56.png。三维原图：3D/outputs/figures/。',small=True)
para('t31至t81为3D/publication/tables/下对应*_data.csv的短名；优化主对照直接读取outputs下保存的CSV，纠正部分旧摘要描述和字段陷阱。每个实验节末列出数据记录。',small=True)
para('表中“缺失”表示没有保存的数据，“不适用”表示指标在该条件下没有意义，0表示明确统计为零；不以0替代缺失。未完成的全板不生成损耗排名。损耗通常保留4位、时间与长度汇总2位小数；微小耗时和单线长度探针保留6位，印刷精度单独说明。',small=True)
para('平均和P95描述一次布局内各条路线的统计，不是多次随机运行的置信区间。报告包含41项原清单记录以及清单之外已经执行的早期基线和交叉诊断；纯计划、未执行策略不计作实验。本次整理没有重新运行路由。',small=True)

inventory=csvread(P3/'publication/manifest/experiment_inventory.csv')
missing={r['experiment_id'] for r in inventory}-set(COVERAGE)
if missing: raise ValueError('Uncovered inventory experiments '+str(sorted(missing)))
doc.core_properties.title=NAME; doc.core_properties.author='李昊伦'
doc.core_properties.subject='已完成的二维复现 轨道和路径优化 交叉诊断 三维抬层实验'
# Set the East Asian face at every inheritance level, including table text and fields.
for style in doc.styles:
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'),'宋体')
    fonts.attrib.pop(qn('w:eastAsiaTheme'),None)
defaults=doc.styles.element.find(qn('w:docDefaults'))
if defaults is not None:
    for fonts in defaults.iter(qn('w:rFonts')):
        fonts.set(qn('w:eastAsia'),'宋体')
        fonts.attrib.pop(qn('w:eastAsiaTheme'),None)
parts=[doc.part]
for section in doc.sections:
    parts.extend([section.header.part,section.footer.part])
for part in parts:
    for run in part.element.iter(qn('w:r')):
        fonts=run.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:eastAsia'),'宋体')
        fonts.attrib.pop(qn('w:eastAsiaTheme'),None)
doc.save(OUT/(NAME+'.docx'))
(OUT/'all_experiments_figure_provenance.json').write_text(json.dumps(IMAGES,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'all_experiments_coverage.json').write_text(json.dumps({'inventory_coverage':COVERAGE,'additional_sections':['早期输入与二维骨架验证','排序反转与硬约束负结果','局部单线恢复沙箱'],'tables':TABLES,'images':IMAGES},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(OUT/(NAME+'.docx')),'images':len(IMAGES),'tables':len(TABLES),'inventory_covered':len(COVERAGE)},ensure_ascii=False))
