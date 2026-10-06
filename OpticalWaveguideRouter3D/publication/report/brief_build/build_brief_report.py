from pathlib import Path
import csv
import json
import hashlib
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[4]
P3 = ROOT / 'OpticalWaveguideRouter3D'
P2 = ROOT / 'OpticalWaveguideRouter2D'
OUT = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)

def csv_rows(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

repro = csv_rows(P3 / 'publication/tables/t32_reproduction_compare_data.csv')
radius = csv_rows(P3 / 'publication/tables/t33_radius_sweep_512_data.csv')
step12 = csv_rows(P3 / 'publication/tables/t51_step12_main_data.csv')
step13 = csv_rows(P3 / 'publication/tables/t61_step13_main_data.csv')
step14 = {str(n): csv_rows(P3 / f'outputs/opt2d_step14/{n}/comparison.csv') for n in (256, 512)}
layers = csv_rows(P3 / 'publication/tables/t81_3d_experiments_data.csv')
USED_IMAGES = []

def scheme(rows, name, n=None):
    return next(r for r in rows if r['scheme'] == name and (n is None or r.get('channels') == str(n)))

def num(value, places=4):
    return f'{float(value):.{places}f}'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Cm(1.6)
sec.left_margin = sec.right_margin = Cm(2)
sec.header_distance = sec.footer_distance = Cm(0.7)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption', 'Header', 'Footer']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    st.paragraph_format.widow_control = True
    for border in st.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Normal'].paragraph_format.space_after = Pt(4)
doc.styles['Title'].font.size = Pt(20)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(9)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(7)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)
doc.styles['Heading 2'].font.size = Pt(11.5)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].font.bold = False
doc.styles['Caption'].paragraph_format.line_spacing = 1.05
doc.styles['Caption'].paragraph_format.space_after = Pt(5)

header = sec.header.paragraphs[0]
header.text = '光波导自动布线实验报告'
header.style = 'Header'
header.runs[0].font.size = Pt(9)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.style = 'Footer'
footer.add_run('第 ')
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)
footer.add_run(' 页')

def para(text, label=None, small=False):
    p = doc.add_paragraph()
    if label:
        p.add_run(label + '：').bold = True
    p.add_run(text)
    if small:
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            r.font.size = Pt(9)
    return p

def heading(text, level=1):
    return doc.add_heading(text, level)

def table(headers, rows, widths, left_cols=(0,)):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Cm(width)
    for i, (cell, text) in enumerate(zip(t.rows[0].cells, headers)):
        cell.text = str(text)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    for ri, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        no_break = OxmlElement('w:cantSplit')
        pr.append(no_break)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader')
            pr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, val in [('top', 75), ('bottom', 75), ('left', 95), ('right', 95)]:
                el = OxmlElement('w:' + side)
                el.set(qn('w:w'), str(val))
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + side)
                el.set(qn('w:val'), 'single')
                el.set(qn('w:sz'), '4')
                el.set(qn('w:color'), 'D9D9D9')
                borders.append(el)
            tcpr.append(borders)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E9EEF5' if ri == 0 else ('F7F8FA' if ri % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci in left_cols else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.bold = ri == 0
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = Pt(2)
    after.add_run().font.size = Pt(2)
    return t

def picture(path, width, caption, alt):
    USED_IMAGES.append(Path(path))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Cm(width))
    shape._inline.docPr.set('descr', alt)
    cp = doc.add_paragraph(caption, style='Caption')
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

SOURCE_NOTES = []
def source(text):
    SOURCE_NOTES.append(text)

def newpage():
    doc.add_page_break()

# Page 1: definition and actual reproduction layout.
doc.add_paragraph('光波导自动布线实验报告', style='Title')
para('我们复现了二维波导排布，比较了半径、轨道和路径形状，再进行了三维抬层实验。主要结果是：扩大半径和减少转弯角度能降低模型损耗；三维抬层能减少近距路线对，但仍存在交叉和间距问题。')
heading('R5 和 R6 的含义', 2)
para('R 表示圆弧转弯半径。R5 是 5 mm，R6 是 6 mm。R6 转弯更缓，但需要更大空间。模型中一次 90° 转弯的损耗约为 2.38 dB 和 1.90 dB。F56 表示自由弯角路径与 5/6 mm 半径选择；清单编号 R05、R06 则分别对应 2 mm、3 mm 半径实验。')
heading('实验1 原二维布线复现')
para('恢复原程序的轨道分配与圆弧构造，在 150 × 150 mm 区域分别布置 256、512 条连接，并计算直线、弯曲和交叉损耗。', '做了什么')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
p.paragraph_format.keep_with_next = True
for i, n in enumerate([256, 512]):
    if i:
        p.add_run('  ')
    shape = p.add_run().add_picture(str(P2 / f'results/fiberBoard{n}bend.png'), width=Cm(8.25))
    USED_IMAGES.append(P2 / f'results/fiberBoard{n}bend.png')
    shape._inline.docPr.set('descr', f'{n} 条波导半径 5 mm 的二维复现布局')
cp = doc.add_paragraph('图1 二维复现布局  左为256条  右为512条  红线表示波导', style='Caption')
cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = []
for case, label in [('256', '256条 R5'), ('512', '512条 R5'), ('512_R4', '512条 R4')]:
    a = next(r for r in repro if r['case'] == case and r['metric'] == '平均损耗')
    b = next(r for r in repro if r['case'] == case and r['metric'] == '最大损耗')
    rr.append([label, num(a['repro_db']), num(b['repro_db']), f"{float(a['thesis_db']):.1f} / {float(b['thesis_db']):.1f}"])
table(['实验', '平均损耗 dB', '最大损耗 dB', '原报告平均 / 最大'], rr, [4, 3.7, 3.7, 5.3])
para('全部连接完成。512 条路线的轨道与弯曲几何和恢复的原程序逐条一致；主要损耗数值接近原报告。交叉损耗使用原报告图线提取的近似表。所有损耗均为模型计算值，没有制作器件实测。', '结论')
source('2D/results/fiberBoard256_loss_summary.json、fiberBoard512_loss_summary.json；t32_reproduction_compare_data.csv。2D和3D的完整目录见末页。')

# Page 2: radius and crossing-table sensitivity.
newpage()
heading('实验2 转弯半径对照')
para('512 条连接分别采用 2、3、4、5 mm 半径重新布线，再计算每条路线的损耗。每个半径独立构造几何。', '做了什么')
table(['转弯半径 mm', '平均损耗 dB', '最大损耗 dB'], [[r['radius_mm'], num(r['mean_loss_db']), num(r['max_loss_db'])] for r in radius], [5.7, 5.5, 5.5])
picture(P2 / 'results/loss_vs_radius.png', 16.7, '图2 半径与损耗  左图为平均及最大损耗  右图为逐条波导分布', '512 条波导半径从 2 mm 增到 5 mm 时平均和最大损耗下降')
para('平均损耗由 16.5873 降至 5.5147 dB。在 R5 布局中，弯曲损耗占约 82.7%，直线占 11.5%，交叉占 5.8%。在这批数据和当前模型下，降低弯曲损耗是主要改进方向。', '结论')
heading('交叉损耗表核对', 2)
para('原始交叉损耗表缺失，因此我们在同一 512 条 R5 几何上比较三种取法，检查数值对近似表的依赖。')
table(['交叉损耗表取法', '平均损耗 dB', '最大损耗 dB'], [
    ['全部按90°交叉估算', '5.4758', '6.5224'],
    ['图线提取并用90°文字值校准', '5.5147', '6.5761'],
    ['图线提取原值', '5.5745', '6.6659'],
], [8.2, 4.25, 4.25])
para('三种取法的平均损耗跨度约 0.099 dB，远小于半径实验的变化。主要趋势一致，但交叉损耗仍含近似，不能宣称恢复了全部原始损耗数据。', '结论')
source('t33_radius_sweep_512_data.csv、t34_crossing_model_sensitivity_data.csv；2D/tools/radius_sweep.py。')

# Page 3: early optimization and controlled shape comparison.
newpage()
heading('实验3 候选轨道与半径选择')
para('比较原版 R5、统一 R6、固定 R5 的候选轨道搜索，以及允许选择 R5/R6 的方案。这是早期探索，允许同端口槽位调整；后续路径结构与保护实验才逐路冻结端点。这里采用解析求交统计，与实验1原版统计的小数略有差别。', '做了什么')
rows = []
for name, label in [('A', '原版R5'), ('B', '统一R6'), ('C', 'R5候选轨道优化'), ('D', 'R5/R6自适应')]:
    r = scheme(step12, name, 512)
    placed = 512 - int(r['unplaced_count'])
    rows.append([label, f'{placed}/512', num(r['mean_loss_db']), num(r['max_loss_db'])])
table(['512条连接的方案', '布通条数', '平均损耗 dB', '最大损耗 dB'], rows, [6.1, 3.0, 3.8, 3.8])
para('固定 R5 搜索轨道，平均只改善约 0.0015 dB；允许选择 R5/R6 后全部布通，平均降至 4.5593 dB。统一 R6 有 2 条未布通，其损耗只统计已布通路线，不能与完整方案直接排名。', '结果和结论')
para('还测试了候选数、布线顺序、拆线重布和位置惩罚。候选数从 1 增至 16，平均改善仍不超过约 0.002 dB；本次拆线重布没有进一步改善。按连接跨度排序时，256、512 规模分别有 76、144 条未布通。增加搜索范围或改变顺序并不保证更好的完整结果。')
heading('实验4 自由弯角路径')
para('保持每条连接的起终点不变，将跨上下两侧的连接改成圆弧与斜线组成的 S 形，允许较小转弯角度，并选择 5/6 mm 半径。同侧连接保留 U 形。', '做了什么')
sr = []
for n in (256, 512):
    for name, label in [('A', '原版R5'), ('D56', '自适应半径 U形'), ('F56', '自适应半径 自由弯角')]:
        r = scheme(step13, name, n)
        sr.append([str(n), label, num(r['mean_loss_db']), num(r['max_loss_db'])])
table(['条数', '路径方案', '平均损耗 dB', '最大损耗 dB'], sr, [1.7, 7.2, 3.9, 3.9])
para('相对原版的总改善同时包含半径与路径结构两部分。在逐路半径相同的条件下，改变路径结构使 256、512 平均损耗再降低约 1.3690、1.2661 dB。这支持减少总转弯角度的作用。', '结论')
source('t51_step12_main_data.csv、t52_step12_ablation_data.csv、t61_step13_main_data.csv。优化基准使用解析求交统计，与实验1原版统计的小数略有差别。')

# Page 4: geometry picture, validation and strict acceptance.
newpage()
heading('最差路线与几何复核')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
p.paragraph_format.keep_with_next = True
for i, name in enumerate(['A', 'F56']):
    if i:
        p.add_run('  ')
    path = P3 / f'outputs/opt2d_step13/512/figures/worst_route_{name}.png'
    USED_IMAGES.append(path)
    shape = p.add_run().add_picture(str(path), width=Cm(8.25))
    shape._inline.docPr.set('descr', f'项目已有的512条{name}方案最差路线局部图')
cp = doc.add_paragraph('图3 项目已有的最差路线局部图  左为原版路线287  右为自由弯角方案路线34', style='Caption')
cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
para('灰线表示周围波导，红线为各方案最差路线的局部关注段。两图分别选择各方案的最差路线，路线编号不同。512条实验补修前后数据一致，因此可以使用项目中已保存的这两张图。', small=True)
wa = next(r for r in csv_rows(P3 / 'outputs/opt2d_step13_fix/512/A/per_route.csv') if r['route_id'] == '287')
wf = next(r for r in csv_rows(P3 / 'outputs/opt2d_step13_fix/512/F56/per_route.csv') if r['route_id'] == '34')
table(['各方案最差路线的指标', '原版路线287', '自由弯角方案路线34'], [
    [label, num(wa[key]), num(wf[key])] for key, label in [
        ('total_loss_db', '总损耗 dB'),
        ('bend_loss_db', '弯曲损耗 dB'),
        ('crossing_loss_db', '交叉损耗 dB'),
    ]
], [6.5, 5.1, 5.1])
para('自由弯角降低了弯曲损耗，但 512 条布局的平均交叉损耗从 0.3219 增至 0.5650 dB，间距违规路线对从 692 增至 1231。全部布通和无重合，并不表示满足全部间距要求。')
para('早期 256 条自由弯角布局发现路线7与31重合约 152.007 mm。修复求交判断并重新布线后，重合数为0，平均损耗由 2.9689 变为2.9703 dB。本报告采用补修后的结果。')
newpage()
heading('实验5 小角交叉与关键路线保护')
para('冻结端点和每条路线的半径，惩罚间距不足、小于20°的交叉，并局部重布。接受改动时要求全局最大损耗和固定11条保护路线不恶化。', '做了什么')
table(['规模', '尝试 / 接受候选', '平均损耗 dB', '最大损耗 dB'], [
    ['256', '68 / 0', '2.9703 → 2.9703', '5.0783 → 5.0783'],
    ['512', '74 / 6', '3.2947 → 3.2936', '5.4995 → 5.4654'],
], [2.2, 4.0, 5.25, 5.25])
para('保护性优化在 512 条布局找到小幅改善，256 条布局没有候选满足全部条件。单独加入间距、小角度惩罚或两者同时加入，虽然改善了部分指标，但均未通过全部保护条件。', '结论')
heading('小角交叉压力测试', 2)
para('固定几何，仅把小于20°的交叉损耗乘以5或10，比较最差路线的结果。这是人为压力情景，用来检查方案对交叉项不确定性的敏感程度。')
table(['512条布局的情景', '优化前最大 dB', '优化后最大 dB'], [
    ['小角交叉损耗乘以5', '5.7264', '5.6340'],
    ['小角交叉损耗乘以10', '8.4333', '8.4996'],
], [7.0, 4.85, 4.85])
para('乘以5时改善，乘以10时略恶化，因此不能说方案对所有风险情景都更好。最终 256 自由弯角方案仍有415对间距违规；512保护优化后仍有1228对。')
source('outputs/opt2d_step13_fix、outputs/opt2d_step14；step_13_opt2d_fix_and_reverification.md。')

# Page 5: actual 3D layout and comparable geometry metrics.
newpage()
heading('实验6 三维抬层布线')
para('先在同一 512 条布局上比较两层与三层，各尝试50次；再把固定输入确定性扩展为1024条连接，在300 × 200 mm区域使用0、1、2 mm三个高度层，局部抬高路线。近距判断采用最小中心线间距0.1 mm的实验阈值。', '做了什么')
picture(P3 / 'outputs/figures/step_11_1024_3d_overview_z20.png', 11.4,
        '图4 1024条连接的三维布线  真实层高0 1 2 mm  图中高度放大20倍',
        '灰色为层0蓝色为层1橙色为层2紫色为层间过渡的1024条三维波导')
table(['实验', '尝试 / 成功抬升', '近距路线对变化', '减少比例'], [
    ['512条 两层', '50 / 16', '49518 → 46624', '5.84%'],
    ['512条 三层', '50 / 22', '49518 → 45400', '8.32%'],
    ['1024条 三层', '1024 / 175', '204291 → 138113', '32.39%'],
], [4.0, 4.0, 5.3, 3.4])
para('同样尝试次数下，三层比两层减少更多近距路线对，但额外长度和时间更多。1024条实验运行约19.5分钟，总长增加约87.41 mm；849条未抬升，72条抬到1 mm层，103条抬到2 mm层。', '结果')
para('局部抬层能减少几何碰撞，但仍剩138113对近距路线，另有4对未决，尚未获得无碰撞版图。本实验没有计算三维光学损耗，也没有验证器件制造。', '结论')
source('t81_3d_experiments_data.csv；outputs/step_10_fixed_1024_summary.json。1024输入为合成扩展实例。')

# Page 6: read the shape, practical conclusions and compact source index.
newpage()
heading('三维结果如何看')
picture(P3 / 'outputs/figures/step_11_1024_xz_side.png', 14.5,
        '图5 三维布线的XZ侧视  紫色为抬层与下降过渡',
        '所有1024条路线侧视投影显示0 1 2 mm高度与层间过渡')
para('灰色、蓝色、橙色分别表示三个层的主体段，紫色表示层间过渡。端点保持在底层，选中的路线在中间段抬高后再回到底层。俯视投影仍可能相交，因此要结合高度与近距检测判断。')
heading('这些实验得出的结论')
para('原二维布线的几何与主要损耗数值基本复现。损耗中交叉项使用近似数据，当前证据支持主要数值和趋势，不能把交叉表当作原始测量数据。')
para('在当前模型中，弯曲是损耗的主要来源。扩大半径和减少总转弯角度带来的改善，明显大于单纯增加轨道候选；统一大半径又会遇到空间限制。')
para('自由弯角能降低平均和最大损耗，但增加小角交叉与间距不足。严格保护现有路线后，进一步改进较小；未布通或未通过验收的配置不能当作最终成功方案。')
para('三维抬层已经显示出减少近距路线对的作用。1024条实验减少约三分之一，但仍有大量剩余近距对。目前能得出的结论是几何改善，不能据此声称光学损耗或制造问题已经解决。')
heading('结果图与数据位置', 2)
para('文件均位于 C:\\Users\\lihao\\Desktop\\Graduation Project。以下用2D表示 OpticalWaveguideRouter2D，用3D表示 OpticalWaveguideRouter3D。', small=True)
para('二维原图：2D/results/fiberBoard256bend.png 与 fiberBoard512bend.png。半径对照图：2D/results/loss_vs_radius.png。', small=True)
para('三维原图：3D/outputs/figures/ 中的 step_11_1024_3d_overview_z20.png 和 step_11_1024_xz_side.png。', small=True)
para('最差路线原图：3D/outputs/opt2d_step13/512/figures/ 中的 worst_route_A.png 和 worst_route_F56.png。', small=True)
para('实验数值：2D/results/*_loss_summary.json、3D/publication/tables/*_data.csv。补修与保护明细：3D/outputs/opt2d_step13_fix/ 和 opt2d_step14/。', small=True)
para('平均、最大和分位数描述一次布局内各条路线的结果。本文整理已有实验及验证记录，本次重写没有重新运行布线实验。', small=True)

doc.core_properties.title = '光波导自动布线实验报告'
doc.core_properties.subject = '二维复现 半径和路径优化 三维抬层实验结果'
doc.core_properties.author = '李昊伦'
doc.core_properties.keywords = '光波导 实验结果 R5 R6 三维布线'
doc.save(OUT / '光波导布线实验简明报告.docx')
(OUT / 'figure_provenance.json').write_text(json.dumps([
    {'path': str(p), 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in USED_IMAGES
], ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'docx': str(OUT / '光波导布线实验简明报告.docx'), 'images': len(doc.inline_shapes), 'tables': len(doc.tables)}, ensure_ascii=False))
