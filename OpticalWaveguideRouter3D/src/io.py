"""连接关系读取、配置读取与布线结果保存接口。"""

from pathlib import Path
from typing import Any

from .models import Route, Waveguide


def load_connections(path: str | Path) -> list[Waveguide]:
    """预留 CSV 连接读取接口；字段规范待定义。"""
    raise NotImplementedError


def load_config(path: str | Path) -> dict[str, Any]:
    """预留 YAML 配置读取接口；当前不解析文件。"""
    raise NotImplementedError


def save_routes(routes: list[Route], path: str | Path) -> None:
    """预留布线结果保存接口；存储格式待定义。"""
    raise NotImplementedError


def load_legacy_fiberboard(path: str | Path) -> list[Waveguide]:
    """读取单工作表、首行 Port1/Port2 表头的前代连接表。

    完全空白行跳过；部分缺值、字符串 ID 或非整数值拒绝。公式仅使用文件保存的缓存值，不重算。
    ID 在单次导入的数据集中唯一，按有效行从零编号；不同文件须独立使用。
    不分配槽位或位置，不改变连接方向，不去重，不保存原文件。
    """
    from math import isfinite
    from openpyxl import load_workbook
    from .models import Port

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Expected exactly one worksheet.")
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(next(rows, ()))
        if headers.count("Port1") != 1 or headers.count("Port2") != 1:
            raise ValueError("Row 1 must contain unique Port1 and Port2 columns.")
        columns = (headers.index("Port1"), headers.index("Port2"))
        waveguides: list[Waveguide] = []
        raw_workbook = load_workbook(Path(path), read_only=True, data_only=False)
        try:
            raw_rows = list(raw_workbook.worksheets[0].iter_rows(min_row=2, values_only=True))
        finally:
            raw_workbook.close()
        for row_number, row in enumerate(rows, start=2):
            if all(value is None for value in raw_rows[row_number - 2]):
                continue
            pmt_ids: list[int] = []
            for column, label in zip(columns, ("Port1", "Port2")):
                value = row[column]
                if (
                    isinstance(value, bool)
                    or not isinstance(value, (int, float))
                    or (isinstance(value, float) and (not isfinite(value) or not value.is_integer()))
                ):
                    raise ValueError(
                        f"{sheet.title} row {row_number} {label}: expected integer PMT ID, got {value!r}."
                    )
                pmt_ids.append(int(value))
            identifier = len(waveguides)
            waveguides.append(Waveguide(
                id=identifier,
                start_port=Port(2 * identifier, pmt_ids[0], None, None),
                end_port=Port(2 * identifier + 1, pmt_ids[1], None, None),
            ))
        return waveguides
    finally:
        workbook.close()


def load_legacy_512_snapshot(
    fiberboard_path: str | Path, snapshot_path: str | Path
) -> list[Waveguide]:
    """Restore this legacy endpoint snapshot in original connection order.

    Supports small fixtures of the same schema, not arbitrary layout generation.
    The unnamed first column is the zero-based input connection index.
    Coordinates retain snapshot values (legacy evidence supports mm); local_id
    remains None. Formulas and ambiguous self-connections are rejected.
    """
    from math import isfinite
    from openpyxl import load_workbook
    from .models import Point2D

    waveguides = load_legacy_fiberboard(fiberboard_path)
    book = load_workbook(snapshot_path, read_only=True, data_only=False)
    try:
        if len(book.worksheets) != 1:
            raise ValueError("Snapshot must have exactly one worksheet.")
        rows = book.worksheets[0].iter_rows(values_only=True)
        headers = tuple(next(rows, ()))
        required = ("Port1", "Port2", "index1", "index2", "sy", "ly", "dz", "sx", "lx", "dx")
        if not headers or headers[0] is not None:
            raise ValueError("Snapshot first column must be the unnamed connection index.")
        if any(headers.count(key) != 1 for key in required):
            raise ValueError("Snapshot requires unique columns: " + ", ".join(required))
        columns = {key: headers.index(key) for key in required}
        seen: set[int] = set()
        for row_number, row in enumerate(rows, 2):
            if all(value is None for value in row):
                continue
            values = {"connection_index": row[0], **{key: row[col] for key, col in columns.items()}}
            for key, value in values.items():
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not isfinite(value):
                    raise ValueError(f"Snapshot row {row_number} {key}: expected finite number.")
            for key in ("connection_index", "Port1", "Port2", "index1", "index2"):
                if values[key] != int(values[key]):
                    raise ValueError(f"Snapshot row {row_number} {key}: expected integer.")
            index = int(values["connection_index"])
            if index < 0 or index >= len(waveguides):
                raise ValueError(f"Snapshot row {row_number}: index {index} out of range.")
            if index in seen:
                raise ValueError(f"Snapshot row {row_number}: duplicate index {index}.")
            seen.add(index)
            if values["dz"] != 0 or abs(values["dx"] - abs(values["sx"] - values["lx"])) > 1e-8:
                raise ValueError(f"Snapshot row {row_number}: inconsistent dx or nonzero dz.")
            waveguide = waveguides[index]
            start_id, end_id = waveguide.start_port.pmt_id, waveguide.end_port.pmt_id
            pair = (int(values["Port1"]), int(values["Port2"]))
            if start_id == end_id:
                raise ValueError(f"Snapshot row {row_number}: ambiguous self-connection.")
            first = Point2D(float(values["sx"]), float(values["sy"]))
            second = Point2D(float(values["lx"]), float(values["ly"]))
            if pair == (start_id, end_id):
                waveguide.start_port.position, waveguide.end_port.position = first, second
            elif pair == (end_id, start_id):
                waveguide.start_port.position, waveguide.end_port.position = second, first
            else:
                raise ValueError(f"Snapshot row {row_number}: PMT pair does not match input index {index}.")
        missing = set(range(len(waveguides))) - seen
        if missing:
            raise ValueError(f"Snapshot missing connection indices: {sorted(missing)}.")
        return waveguides
    finally:
        book.close()
