# -*- coding: utf-8 -*-
"""交付前程序化校验：页数、越界、文字溢出、图形重叠、图片引用、字体、占位符残留。"""
import math
import os
import sys

from pptx import Presentation
from pptx.util import Emu, Inches

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deck_lib import units_of

ROOT = r'C:\Users\lihao\Desktop\Graduation Project'
PATH = os.path.join(ROOT, 'TOA-cos论文精读_讲解型汇报.pptx')
W, H = 13.3333, 7.5
prs = Presentation(PATH)

problems = []
print(f'幻灯片尺寸: {prs.slide_width/914400:.3f} x {prs.slide_height/914400:.3f} in')
print(f'页数: {len(prs.slides)}')


def walk(shapes, depth=0):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes, depth + 1)
        else:
            yield sh


for si, slide in enumerate(prs.slides, 1):
    boxes = []
    n_pic = 0
    for sh in walk(slide.shapes):
        l, t = sh.left / 914400, sh.top / 914400
        w, h = sh.width / 914400, sh.height / 914400
        if sh.shape_type == 13:
            n_pic += 1
        # 越界
        if l < -0.02 or t < -0.02 or l + w > W + 0.02 or t + h > H + 0.02:
            problems.append(f'P{si} 越界: <{sh.name}> '
                            f'({l:.2f},{t:.2f},{w:.2f},{h:.2f})')
        if sh.has_text_frame and sh.text_frame.text.strip():
            txt = sh.text_frame.text
            # 估算所需高度
            size = 18.0
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size:
                        size = max(size, r.font.size.pt)
                        break
            need = 0.0
            for p in sh.text_frame.paragraphs:
                sz = size
                for r in p.runs:
                    if r.font.size:
                        sz = r.font.size.pt
                        break
                per = max(1.0, (w * 72.0) / sz)
                lines = max(1, math.ceil(units_of(p.text) / per))
                need += lines * sz * 1.36 / 72.0
            if need > h * 1.10 + 0.06:
                problems.append(f'P{si} 可能溢出: <{sh.name}> 需 {need:.2f}in > 框 {h:.2f}in  '
                                f'| {txt[:34]!r}')
            if t < 0.78 and sh.is_placeholder is False and l > 1.0:
                pass
            boxes.append((l, t, w, h, sh.name, txt[:24]))
    # 文本框互压检查
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.12 and oy > 0.12:
                problems.append(f'P{si} 文本框重叠 {ox:.2f}x{oy:.2f}in: '
                                f'<{a[4]}> {a[5]!r} × <{b[4]}> {b[5]!r}')
    print(f'  P{si:2d}  形状 {len(list(walk(slide.shapes))):3d}  图片 {n_pic}  '
          f'备注 {"有" if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip() else "无"}')

# 图片与关系完整性
bad = 0
for si, slide in enumerate(prs.slides, 1):
    for rid, rel in slide.part.rels.items():
        if rel.reltype.endswith('/image'):
            tp = rel.target_part
            if tp is None or not getattr(tp, 'blob', b''):
                problems.append(f'P{si} 图片关系损坏 rId={rid}')
                bad += 1

# 占位符残留
for si, slide in enumerate(prs.slides, 1):
    for sh in walk(slide.shapes):
        if sh.has_text_frame:
            tx = sh.text_frame.text
            for token in ('Click to add', '单击此处', 'xxx', 'lorem', 'TODO', '[insert', '请在此'):
                if token in tx:
                    problems.append(f'P{si} 残留占位文本 {token!r}: <{sh.name}>')

# 字体规范抽查
fonts = {}
for slide in prs.slides:
    for sh in walk(slide.shapes):
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    fonts.setdefault(r.font.name, 0)
                    fonts[r.font.name] += 1
print('\n拉丁字体使用统计:', fonts)

print('\n' + '=' * 60)
if problems:
    print(f'发现 {len(problems)} 个问题:')
    for p in problems:
        print('  -', p)
else:
    print('校验通过：无越界 / 无溢出 / 无重叠 / 图片引用完好 / 无占位残留')
print(f'图片关系损坏数: {bad}')
