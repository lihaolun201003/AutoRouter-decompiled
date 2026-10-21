#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Independent acceptance audit for the V3-C4 four-fold symmetric connectivity deliverables.

This script deliberately does **not** import ``generate_v3_c4`` (or any other generation
module).  It re-reads the files that were actually written to disk, re-implements the
physical port model from the documented formulas, and re-checks every mandatory
acceptance item of the task:

* CSV exactly 1024 connections, project-compatible 12-column format, UTF-8-SIG, CRLF;
* all 2048 physical ports used exactly once; 64 connectors x 32 channels;
* NS/EW each 256, NE/ES/SW/WN each 128; row1=opposite 512, row2=adjacent 512;
* port coordinates consistent with four_side_1024/fslw/ports.py (independent formula
  + direct import of the repository port model as a second opinion);
* the whole connection set mapped by rot90/rot180/rot270 is still the same unordered
  connection set (port level, not connector level);
* every connection has the other three members of its 4-member rotation orbit;
* the 1024 connections in the Excel workbook match the CSV exactly;
* the V2 opposite 512 connections are preserved field-by-field;
* the recorded SHA256 values in validation_v3_c4.json match the actual files.

Run:
    python validate_v3_c4.py
Exit code 0 = all PASS, 1 = at least one FAIL.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
import sys
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
DESKTOP = Path(r'C:\Users\lihao\Desktop')

OUT_CSV = HERE / 'optical_1024_multitomulti_connectivity_v3_c4_symmetric.csv'
OUT_XLSX = HERE / 'fiberBoard1024_V3_C4_四边旋转对称连接表.xlsx'
OUT_JSON = HERE / 'validation_v3_c4.json'
OUT_MD = HERE / 'generation_report.md'
OUT_COMPARISON = HERE / 'pairing_comparison_v3_c4.json'
OUT_INDEPENDENT = HERE / 'independent_acceptance_v3_c4.json'
GENERATOR = HERE / 'generate_v3_c4.py'

V2_CSV = DESKTOP / 'optical_1024_multitomulti_connectivity_v2_symmetric.csv'
V1_CSV = DESKTOP / 'optical_1024_multitomulti_connectivity_v1.csv'

V1_SHA256 = '0ecce47109f9223e8407928c423ad5d9272186938cc267877a6158936e0e428b'
V2_SHA256 = '49dd9e89c4ea506d04638ae34eb06a0d8163d479de80869804c881ee1c43915e'

CSV_FIELDS = ['route_id', 'source_side', 'source_connector', 'source_row',
              'source_column', 'source_channel', 'target_side', 'target_connector',
              'target_row', 'target_column', 'target_channel', 'relation']

BOARD_MM = 150.0
CONNECTOR_PITCH_MM = 9.0
FIRST_CONNECTOR_CENTER_MM = 7.5
COLUMN_PITCH_MM = 0.25
ROWS = 2
COLUMNS = 16
CONNECTORS_PER_SIDE = 16
SIDES = ('N', 'E', 'S', 'W')
Z_MM = {1: 0.0, 2: 0.5}

CHECKS = []


def add(cid, category, name, expected, actual, passed=None, detail=''):
    if passed is None:
        passed = (expected == actual)
    CHECKS.append(dict(id=cid, category=category, name=name, expected=expected,
                       actual=actual, pass_=bool(passed), detail=detail))


def sha256_file(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def along_mm(idx, column):
    return (FIRST_CONNECTOR_CENTER_MM + (idx - 1) * CONNECTOR_PITCH_MM
            + (column - (COLUMNS + 1) / 2.0) * COLUMN_PITCH_MM)


def model_xyz(side, connector_index, row, column):
    a = along_mm(connector_index, column)
    x, y = {'N': (a, BOARD_MM), 'S': (a, 0.0),
            'E': (BOARD_MM, a), 'W': (0.0, a)}[side]
    return x, y, Z_MM[row]


PORT_COORDS = {}
for side in SIDES:
    for idx in range(1, CONNECTORS_PER_SIDE + 1):
        for row in (1, 2):
            for column in range(1, COLUMNS + 1):
                PORT_COORDS[(side, f'{side}{idx:02d}', row, column)] = \
                    model_xyz(side, idx, row, column)
COORD_TO_EID = {(round(x, 6), round(y, 6), z): eid
                for eid, (x, y, z) in PORT_COORDS.items()}
assert len(PORT_COORDS) == 2048 and len(COORD_TO_EID) == 2048


def rot90_eid(eid, times=1):
    x, y, z = PORT_COORDS[eid]
    for _ in range(times):
        x, y = round(BOARD_MM - y, 6), round(x, 6)
    return COORD_TO_EID[(round(x, 6), round(y, 6), z)]


def rot90_conn(conn):
    return frozenset(rot90_eid(e) for e in conn)


def read_csv_rows(path):
    raw = Path(path).read_bytes()
    rows = list(csv.DictReader(raw.decode('utf-8-sig').splitlines()))
    return rows, raw


def eid_of(row, role):
    return (row[f'{role}_side'], row[f'{role}_connector'],
            int(row[f'{role}_row']), int(row[f'{role}_column']))


def main():
    print('=' * 78)
    print('V3-C4 connectivity independent acceptance audit')
    print('=' * 78)
    print('output dir :', HERE)

    # ------------------------------------------------------------------ CSV
    raw = OUT_CSV.read_bytes()
    rows = list(csv.DictReader(raw.decode('utf-8-sig').splitlines()))
    header = list(rows[0].keys()) if rows else []
    add('V01', 'CSV', 'CSV 文件存在且使用 UTF-8-SIG（含 BOM）',
        True, raw.startswith(b'\xef\xbb\xbf'), passed=raw.startswith(b'\xef\xbb\xbf'))
    add('V02', 'CSV', 'CSV 表头为项目固定 12 列且顺序一致',
        CSV_FIELDS, header)
    add('V03', 'CSV', 'CSV 恰好 1024 条数据行', 1024, len(rows))
    add('V04', 'CSV', 'CSV 行尾为 CRLF（与现有 V1/V2 一致）',
        True, raw.count(b'\r\n') == len(rows) + 1)
    rids = [int(r['route_id']) for r in rows]
    add('V05', 'CSV', 'route_id 为 1..1024 连续且唯一', '1..1024',
        f'count={len(set(rids))}, min={min(rids)}, max={max(rids)}',
        passed=(sorted(rids) == list(range(1, 1025)) and len(set(rids)) == 1024))

    bad_fmt = []
    rel_of = {}
    pair_counts = Counter()
    directed = Counter()
    for r in rows:
        sides = {r['source_side'], r['target_side']}
        for role in ('source', 'target'):
            side = r[f'{role}_side']
            conn = r[f'{role}_connector']
            row = int(r[f'{role}_row'])
            col = int(r[f'{role}_column'])
            ch = int(r[f'{role}_channel'])
            idx = int(conn[1:]) if conn[1:].isdigit() else -1
            if not (side in SIDES and conn[0] == side and 1 <= idx <= 16
                    and row in (1, 2) and 1 <= col <= 16
                    and ch == (row - 1) * 16 + col):
                bad_fmt.append((r['route_id'], role))
        expected_rel = 'opposite' if sides in ({'N', 'S'}, {'E', 'W'}) else 'adjacent'
        rel_of[r['route_id']] = expected_rel
        pair_counts[''.join(sorted(sides))] += 1
        directed[f"{r['source_side']}->{r['target_side']}"] += 1
    add('V06', 'CSV', 'connector/row/column/channel 符合物理端口编号规则', [], bad_fmt[:5],
        passed=not bad_fmt, detail=f'violations={len(bad_fmt)}')
    bad_rel = [r['route_id'] for r in rows if r['relation'] != rel_of[r['route_id']]]
    add('V07', 'CSV', 'relation 与两端边类别一致（对边 opposite / 邻边 adjacent）',
        [], bad_rel[:5], passed=not bad_rel)
    add('V08', 'CSV', 'source 与 target 位于不同边', 1024,
        sum(1 for r in rows if r['source_side'] != r['target_side']))
    add('V09', 'CSV', '类别计数 NS 256, EW 256, EN 128, ES 128, SW 128, NW 128',
        {'NS': 256, 'EW': 256, 'EN': 128, 'ES': 128, 'SW': 128, 'NW': 128},
        dict(pair_counts))
    add('V10', 'CSV', '有向边对 N→S 256, E→W 256, N→E/E→S/S→W/W→N 各 128',
        {'N->S': 256, 'E->W': 256, 'N->E': 128, 'E->S': 128,
         'S->W': 128, 'W->N': 128}, dict(directed))
    bad_row = [r['route_id'] for r in rows
               if int(r['source_row']) != int(r['target_row'])
               or int(r['source_row']) != (1 if r['relation'] == 'opposite' else 2)]
    add('V11', 'CSV', 'row1 全为对边（512）、row2 全为邻边（512），同行连接', [], bad_row[:5],
        passed=not bad_row, detail=f'violations={len(bad_row)}')

    # --------------------------------------------------------- port usage
    usage = Counter()
    for r in rows:
        usage[eid_of(r, 'source')] += 1
        usage[eid_of(r, 'target')] += 1
    add('V12', '端口', '2048 个物理端口各被使用且仅使用一次',
        {'unique': 2048, 'used_once': True},
        {'unique': len(usage), 'used_once': set(usage.values()) == {1}},
        passed=(len(usage) == 2048 and set(usage.values()) == {1}))
    side_use = Counter(eid[0] for eid in usage)
    add('V13', '端口', '每边 512 个端口全部使用', {'N': 512, 'E': 512, 'S': 512, 'W': 512},
        {s: side_use[s] for s in SIDES})
    conn_ch = defaultdict(set)
    for eid in usage:
        conn_ch[eid[1]].add((eid[2], eid[3]))
    add('V14', '端口', '64 个接头各使用 32 个通道', {'connectors': 64, 'channels': 32},
        {'connectors': len(conn_ch),
         'channels': sorted({len(v) for v in conn_ch.values()})},
        passed=(len(conn_ch) == 64 and all(len(v) == 32 for v in conn_ch.values())))

    # --------------------------------------------------- coordinate model
    bad_coord = []
    for eid in usage:
        x, y, z = model_xyz(eid[0], int(eid[1][1:]), eid[2], eid[3])
        if abs(x - PORT_COORDS[eid][0]) > 1e-9 or abs(y - PORT_COORDS[eid][1]) > 1e-9 \
                or z != PORT_COORDS[eid][2]:
            bad_coord.append(eid)
        if not (0.0 <= x <= BOARD_MM and 0.0 <= y <= BOARD_MM):
            bad_coord.append(eid)
    add('V15', '坐标', '端口身份→(x,y,z) 物理模型独立复算一致且均在 150×150 板内',
        [], bad_coord[:5], passed=not bad_coord)
    # compare with the repository fslw/ports.py model
    fslw_ok, fslw_err = False, None
    fslw_bad = []
    try:
        if str(REPO) not in sys.path:
            sys.path.insert(0, str(REPO))
        from four_side_1024.fslw.ports import build_plane_endpoints
        eps = build_plane_endpoints()
        fslw_table = {}
        for ep in eps.values():
            fslw_table[(ep['side'], ep['connector'], ep['row'], ep['column'])] = (
                float(ep['routing_xy'][0]), float(ep['routing_xy'][1]), float(ep['z_mm']))
        for eid in usage:
            ref = fslw_table.get(eid)
            if ref is None or abs(ref[0] - PORT_COORDS[eid][0]) > 1e-9 \
                    or abs(ref[1] - PORT_COORDS[eid][1]) > 1e-9 \
                    or abs(ref[2] - PORT_COORDS[eid][2]) > 1e-9:
                fslw_bad.append((eid, ref))
        fslw_ok = not fslw_bad
    except Exception as exc:                                # noqa: BLE001
        fslw_err = f'{type(exc).__name__}: {exc}'
    add('V16', '坐标', '与仓库 four_side_1024/fslw/ports.py build_plane_endpoints() 完全一致',
        True, fslw_ok, passed=fslw_ok,
        detail='' if fslw_err is None else fslw_err)
    bad_z = [eid for eid in usage if PORT_COORDS[eid][2] != Z_MM[eid[2]]]
    add('V17', '坐标', '物理高度 row1→0.0 mm、row2→0.5 mm', [], bad_z[:5],
        passed=not bad_z)
    add('V18', '坐标', '2048 个端口的三维坐标互不相同（两排 XY 可重合、z 不同）',
        2048, len({PORT_COORDS[eid] for eid in usage}))

    # ------------------------------------------------------- symmetry
    conns = [frozenset((eid_of(r, 'source'), eid_of(r, 'target'))) for r in rows]
    conn_set = set(conns)
    missing = {}
    for k in (1, 2, 3):
        missing[k] = sum(1 for c in conns
                         if frozenset(rot90_eid(e, k) for e in c) not in conn_set)
    add('V19', '对称', '全连接集合 rot90 后仍为同一无序连接集合（端口级）', 0, missing[1])
    add('V20', '对称', '全连接集合 rot180 后仍为同一无序连接集合（端口级）', 0, missing[2])
    add('V21', '对称', '全连接集合 rot270 后仍为同一无序连接集合（端口级）', 0, missing[3])
    orbits = []
    seen = set()
    bad_orbit = []
    for c in conns:
        if c in seen:
            continue
        orbit = [c]
        cur = c
        for _ in range(3):
            cur = rot90_conn(cur)
            orbit.append(cur)
        if len(set(orbit)) != 4 or rot90_conn(cur) != c:
            bad_orbit.append(sorted(orbit))
        for x in orbit:
            seen.add(x)
        orbits.append(orbit)
    add('V22', '对称', '每条连接在其四元旋转轨道中有另外 3 条连接，共 256 组',
        {'groups': 256, 'orbit_errors': 0},
        {'groups': len(orbits), 'orbit_errors': len(bad_orbit)},
        passed=(len(orbits) == 256 and not bad_orbit and len(seen) == 1024))

    # ------------------------------------------------------------- V2
    v2_raw = V2_CSV.read_bytes()
    v2_rows = list(csv.DictReader(v2_raw.decode('utf-8-sig').splitlines()))
    add('V23', 'V2', 'V2 输入 CSV SHA256 与冻结值一致', V2_SHA256,
        hashlib.sha256(v2_raw).hexdigest())
    v2_by_id = {int(r['route_id']): r for r in v2_rows}
    v3_by_id = {int(r['route_id']): r for r in rows}
    diff = []
    for rid in range(1, 513):
        a, b = v3_by_id.get(rid), v2_by_id.get(rid)
        if a is None or b is None:
            diff.append((rid, 'missing'))
            continue
        for f in CSV_FIELDS:
            if f in ('source_row', 'source_column', 'source_channel',
                     'target_row', 'target_column', 'target_channel'):
                if int(a[f]) != int(b[f]):
                    diff.append((rid, f, a[f], b[f]))
            elif str(a[f]) != str(b[f]):
                diff.append((rid, f, a[f], b[f]))
    add('V24', 'V2', 'V2 对边 512 条（route_id 1..512）在 V3 CSV 中逐字段完全一致', 0, len(diff),
        detail=str(diff[:5]))
    v2_conn = set()
    for r in v2_rows:
        if int(r['source_row']) == 1:
            v2_conn.add(frozenset((eid_of(r, 'source'), eid_of(r, 'target'))))
    v2_missing = {}
    for k in (1, 2, 3):
        v2_missing[k] = sum(1 for c in v2_conn
                            if frozenset(rot90_eid(e, k) for e in c) not in v2_conn)
    add('V25', 'V2', 'V2 原对边 512 条在真实物理坐标下严格 90° 旋转对称',
        {'rot90': 0, 'rot180': 0, 'rot270': 0},
        {'rot90': v2_missing[1], 'rot180': v2_missing[2], 'rot270': v2_missing[3]})
    v2_adj = set()
    for r in v2_rows:
        if int(r['source_row']) == 2:
            v2_adj.add(frozenset((eid_of(r, 'source'), eid_of(r, 'target'))))
    add('V26', 'V2', 'V2 原邻边 512 条不具备 90° 旋转对称（说明必须重设计）',
        {'missing_rot90': 512},
        {'missing_rot90': sum(1 for c in v2_adj
                             if frozenset(rot90_eid(e) for e in c) not in v2_adj)})

    # ----------------------------------------------------------- Excel
    try:
        import openpyxl
    except Exception as exc:                                # noqa: BLE001
        add('V27', 'Excel', 'openpyxl 可用', True, False, passed=False, detail=str(exc))
        openpyxl = None
    if openpyxl is not None:
        wb = openpyxl.load_workbook(OUT_XLSX, data_only=True)
        expect_sheets = ['1024连接总表', '四边接口表', '旋转对称组', '统计与验证', '来源与说明']
        add('V27', 'Excel', '工作簿可打开且恰好 5 个指定 Sheet', expect_sheets, wb.sheetnames)

        ws1 = wb['1024连接总表']
        rows1 = list(ws1.iter_rows(min_row=2, values_only=True))
        hdr1 = [c.value for c in ws1[1]]
        add('V28', 'Excel', 'Sheet1 连接总表 1024 行、前 12 列表头与 CSV 相同',
            {'rows': 1024, 'fields': CSV_FIELDS},
            {'rows': len(rows1), 'fields': hdr1[:12]})
        bad1 = []
        for row in rows1:
            rid = row[0]
            r = v3_by_id.get(rid)
            if r is None:
                bad1.append(('route', rid))
                continue
            actual = [int(row[0]), row[1], row[2], int(row[3]), int(row[4]), int(row[5]),
                      row[6], row[7], int(row[8]), int(row[9]), int(row[10]), row[11]]
            expected = [int(r['route_id']), r['source_side'], r['source_connector'],
                        int(r['source_row']), int(r['source_column']),
                        int(r['source_channel']), r['target_side'], r['target_connector'],
                        int(r['target_row']), int(r['target_column']),
                        int(r['target_channel']), r['relation']]
            if actual != expected:
                bad1.append((rid, actual, expected))
            for ci, eid in ((13, eid_of(r, 'source')), (17, eid_of(r, 'target'))):
                x, y, z = PORT_COORDS[eid]
                if abs(float(row[ci]) - x) > 1e-9 or abs(float(row[ci + 1]) - y) > 1e-9 \
                        or abs(float(row[ci + 2]) - z) > 1e-9:
                    bad1.append((rid, 'coords', eid))
        add('V29', 'Excel', 'Sheet1 的 1024 条连接与 CSV 逐字段一致、坐标/Z 高度符合物理模型',
            [], bad1[:5], passed=not bad1, detail=f'violations={len(bad1)}')

        ws2 = wb['四边接口表']
        rows2 = list(ws2.iter_rows(min_row=2, values_only=True))
        add('V30', 'Excel', 'Sheet2 四边接口表 2048 行', 2048, len(rows2))
        peer_of = {}
        for r in rows:
            a, b = eid_of(r, 'source'), eid_of(r, 'target')
            peer_of[a] = (int(r['route_id']), b)
            peer_of[b] = (int(r['route_id']), a)
        seen2 = Counter()
        bad2 = []
        for row in rows2:
            eid = (row[0], row[1], int(row[2]), int(row[3]))
            seen2[eid] += 1
            if eid not in PORT_COORDS:
                bad2.append(('unknown', eid))
                continue
            x, y, z = PORT_COORDS[eid]
            if abs(float(row[5]) - x) > 1e-9 or abs(float(row[6]) - y) > 1e-9 \
                    or abs(float(row[7]) - z) > 1e-9:
                bad2.append(('coords', eid))
            rid, other = peer_of[eid]
            if int(row[10]) != rid:
                bad2.append(('route_id', eid))
            if (row[13], row[14], int(row[15]), int(row[16])) != other:
                bad2.append(('peer', eid))
            ox, oy, oz = PORT_COORDS[other]
            if abs(float(row[19]) - ox) > 1e-9 or abs(float(row[20]) - oy) > 1e-9 \
                    or abs(float(row[21]) - oz) > 1e-9:
                bad2.append(('peer_coords', eid))
        add('V31', 'Excel', 'Sheet2 的 2048 个端口各出现一次且对端端口/坐标与 CSV 一致',
            {'unique': 2048, 'used_once': True},
            {'unique': len(seen2), 'used_once': set(seen2.values()) == {1}},
            passed=(len(seen2) == 2048 and set(seen2.values()) == {1} and not bad2),
            detail=f'violations={len(bad2)}')

        ws3 = wb['旋转对称组']
        rows3 = list(ws3.iter_rows(min_row=2, values_only=True))
        conn_by_rid = {int(r['route_id']): frozenset((eid_of(r, 'source'), eid_of(r, 'target')))
                       for r in rows}
        all_rids = []
        bad3 = []
        for row in rows3:
            rids3 = [int(row[2]), int(row[3]), int(row[4]), int(row[5])]
            all_rids.extend(rids3)
            if any(x not in conn_by_rid for x in rids3):
                bad3.append(('unknown', rids3))
                continue
            cur = conn_by_rid[rids3[0]]
            for k in range(1, 4):
                cur = rot90_conn(cur)
                if cur != conn_by_rid[rids3[k]]:
                    bad3.append(('phase', rids3, k))
            if rot90_conn(conn_by_rid[rids3[3]]) != conn_by_rid[rids3[0]]:
                bad3.append(('closure', rids3))
        add('V32', 'Excel', 'Sheet3 恰好 256 组、每组 4 条、覆盖全部 1024 条且 90° 相位正确',
            {'groups': 256, 'unique_routes': 1024, 'errors': 0},
            {'groups': len(rows3), 'unique_routes': len(set(all_rids)),
             'errors': len(bad3)},
            passed=(len(rows3) == 256 and sorted(all_rids) == list(range(1, 1025))
                    and not bad3))

        fmt = {}
        for name in ('1024连接总表', '四边接口表', '旋转对称组'):
            ws = wb[name]
            fmt[name] = dict(freeze=str(ws.freeze_panes),
                             autofilter=str(ws.auto_filter.ref),
                             bold=bool(ws['A1'].font.bold))
        add('V33', 'Excel', '数据表冻结表头、启用自动筛选、表头加粗',
            {n: 'freeze+filter+bold' for n in fmt},
            {n: ('freeze+filter+bold' if (v['freeze'] not in ('None', 'A1')
                                          and v['autofilter'] != 'None' and v['bold'])
                 else f"freeze={v['freeze']} filter={v['autofilter']} bold={v['bold']}")
             for n, v in fmt.items()},
            passed=all(v['freeze'] not in ('None', 'A1') and v['autofilter'] != 'None'
                       and v['bold'] for v in fmt.values()))

        ws4 = wb['统计与验证']
        results4 = [row[5] for row in ws4.iter_rows(min_row=2, values_only=True)
                    if row[5] in ('PASS', 'FAIL')]
        add('V34', 'Excel', 'Sheet4 统计与验证：全部校验结果为 PASS、无 FAIL',
            {'fail': 0, 'pass>=30': True},
            {'fail': sum(1 for x in results4 if x == 'FAIL'),
             'pass': sum(1 for x in results4 if x == 'PASS')},
            passed=(sum(1 for x in results4 if x == 'FAIL') == 0
                    and sum(1 for x in results4 if x == 'PASS') >= 30))

        ws5 = wb['来源与说明']
        texts5 = [str(cell) for row in ws5.iter_rows(min_row=2, values_only=True)
                  for cell in row if cell is not None]
        csv_sha = hashlib.sha256(raw).hexdigest()
        add('V35', 'Excel', 'Sheet5 来源与说明记录输出 CSV SHA256 与 V2 输入 SHA256',
            {'csv_sha_recorded': True, 'v2_sha_recorded': True},
            {'csv_sha_recorded': any(csv_sha in t for t in texts5),
             'v2_sha_recorded': any(V2_SHA256 in t for t in texts5)})

    # -------------------------------------------- validation JSON integrity
    try:
        payload = json.loads(OUT_JSON.read_text(encoding='utf-8'))
        add('V36', '记录', 'validation_v3_c4.json 存在且 summary.all_pass 为 true',
            {'all_pass': True}, {'all_pass': payload.get('summary', {}).get('all_pass')},
            passed=bool(payload.get('summary', {}).get('all_pass')))
        recorded_csv = payload.get('files', {}).get('csv', {}).get('sha256')
        recorded_xlsx = payload.get('files', {}).get('xlsx', {}).get('sha256')
        actual_csv = sha256_file(OUT_CSV)
        actual_xlsx = sha256_file(OUT_XLSX)
        add('V37', '记录', 'validation JSON 记录的 CSV/XLSX SHA256 与实际文件一致',
            {'csv': actual_csv, 'xlsx': actual_xlsx},
            {'csv': recorded_csv, 'xlsx': recorded_xlsx},
            passed=(recorded_csv == actual_csv and recorded_xlsx == actual_xlsx))
    except Exception as exc:                                # noqa: BLE001
        add('V36', '记录', 'validation_v3_c4.json 可读取', True, False,
            passed=False, detail=f'{type(exc).__name__}: {exc}')
        add('V37', '记录', 'validation JSON 记录的哈希与实际文件一致', True, False,
            passed=False, detail='skipped')

    # ------------------------------------------------------- file presence
    add('V38', '文件', '生成脚本、生成报告、配对比较文件均存在',
        {'generator': True, 'report': True, 'comparison': True},
        {'generator': GENERATOR.exists(), 'report': OUT_MD.exists(),
         'comparison': OUT_COMPARISON.exists()})
    add('V39', '文件', '历史 V1/V2 CSV 未被修改（SHA256 不变）',
        {'V1': V1_SHA256, 'V2': V2_SHA256},
        {'V1': sha256_file(V1_CSV) if V1_CSV.exists() else 'missing',
         'V2': sha256_file(V2_CSV) if V2_CSV.exists() else 'missing'})

    # ----------------------------------------------------------- summary
    n_pass = sum(1 for c in CHECKS if c['pass_'])
    n_fail = sum(1 for c in CHECKS if not c['pass_'])
    print()
    for c in CHECKS:
        print(f"  [{'PASS' if c['pass_'] else 'FAIL'}] {c['id']:<4} {c['name']}")
    print()
    print(f'RESULT: {n_pass} PASS / {n_fail} FAIL')
    OUT_INDEPENDENT.write_text(json.dumps(dict(
        schema='v3_c4_independent_acceptance/1.0',
        script=Path(__file__).name,
        verified_from_saved_files=True,
        independent_of_generator=True,
        checks=CHECKS,
        summary=dict(total=len(CHECKS), pass_count=n_pass, fail_count=n_fail,
                     all_pass=(n_fail == 0)),
    ), indent=2, ensure_ascii=False), encoding='utf-8')
    print('written:', OUT_INDEPENDENT)
    return 0 if n_fail == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
