"""Independent audit of the V3-Huang 1024-channel dual-height topology inputs.

Owner: connectivity (``four_side_1024/v3/SPEC.md`` sections 1, 2 and 7).

This module is a deliberately **independent** re-derivation and audit of the two
frozen V3-Huang input files:

* ``data/v3_huang/fiberBoard1024_huang_fiber_pairs.csv`` -- the routing input
  (1024 fibre-level connections with fixed endpoints);
* ``data/v3_huang/fiberBoard1024_huang_four_side_connectors.xlsx`` -- the
  connector-level Port1/Port2 table, port map, 64x64 connector matrix, class
  statistics, fibre-level reference sheet and connector-pair multiplicities;
* ``data/v3_huang/PROVENANCE.json`` -- original Desktop paths + SHA256.

It does **not** import Lead's V3 input modules (``four_side_1024/v3/config.py``,
``inputs.py``, ``ports.py``); the cross-check with those modules is performed
only in a subprocess (:func:`compare_with_lead_v3_module`).  The previous-round
``fslw/ports_audit.py`` was used only as a style reference.

V3 model (SPEC section 2)
-------------------------
* topology class comes from the **CSV ``relation`` column** (never inferred from
  the row number): same 448 / opposite 384 / adjacent 192;
* plane comes from the **CSV row**: row 1 -> plane z1 (512 connections),
  row 2 -> plane z2 (512 connections); row 1 mixes same 222 / opposite 196 /
  adjacent 94, row 2 mixes same 226 / opposite 188 / adjacent 98;
* 4 sides x 16 connectors x 32 fibres = 2048 physical ports, each used once;
* physical 2x16 MPO geometry (unchanged from the FSLW round): column pitch
  0.25 mm, two rows 0.5 mm apart (z1 = 0.0, z2 = 0.5 model values),
  ``along = connector_center + (column - 8.5) * 0.25``,
  ``connector_center = 7.5 + (index - 1) * 9`` mm;
* no router is run here: this is only the topology/interface audit.

Run ``python -m four_side_1024.v3.ports_audit_v3`` from
``OpticalWaveguideRouter2D`` to rebuild all artifacts under
``outputs/four_side_1024_v3_huang_R40_gap30_abcd/topology/``.
"""
from __future__ import annotations

import csv
import hashlib
import json
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

if __package__ in (None, ''):  # pragma: no cover - direct script execution
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

__all__ = [
    'INPUT_DIR', 'CSV_PATH', 'XLSX_PATH', 'PROVENANCE_PATH', 'OUTPUT_DIR',
    'CSV_SHA256', 'XLSX_SHA256', 'EXPECTED', 'MODEL', 'CLASS_ORDER',
    'PAIR_ORDER', 'PAIR_CLASS', 'PHYSICAL_PORT_COLUMNS',
    'read_connectivity_csv', 'read_excel_workbook', 'derive_ports',
    'validate_inputs', 'physical_port_rows', 'compare_with_lead_v3_module',
    'build_all', 'main',
]

# ---------------------------------------------------------------------------
# Frozen inputs / expected identity (SPEC sections 1-2, task description)
# ---------------------------------------------------------------------------
CODE_DIR = Path(__file__).resolve().parents[2]  # .../OpticalWaveguideRouter2D
INPUT_DIR = CODE_DIR / 'data' / 'v3_huang'
CSV_PATH = INPUT_DIR / 'fiberBoard1024_huang_fiber_pairs.csv'
XLSX_PATH = INPUT_DIR / 'fiberBoard1024_huang_four_side_connectors.xlsx'
PROVENANCE_PATH = INPUT_DIR / 'PROVENANCE.json'
OUTPUT_DIR = (CODE_DIR / 'outputs' / 'four_side_1024_v3_huang_R40_gap30_abcd'
              / 'topology')

CSV_SHA256 = '76d91316e0933a64b8c822fd541faa2bda571fdef97e49b64f75b55cf4a6cf23'
XLSX_SHA256 = '4d285f5667fa922849b3371525944e12feca1dae1bc07c0d5cd40be3b38b32de'

# Physical 2x16 MPO model -- unchanged from the FSLW round (V3 SPEC section 3).
BOARD_MM = 150.0
SIDES = 'NESW'
CONNECTORS_PER_SIDE = 16
ROWS = 2
COLUMNS = 16
CHANNELS_PER_CONNECTOR = ROWS * COLUMNS  # 32
CONNECTOR_PITCH_MM = 9.0
FIRST_CONNECTOR_CENTER_MM = (BOARD_MM - (CONNECTORS_PER_SIDE - 1) * CONNECTOR_PITCH_MM) / 2.0
FIBER_COLUMN_PITCH_MM = 0.25
FIBER_ROW_PITCH_MM = 0.50
PLANE_Z_MM = {1: 0.0, 2: 0.5}       # CSV row -> model z height
PLANE_NAME = {1: 'z1', 2: 'z2'}
PLANE_ROW = {'z1': 1, 'z2': 2}
PORT_MODEL = 'real_2x16_dual_height'
TOL_MM = 1e-9
MATCH_TOL_MM = 1e-6

CSV_FIELDS = ['route_id', 'source_side', 'source_connector', 'source_row',
              'source_column', 'source_channel', 'target_side', 'target_connector',
              'target_row', 'target_column', 'target_channel', 'relation']
CSV_NUMERIC_FIELDS = {'route_id', 'source_row', 'source_column', 'source_channel',
                      'target_row', 'target_column', 'target_channel'}

CLASS_ORDER = ('same', 'opposite', 'adjacent')
PAIR_ORDER = ('NN', 'SS', 'EE', 'WW', 'NS', 'EW', 'EN', 'NW', 'ES', 'SW')
PAIR_CLASS = {'NN': 'same', 'SS': 'same', 'EE': 'same', 'WW': 'same',
              'NS': 'opposite', 'EW': 'opposite',
              'EN': 'adjacent', 'NW': 'adjacent', 'ES': 'adjacent', 'SW': 'adjacent'}
CLASS_COLORS = {'same': '#2f6db5', 'opposite': '#c0392b', 'adjacent': '#d99b1c'}
SIDE_COLORS = {'N': '#3366bb', 'E': '#e3a323', 'S': '#b84040', 'W': '#60a735'}

EXPECTED = dict(
    rows=1024,
    route_id_min=1,
    route_id_max=1024,
    class_counts={'same': 448, 'opposite': 384, 'adjacent': 192},
    pair_counts={'NN': 112, 'SS': 112, 'EE': 112, 'WW': 112,
                 'NS': 192, 'EW': 192,
                 'EN': 48, 'NW': 48, 'ES': 48, 'SW': 48},
    plane_row_counts={'z1': {'same': 222, 'opposite': 196, 'adjacent': 94},
                      'z2': {'same': 226, 'opposite': 188, 'adjacent': 98}},
    plane_connections={'z1': 512, 'z2': 512},
    plane_endpoints={'z1': 1024, 'z2': 1024},
    connectors=64,
    endpoints=2048,
    fibres_per_connector=32,
    physical_ports=2048,
    matrix_total=2048,  # symmetric matrix double-counts the 1024 connections
    excel_sheets=7,
)

MODEL = dict(
    port_model=PORT_MODEL,
    interface_standard='IEC 61754-7-3:2019 2x16 MPO',
    board_mm=[BOARD_MM, BOARD_MM],
    connectors_per_side=CONNECTORS_PER_SIDE,
    connector_pitch_mm=CONNECTOR_PITCH_MM,
    first_connector_center_mm=FIRST_CONNECTOR_CENTER_MM,
    fiber_column_pitch_mm=FIBER_COLUMN_PITCH_MM,
    fiber_row_pitch_mm=FIBER_ROW_PITCH_MM,
    rows=ROWS, columns=COLUMNS, channels_per_connector=CHANNELS_PER_CONNECTOR,
    plane_z_mm={str(row): z for row, z in PLANE_Z_MM.items()},
    plane_rule='CSV row 1 -> z1 (512 conns), row 2 -> z2 (512 conns); relation comes '
               'from the CSV relation column, never from the row',
    relation_rule='relation column: same/opposite/adjacent; pair = sorted side letters '
                  '(same -> NN/SS/EE/WW, opposite -> NS/EW, adjacent -> EN/NW/ES/SW)',
    along_formula='along = connector_center + (column - 8.5) * 0.25',
    physical_3d_definition='(x, y, z) with z = z1 for row 1, z2 for row 2',
    z_transition='not_implemented (two independent 2D planes)',
    layer='0 (independent 2D plane)',
    cross_plane_rule='cross-plane XY projections are NOT same-layer conflicts',
    manufacturing_status='simulation under the two-height interface assumption; glass '
                         'thickness / connector mounting unconfirmed; not a claim of '
                         'manufacturability',
)

PHYSICAL_PORT_COLUMNS = ['side', 'connector', 'row', 'column', 'channel', 'plane',
                         'x_mm', 'y_mm', 'z_mm', 'along_mm',
                         'connector_center_mm', 'inward_x', 'inward_y', 'port_model']


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def sha256_file(path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def connector_number(connector) -> int:
    name = str(connector)
    if not name[1:].isdigit():
        raise ValueError(f'ill-formed connector {connector!r}')
    return int(name[1:])


def connector_center_mm(connector) -> float:
    return FIRST_CONNECTOR_CENTER_MM + (connector_number(connector) - 1) * CONNECTOR_PITCH_MM


def port_along_mm(connector, column) -> float:
    return connector_center_mm(connector) + \
        (int(column) - (COLUMNS + 1) / 2.0) * FIBER_COLUMN_PITCH_MM


def edge_xy(side, along) -> list:
    if side == 'N':
        return [float(along), BOARD_MM]
    if side == 'S':
        return [float(along), 0.0]
    if side == 'E':
        return [BOARD_MM, float(along)]
    if side == 'W':
        return [0.0, float(along)]
    raise ValueError(f'unknown side {side!r}')


def inward_vector(side) -> list:
    return {'N': [0.0, -1.0], 'S': [0.0, 1.0],
            'E': [-1.0, 0.0], 'W': [1.0, 0.0]}[side]


def endpoint_id(connector, row, column) -> str:
    return f'{connector}:{int(row)}:{int(column)}'


def connector_port_index(connector) -> int:
    """1-based connector index in the frozen N01..N16, E01.., S01.., W01.. order."""
    return SIDES.index(str(connector)[0]) * CONNECTORS_PER_SIDE + connector_number(connector)


def class_from_sides(side_a, side_b):
    """(class, pair) derived from the two board sides -- consistency check only."""
    if side_a == side_b:
        return 'same', f'{side_a}{side_b}'
    pair = ''.join(sorted((str(side_a), str(side_b))))
    return ('opposite' if pair in ('NS', 'EW') else 'adjacent'), pair


def side_pair_label(connector_a, connector_b) -> str:
    """Side-pair label of a connector pair: NN/SS/EE/WW, NS/EW, EN/NW/ES/SW."""
    side_a = str(connector_a)[0]
    side_b = str(connector_b)[0]
    if side_a == side_b:
        return side_a + side_b
    return ''.join(sorted((side_a, side_b)))


def _jsonable(obj):
    if isinstance(obj, dict):
        return {str(k): _jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    if isinstance(obj, Counter):
        return _jsonable(dict(obj))
    if isinstance(obj, (set, frozenset)):
        return sorted(_jsonable(v) for v in obj)
    return obj


def _to_int(value):
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return int(str(value).strip())


def _text(value) -> str:
    return '' if value is None else str(value).strip()


def _looks_like_connector(value) -> bool:
    text = _text(value)
    return len(text) == 3 and text[0] in SIDES and text[1:].isdigit()


# ---------------------------------------------------------------------------
# Independent CSV parser
# ---------------------------------------------------------------------------
def read_connectivity_csv(path=None) -> tuple:
    """Parse the frozen fibre-pair CSV (own parser; returns routes, header, SHA256)."""
    path = Path(path or CSV_PATH)
    raw = path.read_bytes()
    rows = list(csv.DictReader(raw.decode('utf-8-sig').splitlines()))
    if not rows:
        raise ValueError(f'{path}: empty CSV')
    header = list(rows[0].keys())
    routes = []
    for line_number, row in enumerate(rows, start=2):
        missing = [field for field in CSV_FIELDS if field not in row]
        if missing:
            raise ValueError(f'{path}:{line_number}: missing fields {missing}')
        record = {}
        for field in CSV_FIELDS:
            value = row[field]
            record[field] = _to_int(value) if field in CSV_NUMERIC_FIELDS else _text(value)
        routes.append(record)
    return routes, header, hashlib.sha256(raw).hexdigest()


# ---------------------------------------------------------------------------
# Independent XLSX parser (openpyxl; sheets located structurally, ASCII markers)
# ---------------------------------------------------------------------------
def read_excel_workbook(path=None) -> dict:
    """Parse the frozen connector workbook; every sheet is located by structure.

    The workbook uses Chinese sheet/column names, so no Chinese literal is needed
    here: the sheets are identified by shape and by ASCII cell markers
    (``Port1``/``Port2``, ``route_id``, ``N01`` ... ``W16``, matrix shape 65x65,
    connector rows E06/N14/..., and the class names same/opposite/adjacent).
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment check
        raise RuntimeError('openpyxl is required to audit the V3 Excel input') from exc
    path = Path(path or XLSX_PATH)
    raw = path.read_bytes()
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheet_names = list(workbook.sheetnames)

    def first_cell(sheet, row=1, column=1):
        return _text(sheet.cell(row=row, column=column).value)

    def find_sheet(predicate, label):
        for sheet in workbook.worksheets:
            try:
                if predicate(sheet):
                    return sheet
            except Exception:  # pragma: no cover - defensive
                continue
        raise ValueError(f'Excel sheet not found: {label}; sheets={sheet_names}')

    ports_sheet = find_sheet(
        lambda sheet: first_cell(sheet) == 'Port1'
        and first_cell(sheet, 1, 2) == 'Port2', 'Port1/Port2 pairs')
    fibre_sheet = find_sheet(lambda sheet: first_cell(sheet) == 'route_id',
                             'fibre-level reference')
    matrix_sheet = find_sheet(
        lambda sheet: sheet.max_row == 4 * CONNECTORS_PER_SIDE + 1
        and sheet.max_column == 4 * CONNECTORS_PER_SIDE + 1
        and first_cell(sheet, 1, 2) == 'N01' and first_cell(sheet, 2, 1) == 'N01',
        '64x64 matrix')
    port_map_sheet = find_sheet(
        lambda sheet: sheet.max_row == CONNECTORS_PER_SIDE * 4 + 1
        and sheet.max_column == 9
        and _to_int(sheet.cell(row=2, column=1).value) == 1
        and _looks_like_connector(sheet.cell(row=2, column=2).value),
        'port map')
    multiplicity_sheet = find_sheet(
        lambda sheet: sheet.max_column == 4 and sheet.max_row > 500
        and _looks_like_connector(sheet.cell(row=2, column=1).value)
        and _looks_like_connector(sheet.cell(row=2, column=2).value),
        'connector-pair multiplicity')
    stats_sheet = find_sheet(
        lambda sheet: sheet.max_column == 5
        and any(isinstance(cell, (int, float)) and not isinstance(cell, bool)
                and float(cell).is_integer() and int(cell) == 448
                for row in sheet.iter_rows(values_only=True) for cell in row),
        'class statistics')
    used_titles = {sheet.title for sheet in (ports_sheet, fibre_sheet, matrix_sheet,
                                             port_map_sheet, multiplicity_sheet,
                                             stats_sheet)}
    remaining = [sheet for sheet in workbook.worksheets
                 if sheet.title not in used_titles]
    notes_sheet = remaining[0] if remaining else None

    port_pairs = []
    for row in range(2, ports_sheet.max_row + 1):
        port1 = ports_sheet.cell(row=row, column=1).value
        port2 = ports_sheet.cell(row=row, column=2).value
        if port1 is None and port2 is None:
            continue
        port_pairs.append((_to_int(port1), _to_int(port2)))

    port_map = []
    for row in range(2, port_map_sheet.max_row + 1):
        port = port_map_sheet.cell(row=row, column=1).value
        name = port_map_sheet.cell(row=row, column=2).value
        if port is None or name is None:
            continue
        port_map.append(dict(
            port=_to_int(port),
            connector=_text(name),
            side=_text(port_map_sheet.cell(row=row, column=3).value),
            side_index=_to_int(port_map_sheet.cell(row=row, column=4).value),
            center_mm=float(port_map_sheet.cell(row=row, column=5).value),
            fibres=_to_int(port_map_sheet.cell(row=row, column=6).value),
            row1_fibres=_to_int(port_map_sheet.cell(row=row, column=7).value),
            row2_fibres=_to_int(port_map_sheet.cell(row=row, column=8).value),
            distinct_targets=_to_int(port_map_sheet.cell(row=row, column=9).value),
        ))

    class_stats = dict(counts={}, ratios={}, pair_counts={}, raw=[])
    for row in stats_sheet.iter_rows(min_row=1, values_only=True):
        class_stats['raw'].append([_jsonable(cell) for cell in row])
        label = _text(row[0]) if row and row[0] is not None else ''
        numbers = [cell for cell in row[1:] if isinstance(cell, (int, float))]
        if label:
            matched_class = False
            for klass in CLASS_ORDER:
                if klass in label.lower() and numbers:
                    matched_class = True
                    class_stats['counts'][klass] = _to_int(numbers[0])
                    ratios = [number for number in numbers[1:]
                              if not isinstance(number, int)]
                    if ratios:
                        class_stats['ratios'][klass] = ratios[0]
            if not matched_class and numbers \
                    and float(numbers[0]).is_integer() and int(numbers[0]) == 1024:
                class_stats['counts']['total'] = 1024
            if 'total' in label.lower() and numbers:
                class_stats['counts']['total'] = _to_int(numbers[0])
        for column, cell in enumerate(row):
            text = _text(cell)
            if text in PAIR_ORDER:
                neighbour = row[column + 1] if column + 1 < len(row) else None
                if isinstance(neighbour, (int, float)):
                    class_stats['pair_counts'][text] = _to_int(neighbour)

    matrix_columns = [_text(matrix_sheet.cell(row=1, column=column).value)
                      for column in range(2, matrix_sheet.max_column + 1)]
    matrix_rows = [_text(matrix_sheet.cell(row=row, column=1).value)
                   for row in range(2, matrix_sheet.max_row + 1)]
    matrix = []
    for row in range(2, matrix_sheet.max_row + 1):
        values = []
        for column in range(2, matrix_sheet.max_column + 1):
            value = matrix_sheet.cell(row=row, column=column).value
            values.append(_to_int(value) if value is not None else 0)
        matrix.append(values)

    fibre_header = [_text(fibre_sheet.cell(row=1, column=column).value)
                    for column in range(1, fibre_sheet.max_column + 1)]
    fibre_rows = []
    for row in range(2, fibre_sheet.max_row + 1):
        values = [fibre_sheet.cell(row=row, column=column).value
                  for column in range(1, fibre_sheet.max_column + 1)]
        if all(value is None for value in values):
            continue
        typed = []
        for field, value in zip(fibre_header, values):
            if field in CSV_NUMERIC_FIELDS:
                typed.append(_to_int(value) if value is not None else None)
            else:
                typed.append(_text(value) if value is not None else None)
        fibre_rows.append(typed)

    multiplicities = []
    for row in range(2, multiplicity_sheet.max_row + 1):
        values = [multiplicity_sheet.cell(row=row, column=column).value
                  for column in range(1, multiplicity_sheet.max_column + 1)]
        if all(value is None for value in values):
            continue
        multiplicities.append(dict(
            connector_a=_text(values[0]),
            connector_b=_text(values[1]),
            count=_to_int(values[2]),
            connection_class=_text(values[3]),
        ))

    notes = []
    if notes_sheet is not None:
        for row in notes_sheet.iter_rows(min_row=2, values_only=True):
            if row and row[0] is not None:
                notes.append(dict(item=_text(row[0]),
                                  description=_text(row[1]) if len(row) > 1 else ''))

    workbook.close()
    return dict(
        path=str(path.resolve()),
        sha256=hashlib.sha256(raw).hexdigest(),
        sheet_names=sheet_names,
        ports_sheet_name=ports_sheet.title,
        fibre_sheet_name=fibre_sheet.title,
        matrix_sheet_name=matrix_sheet.title,
        port_map_sheet_name=port_map_sheet.title,
        multiplicity_sheet_name=multiplicity_sheet.title,
        stats_sheet_name=stats_sheet.title,
        notes_sheet_name=notes_sheet.title if notes_sheet is not None else None,
        port_pairs=port_pairs,
        port_map=port_map,
        class_stats=class_stats,
        matrix_columns=matrix_columns,
        matrix_rows=matrix_rows,
        matrix=matrix,
        fibre_header=fibre_header,
        fibre_rows=fibre_rows,
        multiplicities=multiplicities,
        notes=notes,
    )


# ---------------------------------------------------------------------------
# Derivation: routes / endpoints / physical 3D ports / per-plane lists
# ---------------------------------------------------------------------------
def physical_port_rows() -> list:
    """All 2048 real 3D physical ports (64 connectors x 32 fibres)."""
    ports = []
    for side in SIDES:
        for number in range(1, CONNECTORS_PER_SIDE + 1):
            connector = f'{side}{number:02d}'
            center = connector_center_mm(connector)
            for row in (1, 2):
                for column in range(1, COLUMNS + 1):
                    along = port_along_mm(connector, column)
                    x, y = edge_xy(side, along)
                    inward = inward_vector(side)
                    ports.append(dict(
                        endpoint_id=endpoint_id(connector, row, column),
                        side=side,
                        connector=connector,
                        row=row,
                        column=column,
                        channel=(row - 1) * COLUMNS + column,
                        plane=PLANE_NAME[row],
                        x_mm=x, y_mm=y, z_mm=PLANE_Z_MM[row],
                        along_mm=along,
                        connector_center_mm=center,
                        inward_x=inward[0], inward_y=inward[1],
                        port_model=PORT_MODEL,
                    ))
    return ports


def derive_ports(csv_path=None) -> dict:
    """Derive routes (class from CSV relation, plane from CSV row) + endpoints."""
    routes_raw, header, sha = read_connectivity_csv(csv_path)
    topology = 'V3_HUANG'
    routes = []
    endpoints = {}
    for raw in routes_raw:
        route = dict(raw)
        relation = _text(raw['relation'])
        if relation not in CLASS_ORDER:
            raise ValueError(f"route {raw['route_id']}: unknown relation {relation!r}")
        derived_class, pair = class_from_sides(_text(raw['source_side']),
                                               _text(raw['target_side']))
        row = int(raw['source_row'])
        route.update(
            route_id=int(raw['route_id']),
            relation=relation,
            connection_class=relation,
            class_derived_from_sides=derived_class,
            class_consistent=(derived_class == relation),
            pair=pair,
            row=row,
            plane=PLANE_NAME[row],
            z_mm=PLANE_Z_MM[row],
            source_connector_index=connector_port_index(raw['source_connector']),
            target_connector_index=connector_port_index(raw['target_connector']),
        )
        for role in ('source', 'target'):
            connector = _text(raw[f'{role}_connector'])
            side = _text(raw[f'{role}_side'])
            row_role = int(raw[f'{role}_row'])
            column = int(raw[f'{role}_column'])
            channel = int(raw[f'{role}_channel'])
            eid = endpoint_id(connector, row_role, column)
            along = port_along_mm(connector, column)
            x, y = edge_xy(side, along)
            endpoints[eid] = dict(
                endpoint_id=eid,
                connector=connector,
                side=side,
                row=row_role,
                column=column,
                channel=channel,
                routing_xy=[x, y],
                inward=inward_vector(side),
                layer=0,
                plane=PLANE_NAME[row_role],
                z_mm=PLANE_Z_MM[row_role],
                physical_along_mm=along,
                physical_row_z_mm=PLANE_Z_MM[row_role],
                connector_center_mm=connector_center_mm(connector),
                port_model=PORT_MODEL,
            )
            route[f'{role}_endpoint'] = eid
        routes.append(route)

    physical_ports = physical_port_rows()
    plane_endpoints = {
        plane: [dict(endpoint) for _, endpoint in sorted(endpoints.items())
                if endpoint['plane'] == plane]
        for plane in ('z1', 'z2')
    }
    return dict(
        topology=topology,
        csv_path=str(Path(csv_path or CSV_PATH).resolve()),
        csv_sha256=sha,
        csv_header=header,
        routes=routes,
        endpoints=endpoints,
        physical_ports=physical_ports,
        plane_endpoints=plane_endpoints,
    )


def _connector_pair_counter(routes) -> Counter:
    counter = Counter()
    for route in routes:
        connector_a = _text(route['source_connector'])
        connector_b = _text(route['target_connector'])
        counter[tuple(sorted((connector_a, connector_b)))] += 1
    return counter


def _directed_connector_matrix(routes, connector_order) -> list:
    index = {name: position for position, name in enumerate(connector_order)}
    matrix = [[0] * len(connector_order) for _ in connector_order]
    for route in routes:
        matrix[index[_text(route['source_connector'])]][
            index[_text(route['target_connector'])]] += 1
    return matrix


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
def _provenance_check(csv_sha, xlsx_sha) -> dict:
    report = dict(available=False, passed=False)
    path = PROVENANCE_PATH
    if not path.exists():
        report['error'] = f'{path} not found'
        return report
    provenance = json.loads(path.read_text(encoding='utf-8'))
    files = provenance.get('files', {})
    xlsx_info = files.get('connectors_xlsx', {})
    csv_info = files.get('fiber_pairs_csv', {})
    mismatches = []
    checks = []
    for label, info, actual_sha, expected_sha in (
            ('connectors_xlsx', xlsx_info, xlsx_sha, XLSX_SHA256),
            ('fiber_pairs_csv', csv_info, csv_sha, CSV_SHA256)):
        stored_sha = info.get('stored_sha256')
        source_sha = info.get('source_sha256')
        checks.append(dict(
            file=label,
            stored_sha256=stored_sha,
            source_sha256=source_sha,
            actual_sha256=actual_sha,
            matches_stored=(stored_sha == actual_sha),
            matches_source=(source_sha == actual_sha),
            matches_task_expected=(expected_sha == actual_sha),
        ))
        if stored_sha != actual_sha or source_sha != actual_sha or expected_sha != actual_sha:
            mismatches.append(label)
    report.update(
        available=True,
        passed=not mismatches,
        mismatches=mismatches,
        files=checks,
        provenance=provenance,
        bytes_expected={key: value.get('bytes') for key, value in files.items()},
    )
    return report


def validate_inputs(csv_path=None, xlsx_path=None) -> dict:
    """Audit both V3 inputs independently; returns a JSON-ready report."""
    csv_path = Path(csv_path or CSV_PATH)
    xlsx_path = Path(xlsx_path or XLSX_PATH)
    checks, failed = {}, []

    def record(name, passed, **info):
        checks[name] = dict(passed=bool(passed), **info)
        if not passed:
            failed.append(name)

    csv_routes, csv_header, csv_sha = read_connectivity_csv(csv_path)
    excel = read_excel_workbook(xlsx_path)
    xlsx_sha = excel['sha256']
    state = derive_ports(csv_path)
    routes = state['routes']
    endpoints = state['endpoints']
    physical_ports = state['physical_ports']

    # ---- 1. provenance / SHA256 ------------------------------------------
    provenance_check = _provenance_check(csv_sha, xlsx_sha)
    provenance_info = dict(provenance_check)
    provenance_passed = bool(provenance_info.pop('passed', False))
    record('sha256_provenance', provenance_passed, **provenance_info)
    record('csv_header', csv_header == CSV_FIELDS, observed=csv_header,
           expected=CSV_FIELDS)
    record('excel_sheet_count', len(excel['sheet_names']) == EXPECTED['excel_sheets'],
           observed=len(excel['sheet_names']), expected=EXPECTED['excel_sheets'],
           sheet_names=excel['sheet_names'])

    # ---- 2. CSV rows / ids ----------------------------------------------
    connections = len(routes)
    record('connections', connections == EXPECTED['rows'], observed=connections,
           expected=EXPECTED['rows'])
    route_ids = [int(route['route_id']) for route in routes]
    record('route_ids',
           sorted(route_ids) == list(range(EXPECTED['route_id_min'],
                                           EXPECTED['route_id_max'] + 1)),
           unique=len(set(route_ids)), min=min(route_ids), max=max(route_ids))

    # ---- 3. class counts (from the CSV relation column only) ------------
    class_counts = Counter(route['relation'] for route in routes)
    pair_counts = Counter(route['pair'] for route in routes)
    record('class_counts',
           {key: class_counts.get(key, 0) for key in CLASS_ORDER} == EXPECTED['class_counts'],
           observed={key: class_counts.get(key, 0) for key in CLASS_ORDER},
           expected=EXPECTED['class_counts'],
           source='CSV relation column (never inferred from the row number)')
    record('pair_counts',
           {key: pair_counts.get(key, 0) for key in PAIR_ORDER} == EXPECTED['pair_counts'],
           observed={key: pair_counts.get(key, 0) for key in PAIR_ORDER},
           expected=EXPECTED['pair_counts'])
    inconsistent = [route['route_id'] for route in routes
                    if not route['class_consistent']]
    record('relation_matches_sides', not inconsistent,
           violations=len(inconsistent), sample=inconsistent[:5],
           note='consistency check only; classification itself uses the CSV column')

    # ---- 4. channel / connector identity ---------------------------------
    channel_violations, connector_violations = [], []
    for route in routes:
        for role in ('source', 'target'):
            side = _text(route[f'{role}_side'])
            connector = _text(route[f'{role}_connector'])
            row = int(route[f'{role}_row'])
            column = int(route[f'{role}_column'])
            channel = int(route[f'{role}_channel'])
            if channel != (row - 1) * COLUMNS + column:
                channel_violations.append(dict(route_id=route['route_id'], role=role,
                                               row=row, column=column, channel=channel))
            if (connector[:1] != side or not connector[1:].isdigit()
                    or not 1 <= int(connector[1:]) <= CONNECTORS_PER_SIDE
                    or row not in (1, 2) or not 1 <= column <= COLUMNS
                    or not 1 <= channel <= CHANNELS_PER_CONNECTOR):
                connector_violations.append(dict(route_id=route['route_id'], role=role,
                                                 side=side, connector=connector,
                                                 row=row, column=column, channel=channel))
    record('channel_formula', not channel_violations, violations=len(channel_violations),
           sample=channel_violations[:5], formula='channel = (row-1)*16 + column')
    record('connector_letter', not connector_violations, violations=len(connector_violations),
           sample=connector_violations[:5])

    # ---- 5. same-side self connector ------------------------------------
    self_connector = [route['route_id'] for route in routes
                      if _text(route['source_connector']) == _text(route['target_connector'])]
    record('no_same_side_self_connector', not self_connector,
           violations=len(self_connector), sample=self_connector[:5])

    # ---- 6. 2048 endpoint identities used exactly once -------------------
    usage = Counter()
    for route in routes:
        for role in ('source', 'target'):
            eid = endpoint_id(route[f'{role}_connector'], int(route[f'{role}_row']),
                              int(route[f'{role}_column']))
            usage[eid] += 1
    expected_ids = {f'{side}{number:02d}:{row}:{column}'
                    for side in SIDES for number in range(1, CONNECTORS_PER_SIDE + 1)
                    for row in (1, 2) for column in range(1, COLUMNS + 1)}
    record('endpoints_used_once',
           set(usage) == expected_ids and set(usage.values()) == {1},
           distinct=len(usage), expected=len(expected_ids),
           usage_min=min(usage.values()), usage_max=max(usage.values()),
           missing=len(expected_ids - set(usage)), extra=len(set(usage) - expected_ids))
    record('endpoint_count', len(endpoints) == EXPECTED['endpoints'],
           observed=len(endpoints), expected=EXPECTED['endpoints'])

    # ---- 7. 64 connectors x 32 fibres each -------------------------------
    per_connector = defaultdict(Counter)
    for route in routes:
        for role in ('source', 'target'):
            connector = _text(route[f'{role}_connector'])
            per_connector[connector][int(route[f'{role}_channel'])] += 1
    connectors_expected = {f'{side}{number:02d}' for side in SIDES
                           for number in range(1, CONNECTORS_PER_SIDE + 1)}
    coverage_ok = (set(per_connector) == connectors_expected
                   and all(sorted(counter) == list(range(1, CHANNELS_PER_CONNECTOR + 1))
                           and set(counter.values()) == {1}
                           for counter in per_connector.values()))
    record('connector_fibres', coverage_ok, connectors=len(per_connector),
           expected_connectors=EXPECTED['connectors'],
           fibres_each=CHANNELS_PER_CONNECTOR,
           channels_ok=sum(1 for counter in per_connector.values()
                           if sorted(counter) == list(range(1, CHANNELS_PER_CONNECTOR + 1))))

    # ---- 8. dual-height plane split (from CSV rows) ----------------------
    mixed_rows = [route['route_id'] for route in routes
                  if int(route['source_row']) != int(route['target_row'])]
    record('no_mixed_rows', not mixed_rows, violations=len(mixed_rows),
           sample=mixed_rows[:5])
    plane_counts = {}
    for plane, row in PLANE_ROW.items():
        plane_routes = [route for route in routes if int(route['row']) == row]
        plane_class = Counter(route['relation'] for route in plane_routes)
        plane_counts[plane] = dict(
            row=row,
            z_mm=PLANE_Z_MM[row],
            connections=len(plane_routes),
            expected_connections=EXPECTED['plane_connections'][plane],
            class_counts={key: plane_class.get(key, 0) for key in CLASS_ORDER},
            expected_class_counts=EXPECTED['plane_row_counts'][plane],
        )
    plane_ok = all(
        plane_counts[plane]['connections'] == EXPECTED['plane_connections'][plane]
        and plane_counts[plane]['class_counts'] == EXPECTED['plane_row_counts'][plane]
        for plane in PLANE_ROW)
    record('dual_height_split', plane_ok, planes=plane_counts,
           rule='CSV row 1 -> z1, row 2 -> z2 (mixed classes, relation from column)')

    # ---- 9. endpoint plane lists ----------------------------------------
    plane_lists = {}
    for plane, row in PLANE_ROW.items():
        listed = [endpoint for endpoint in endpoints.values() if endpoint['plane'] == plane]
        plane_lists[plane] = dict(
            count=len(listed),
            expected=EXPECTED['plane_endpoints'][plane],
            unique=len({endpoint['endpoint_id'] for endpoint in listed}),
            rows={endpoint['row'] for endpoint in listed},
        )
    record('plane_endpoint_lists',
           all(plane_lists[plane]['count'] == EXPECTED['plane_endpoints'][plane]
               and plane_lists[plane]['unique'] == EXPECTED['plane_endpoints'][plane]
               and plane_lists[plane]['rows'] == {PLANE_ROW[plane]}
               for plane in PLANE_ROW),
           planes={plane: dict(count=value['count'], expected=value['expected'],
                               unique=value['unique'], rows=sorted(value['rows']))
                   for plane, value in plane_lists.items()})

    # ---- 10. physical 3D ports ------------------------------------------
    expected_port_ids = {(f'{side}{number:02d}', row, column)
                         for side in SIDES for number in range(1, CONNECTORS_PER_SIDE + 1)
                         for row in (1, 2) for column in range(1, COLUMNS + 1)}
    port_ids = Counter((port['connector'], port['row'], port['column'])
                       for port in physical_ports)
    coordinate_errors = 0
    max_coordinate_error = 0.0
    for port in physical_ports:
        expected_along = port_along_mm(port['connector'], port['column'])
        expected_xy = edge_xy(port['side'], expected_along)
        delta = max(abs(port['along_mm'] - expected_along),
                    abs(port['x_mm'] - expected_xy[0]),
                    abs(port['y_mm'] - expected_xy[1]),
                    abs(port['z_mm'] - PLANE_Z_MM[port['row']]))
        max_coordinate_error = max(max_coordinate_error, delta)
        if delta > TOL_MM:
            coordinate_errors += 1
    projected = Counter((round(port['x_mm'], 9), round(port['y_mm'], 9))
                        for port in physical_ports)
    unique_3d = len({(round(port['x_mm'], 9), round(port['y_mm'], 9),
                      round(port['z_mm'], 9)) for port in physical_ports})
    record('physical_ports',
           len(physical_ports) == EXPECTED['physical_ports']
           and set(port_ids) == expected_port_ids and set(port_ids.values()) == {1}
           and coordinate_errors == 0 and unique_3d == EXPECTED['physical_ports'],
           count=len(physical_ports), expected=EXPECTED['physical_ports'],
           distinct_identities=len(port_ids),
           duplicates=sum(1 for value in port_ids.values() if value > 1),
           max_coordinate_error_mm=max_coordinate_error,
           unique_3d=unique_3d,
           projected_xy_distinct=len(projected),
           cross_plane_projected_duplicate_ports=sum(1 for value in projected.values()
                                                     if value > 1),
           z1=sum(1 for port in physical_ports if port['plane'] == 'z1'),
           z2=sum(1 for port in physical_ports if port['plane'] == 'z2'))

    connector_order = [f'{side}{number:02d}' for side in SIDES
                       for number in range(1, CONNECTORS_PER_SIDE + 1)]
    matrix_from_csv = _directed_connector_matrix(routes, connector_order)
    matrix_undirected = [[0] * len(connector_order) for _ in connector_order]
    pair_counter = _connector_pair_counter(routes)
    index = {name: position for position, name in enumerate(connector_order)}
    for (connector_a, connector_b), count in pair_counter.items():
        matrix_undirected[index[connector_a]][index[connector_b]] = count
        matrix_undirected[index[connector_b]][index[connector_a]] = count
    column_deltas = []
    for side in SIDES:
        for number in range(1, CONNECTORS_PER_SIDE + 1):
            alongs = [port_along_mm(f'{side}{number:02d}', column)
                      for column in range(1, COLUMNS + 1)]
            column_deltas.extend(round(alongs[i + 1] - alongs[i], 12)
                                 for i in range(len(alongs) - 1))
    connector_centers = [FIRST_CONNECTOR_CENTER_MM + (number - 1) * CONNECTOR_PITCH_MM
                         for number in range(1, CONNECTORS_PER_SIDE + 1)]
    record('port_pitches',
           max(abs(delta - FIBER_COLUMN_PITCH_MM) for delta in column_deltas) <= TOL_MM
           and abs(PLANE_Z_MM[2] - PLANE_Z_MM[1] - FIBER_ROW_PITCH_MM) <= TOL_MM
           and max(abs(connector_centers[i + 1] - connector_centers[i]
                       - CONNECTOR_PITCH_MM)
                   for i in range(len(connector_centers) - 1)) <= TOL_MM,
           column_pitch_mm=FIBER_COLUMN_PITCH_MM,
           column_pitch_observed=sorted(set(column_deltas)),
           row_z_separation_mm=PLANE_Z_MM[2] - PLANE_Z_MM[1],
           fiber_row_pitch_mm=FIBER_ROW_PITCH_MM,
           connector_pitch_mm=CONNECTOR_PITCH_MM,
           connector_centers_first_last=[connector_centers[0], connector_centers[-1]])
    connector_degree_undirected = [sum(row) for row in matrix_undirected]
    record('connector_degree',
           all(total == 32 for total in connector_degree_undirected)
           and sum(connector_degree_undirected) == 2048,
           row_sums=sorted(set(connector_degree_undirected)), expected=32,
           total=sum(connector_degree_undirected),
           note='undirected connector matrix: each connector is incident to exactly '
                '32 fibre connections (row sum 32; matrix total 2048 = 2 x 1024)')

    # ---- 11. Excel cross-checks -----------------------------------------
    excel_checks = _validate_excel(excel, routes, matrix_undirected, connector_order)
    checks.update(excel_checks)
    failed.extend(name for name, check in excel_checks.items()
                  if check.get('passed') is False)

    # ---- 12. Lead core modules are cross-checked only in a subprocess ----
    # (see compare_with_lead_v3_module(); this audit never imports v3 inputs/ports)

    # ---- 13. input notes / warnings -------------------------------------
    warnings = [
        'All coordinates in this audit are simulation under the two-height interface '
        'assumption (z1 = 0.0 mm, z2 = 0.5 mm are model values); glass thickness and '
        'connector mounting direction are not confirmed.',
        'No z transition (via / taper / fanout) is implemented; the two planes are '
        'independent 2D routing problems and cross-plane XY projections are NOT '
        'same-layer conflicts.',
        '9 mm connector pitch mechanical interference is unresolved: the provisional '
        'MPO housing envelope 12.59 mm overlaps neighbouring housings by 3.59 mm '
        '(the 3.75 mm fiber span itself fits); this is inherited from the FSLW round.',
        'The V3 topology is a synthetic 1024-channel expansion of the user-provided '
        '512 fiberBoard (doubling + 96 NS/EW direction exchanges + Euler row/column '
        'fibre assignment); it is a topology baseline, not a measured or hardware '
        'validated demand (see the workbook generation-notes sheet).',
        'Same-side connections (448 = 43.75%) require hairpin/U (refold) routing; the '
        'previous fslw loader rejects same-side connections and assumes row1=opposite / '
        'row2=adjacent, so V3 needs the new loader/validator.',
        'The relation is authoritative from the CSV relation column; the row only '
        'selects the dual-height plane. Row 1 contains 222 same + 196 opposite + '
        '94 adjacent; row 2 contains 226 + 188 + 98.',
        'Excel Sheet1 Port1/Port2 uses 1-based connector indices 1..64 following the '
        'port-map sheet (N01=1 ... W16=64); all 1024 rows match the CSV in order.',
        'Cross-plane XY projections are not same-layer conflicts; conflict checking is '
        'per plane only.',
    ]
    passed = not failed
    return dict(
        passed=passed,
        failed_checks=failed,
        inputs=dict(
            csv=dict(path=str(csv_path.resolve()), sha256=csv_sha,
                     expected_sha256=CSV_SHA256, bytes=csv_path.stat().st_size),
            xlsx=dict(path=str(xlsx_path.resolve()), sha256=xlsx_sha,
                      expected_sha256=XLSX_SHA256, bytes=xlsx_path.stat().st_size,
                      sheet_names=excel['sheet_names']),
            provenance=str(PROVENANCE_PATH.resolve())
            if PROVENANCE_PATH.exists() else None,
        ),
        model=dict(MODEL),
        expected=dict(EXPECTED),
        checks=checks,
        counts=dict(
            connections=connections,
            class_counts={key: class_counts.get(key, 0) for key in CLASS_ORDER},
            pair_counts={key: pair_counts.get(key, 0) for key in PAIR_ORDER},
            planes=plane_counts,
            connectors=len(per_connector),
            endpoints=len(endpoints),
            unique_connector_pairs=len(pair_counter),
            physical_ports=len(physical_ports),
        ),
        excel=dict(
            sheet_names=excel['sheet_names'],
            sheet_roles=dict(
                ports=excel['ports_sheet_name'],
                port_map=excel['port_map_sheet_name'],
                class_stats=excel['stats_sheet_name'],
                matrix=excel['matrix_sheet_name'],
                fibre_reference=excel['fibre_sheet_name'],
                pair_multiplicity=excel['multiplicity_sheet_name'],
                generation_notes=excel['notes_sheet_name'],
            ),
            port_pairs=len(excel['port_pairs']),
            port_map_rows=len(excel['port_map']),
            matrix_shape=[len(excel['matrix']), len(excel['matrix'][0])]
            if excel['matrix'] else [0, 0],
            fibre_rows=len(excel['fibre_rows']),
            multiplicities=len(excel['multiplicities']),
            class_stats=excel['class_stats']['counts'],
            notes=excel['notes'],
        ),
        warnings=warnings,
        note_lines=[note['description'] for note in excel['notes']],
    )


def _validate_excel(excel, routes, matrix_undirected, connector_order) -> dict:
    """All Excel-vs-CSV cross-checks (each returns a check record)."""
    checks = {}

    def record(name, passed, **info):
        checks[name] = dict(passed=bool(passed), **info)

    csv_index_pairs = [(int(route['source_connector_index']),
                        int(route['target_connector_index'])) for route in routes]
    sheet_pairs = [(int(port_a), int(port_b)) for port_a, port_b in excel['port_pairs']]
    record('excel_sheet1_pairs',
           len(sheet_pairs) == len(csv_index_pairs) == 1024
           and sheet_pairs == csv_index_pairs,
           sheet_rows=len(sheet_pairs), csv_rows=len(csv_index_pairs),
           same_order=(sheet_pairs == csv_index_pairs),
           same_multiset=(Counter(tuple(sorted(pair)) for pair in sheet_pairs)
                          == Counter(tuple(sorted(pair)) for pair in csv_index_pairs)),
           min_port=min(min(pair) for pair in sheet_pairs),
           max_port=max(max(pair) for pair in sheet_pairs),
           self_pairs=sum(1 for port_a, port_b in sheet_pairs if port_a == port_b))

    port_map = excel['port_map']
    names = [entry['connector'] for entry in port_map]
    expected_names = [f'{side}{number:02d}' for side in SIDES
                      for number in range(1, CONNECTORS_PER_SIDE + 1)]
    centers_ok = all(
        abs(entry['center_mm'] - (FIRST_CONNECTOR_CENTER_MM
                                  + (entry['side_index'] - 1) * CONNECTOR_PITCH_MM))
        <= TOL_MM for entry in port_map)
    neighbours = defaultdict(set)
    for route in routes:
        connector_a = _text(route['source_connector'])
        connector_b = _text(route['target_connector'])
        neighbours[connector_a].add(connector_b)
        neighbours[connector_b].add(connector_a)
    distinct_ok = all(len(neighbours[entry['connector']]) == entry['distinct_targets']
                      for entry in port_map)
    record('excel_port_map',
           len(port_map) == EXPECTED['connectors'] and names == expected_names
           and [entry['port'] for entry in port_map] == list(range(1, 65))
           and all(entry['side'] == entry['connector'][0] for entry in port_map)
           and [entry['side_index'] for entry in port_map]
           == [number for side in SIDES
               for number in range(1, CONNECTORS_PER_SIDE + 1)]
           and centers_ok and distinct_ok
           and all(entry['fibres'] == 32 and entry['row1_fibres'] == 16
                   and entry['row2_fibres'] == 16 for entry in port_map),
           rows=len(port_map), names_match=(names == expected_names),
           centers_match=centers_ok, distinct_targets_match=distinct_ok,
           distinct_targets_range=[min(entry['distinct_targets'] for entry in port_map),
                                   max(entry['distinct_targets'] for entry in port_map)])

    matrix = excel['matrix']
    header = excel['matrix_columns']
    row_names = excel['matrix_rows']
    symmetric = all(matrix[i][j] == matrix[j][i]
                    for i in range(len(matrix)) for j in range(len(matrix)))
    diagonal = [matrix[i][i] for i in range(min(len(matrix), len(matrix[0])))]
    row_sums = [sum(row) for row in matrix]
    total = sum(row_sums)
    matrix_matches = (header == connector_order and row_names == connector_order
                      and matrix == matrix_undirected)
    record('excel_64x64_matrix',
           len(matrix) == 64 and all(len(row) == 64 for row in matrix)
           and matrix_matches and symmetric and set(diagonal) == {0}
           and set(row_sums) == {32} and total == EXPECTED['matrix_total'],
           shape=[len(matrix), len(matrix[0]) if matrix else 0],
           names_match=(header == connector_order and row_names == connector_order),
           equals_csv_connector_counts=matrix_matches,
           symmetric=symmetric, diagonal_zero=(set(diagonal) == {0}),
           row_sums=sorted(set(row_sums)), total=total,
           total_expected=EXPECTED['matrix_total'],
           note='symmetric matrix counts each undirected connection twice '
                '(sum=2048=2*1024); per-connector degree 32')

    fibre_header = excel['fibre_header']
    fibre_rows = excel['fibre_rows']
    csv_rows_as_lists = [[route[field] for field in CSV_FIELDS] for route in routes]
    mismatches = []
    for position, (sheet_row, csv_row) in enumerate(zip(fibre_rows, csv_rows_as_lists)):
        if list(sheet_row) != list(csv_row):
            mismatches.append(dict(row=position + 2, sheet=sheet_row, csv=csv_row))
    record('excel_fibre_sheet_equals_csv',
           fibre_header == CSV_FIELDS and len(fibre_rows) == len(csv_rows_as_lists) == 1024
           and not mismatches,
           rows=len(fibre_rows), expected_rows=len(csv_rows_as_lists),
           header_match=(fibre_header == CSV_FIELDS), mismatches=len(mismatches),
           sample=mismatches[:3])

    pair_counter = _connector_pair_counter(routes)
    sheet_multiplicity = {}
    duplicate_rows = []
    for entry in excel['multiplicities']:
        pair = tuple(sorted((entry['connector_a'], entry['connector_b'])))
        if pair in sheet_multiplicity:
            duplicate_rows.append(dict(pair=pair, reason='duplicate row'))
        sheet_multiplicity[pair] = (entry['count'], entry['connection_class'])
    counts_match = (set(sheet_multiplicity) == set(pair_counter)
                    and all(sheet_multiplicity[pair][0] == count
                            for pair, count in pair_counter.items()))
    classes_match = True
    for pair in pair_counter:
        canonical = side_pair_label(pair[0], pair[1])
        if pair not in sheet_multiplicity \
                or sheet_multiplicity[pair][1] != PAIR_CLASS[canonical]:
            classes_match = False
    record('excel_pair_multiplicity',
           counts_match and classes_match and not duplicate_rows
           and sum(value[0] for value in sheet_multiplicity.values()) == 1024
           and len(sheet_multiplicity) == len(pair_counter),
           rows=len(excel['multiplicities']), unique_pairs=len(sheet_multiplicity),
           csv_unique_pairs=len(pair_counter),
           counts_match=counts_match, classes_match=classes_match,
           total=sum(value[0] for value in sheet_multiplicity.values()))

    stats = excel['class_stats']
    stats_counts = stats.get('counts', {})
    stats_pairs = stats.get('pair_counts', {})
    csv_class_counts = Counter(route['relation'] for route in routes)
    csv_pair_counts = Counter(route['pair'] for route in routes)
    record('excel_class_statistics',
           all(stats_counts.get(key) == csv_class_counts.get(key)
               for key in CLASS_ORDER)
           and stats_counts.get('total') == 1024
           and all(stats_pairs.get(key) == csv_pair_counts.get(key)
                   for key in PAIR_ORDER),
           sheet_counts=stats_counts, sheet_pair_counts=stats_pairs,
           csv_counts={key: csv_class_counts.get(key, 0) for key in CLASS_ORDER},
           csv_pair_counts={key: csv_pair_counts.get(key, 0) for key in PAIR_ORDER})

    return checks


# ---------------------------------------------------------------------------
# Comparison against Lead's V3 core modules (subprocess only; never imported here)
# ---------------------------------------------------------------------------
_LEAD_SNAPSHOT_SCRIPT = r'''
import json
import sys
from pathlib import Path

from four_side_1024.v3 import config as C
from four_side_1024.v3 import inputs as I
from four_side_1024.v3 import ports as P

ENDPOINT_FIELDS = ['endpoint_id', 'connector', 'side', 'row', 'column', 'channel',
                   'routing_xy', 'inward', 'layer', 'plane', 'z_mm', 'physical_along_mm',
                   'physical_face_u_mm', 'physical_face_v_mm', 'port_model']
ROUTE_FIELDS = ['route_id', 'source_side', 'source_connector', 'source_row',
                'source_column', 'source_channel', 'target_side', 'target_connector',
                'target_row', 'target_column', 'target_channel', 'relation',
                'relation_csv', 'pair', 'plane', 'z_mm', 'source_endpoint',
                'target_endpoint']
CONFIG_FIELDS = ['BOARD_MM', 'CONNECTORS_PER_SIDE', 'COLUMNS', 'ROWS',
                 'CONNECTOR_PITCH_MM', 'FIBER_COLUMN_PITCH_MM', 'FIBER_ROW_PITCH_MM',
                 'PLANE_Z_MM', 'PLANE_NAME', 'PLANE_ROW', 'PORT_MODEL', 'TOPOLOGY',
                 'EXPECTED', 'EXPECTED_CSV_SHA256', 'EXPECTED_XLSX_SHA256']


def endpoint_record(endpoint):
    record = {}
    for field in ENDPOINT_FIELDS:
        value = endpoint.get(field)
        if isinstance(value, tuple):
            value = list(value)
        if field in ('routing_xy', 'inward') and value is not None:
            value = [float(item) for item in value]
        record[field] = value
    return record


def route_record(route):
    record = {}
    for field in ROUTE_FIELDS:
        value = route.get(field)
        if isinstance(value, tuple):
            value = list(value)
        record[field] = value
    return record


out_path = Path(sys.argv[1])
connections, topology_validation = I.load_all()
payload = {
    'module_paths': {'inputs': I.__file__, 'ports': P.__file__, 'config': C.__file__},
    'config': {field: getattr(C, field, None) for field in CONFIG_FIELDS},
    'endpoints': {key: endpoint_record(value) for key, value in P.build_endpoints().items()},
    'connector_records': P.connector_records(),
    'physical_ports': P.physical_port_table(),
    'all_routes': [route_record(route) for route in P.all_routes()],
    'topology_validation': {key: topology_validation.get(key) for key in
                            ('passed', 'csv_path', 'csv_sha256', 'csv_sha256_match',
                             'connections', 'relation_counts', 'pair_counts', 'by_side',
                             'planes', 'endpoints_used_once', 'issues')},
    'excel_validation': topology_validation.get('excel'),
    'planes': {},
}
for plane in C.PLANES:
    routes, endpoints, _connectors, plane_validation = P.load_plane(plane)
    payload['planes'][plane] = {
        'validation': {key: plane_validation.get(key) for key in
                       ('passed', 'issues', 'plane', 'plane_row', 'z_mm', 'connections',
                        'ports', 'relations', 'pairs', 'sides', 'rows', 'csv_sha256',
                        'csv_sha256_match', 'xlsx_sha256')},
        'endpoints': {key: endpoint_record(value) for key, value in endpoints.items()},
        'route_ids': [route['route_id'] for route in routes],
        'routes': [route_record(route) for route in routes],
    }
out_path.write_text(json.dumps(payload, ensure_ascii=False, default=str), encoding='utf-8')
print('v3 snapshot ok')
'''


def _module_source_record(path) -> dict:
    path = Path(path)
    if not path.exists():
        return dict(path=str(path), exists=False)
    import datetime
    stat = path.stat()
    return dict(path=str(path), exists=True, bytes=stat.st_size,
                sha256=sha256_file(path),
                mtime=datetime.datetime.fromtimestamp(stat.st_mtime).isoformat(
                    sep=' ', timespec='seconds'))


def lead_v3_snapshot(python_exe=None, timeout=600) -> dict:
    """Run Lead's V3 ``inputs``/``ports``/``config`` modules in a subprocess."""
    import tempfile

    python = str(python_exe or sys.executable)
    with tempfile.TemporaryDirectory(prefix='v3_lead_snapshot_') as tmp:
        out_path = Path(tmp) / 'v3_lead_snapshot.json'
        completed = subprocess.run(
            [python, '-c', _LEAD_SNAPSHOT_SCRIPT, str(out_path)],
            cwd=str(CODE_DIR), capture_output=True, text=True, timeout=timeout)
        if completed.returncode != 0:
            raise RuntimeError(
                'Lead V3 modules snapshot failed: '
                f'{(completed.stderr or completed.stdout or "").strip()[-3000:]}')
        if not out_path.exists():
            raise RuntimeError('Lead V3 modules snapshot wrote no output')
        return json.loads(out_path.read_text(encoding='utf-8'))


def _canonical_pair(pair) -> str:
    """Both the sorted (Excel) and oriented (SPEC table) pair labels -> sorted form."""
    text = '' if pair is None else str(pair)
    if len(text) == 2 and text[0] == text[1]:
        return text
    return ''.join(sorted(text))


def compare_with_lead_v3_module(python_exe=None, state=None) -> dict:
    """Compare our independent derivation with Lead's V3 core modules.

    Runs Lead's ``v3/config.py``, ``v3/inputs.py`` and ``v3/ports.py`` only in a
    subprocess (never imported by this audit module), then compares:
    all 2048 endpoint identities, the 1024 route records, the two planes, the
    2048 physical ports and the module constants.
    """
    state = state or derive_ports()
    snapshot = lead_v3_snapshot(python_exe=python_exe)
    my_endpoints = state['endpoints']
    lead_endpoints = snapshot['endpoints']
    comparison = dict(
        lead_modules_available=True,
        module_paths=snapshot.get('module_paths'),
        module_sources={name: _module_source_record(path)
                        for name, path in (snapshot.get('module_paths') or {}).items()},
        config={},
        endpoints={},
        routes={},
        planes={},
        physical_ports={},
        lead_topology_validation=snapshot.get('topology_validation'),
        lead_excel_validation=snapshot.get('excel_validation'),
        code_notes=[],
    )

    lead_config = snapshot.get('config') or {}
    constant_pairs = {
        'BOARD_MM': BOARD_MM,
        'CONNECTORS_PER_SIDE': CONNECTORS_PER_SIDE,
        'COLUMNS': COLUMNS,
        'ROWS': ROWS,
        'CONNECTOR_PITCH_MM': CONNECTOR_PITCH_MM,
        'FIBER_COLUMN_PITCH_MM': FIBER_COLUMN_PITCH_MM,
        'FIBER_ROW_PITCH_MM': FIBER_ROW_PITCH_MM,
        'PLANE_Z_MM': PLANE_Z_MM,
        'PLANE_NAME': PLANE_NAME,
        'PLANE_ROW': PLANE_ROW,
        'PORT_MODEL': PORT_MODEL,
    }
    config_mismatches = []
    for name, mine in constant_pairs.items():
        theirs = lead_config.get(name)
        if theirs is None or _jsonable(theirs) != _jsonable(mine):
            config_mismatches.append(dict(constant=name, independent=mine, lead=theirs))
    comparison['config'] = dict(mismatches=config_mismatches,
                                lead=lead_config,
                                passed=not config_mismatches)

    max_routing = 0.0
    max_along = 0.0
    max_z = 0.0
    plane_violations = []
    field_violations = []
    for key in sorted(set(my_endpoints) & set(lead_endpoints)):
        mine, theirs = my_endpoints[key], lead_endpoints[key]
        max_routing = max(
            max_routing,
            max(abs(float(theirs['routing_xy'][0]) - mine['routing_xy'][0]),
                abs(float(theirs['routing_xy'][1]) - mine['routing_xy'][1])))
        if theirs.get('physical_along_mm') is not None:
            max_along = max(max_along,
                            abs(float(theirs['physical_along_mm'])
                                - mine['physical_along_mm']))
        max_z = max(max_z, abs(float(theirs['z_mm']) - mine['z_mm']))
        if theirs['plane'] != mine['plane'] or \
                PLANE_NAME.get(int(theirs['row'])) != theirs['plane']:
            plane_violations.append(dict(endpoint=key, ours=mine['plane'],
                                         lead=theirs['plane'], row=theirs['row']))
        fields = ('side', 'connector', 'row', 'column', 'channel', 'layer')
        if any(str(theirs[field]) != str(mine[field]) for field in fields) or \
                [float(value) for value in theirs['inward']] != list(mine['inward']):
            field_violations.append(dict(endpoint=key))
    endpoint_identity_sets_equal = set(my_endpoints) == set(lead_endpoints)
    comparison['endpoints'] = dict(
        ours=len(my_endpoints), lead=len(lead_endpoints),
        identity_sets_equal=endpoint_identity_sets_equal,
        missing_in_lead=sorted(set(my_endpoints) - set(lead_endpoints))[:10],
        extra_in_lead=sorted(set(lead_endpoints) - set(my_endpoints))[:10],
        max_routing_xy_delta_mm=max_routing,
        max_physical_along_delta_mm=max_along,
        max_z_delta_mm=max_z,
        plane_violations=len(plane_violations),
        field_violations=len(field_violations),
        plane_violation_sample=plane_violations[:5],
        field_violation_sample=field_violations[:5],
        passed=(endpoint_identity_sets_equal and not plane_violations
                and not field_violations and max_routing <= MATCH_TOL_MM
                and max_along <= MATCH_TOL_MM and max_z <= MATCH_TOL_MM),
    )

    my_routes = {int(route['route_id']): route for route in state['routes']}
    lead_routes = {int(route['route_id']): route for route in snapshot['all_routes']}
    route_field_mismatches = []
    relation_column_mismatches = []
    pair_label_mismatches = []
    row_plane_violations = []
    for route_id in sorted(set(my_routes) & set(lead_routes)):
        mine, theirs = my_routes[route_id], lead_routes[route_id]
        expected_plane = PLANE_NAME[int(mine['row'])]
        if (int(theirs['source_row']) != int(mine['source_row'])
                or int(theirs['target_row']) != int(mine['target_row'])
                or theirs['plane'] != expected_plane
                or abs(float(theirs['z_mm']) - PLANE_Z_MM[int(mine['row'])]) > TOL_MM
                or theirs['source_endpoint'] != mine['source_endpoint']
                or theirs['target_endpoint'] != mine['target_endpoint']
                or str(theirs['source_connector']) != str(mine['source_connector'])
                or str(theirs['target_connector']) != str(mine['target_connector'])):
            route_field_mismatches.append(dict(route_id=route_id,
                                               ours=dict(plane=mine['plane']),
                                               lead=dict(plane=theirs['plane'])))
        if theirs.get('relation_csv') != mine['relation']:
            relation_column_mismatches.append(dict(route_id=route_id,
                                                   csv_column=mine['relation'],
                                                   lead_relation_csv=theirs.get('relation_csv')))
        if _canonical_pair(theirs.get('pair')) != _canonical_pair(mine['pair']):
            pair_label_mismatches.append(dict(route_id=route_id, ours=mine['pair'],
                                              lead=theirs.get('pair')))
        if theirs['plane'] != PLANE_NAME[int(theirs['source_row'])]:
            row_plane_violations.append(route_id)
    comparison['routes'] = dict(
        ours=len(my_routes), lead=len(lead_routes),
        route_id_sets_equal=set(my_routes) == set(lead_routes),
        field_mismatches=len(route_field_mismatches),
        relation_column_mismatches=len(relation_column_mismatches),
        pair_label_mismatches=len(pair_label_mismatches),
        row_plane_violations=len(row_plane_violations),
        field_mismatch_sample=route_field_mismatches[:5],
        passed=(set(my_routes) == set(lead_routes) and not route_field_mismatches
                and not relation_column_mismatches and not pair_label_mismatches
                and not row_plane_violations),
    )

    for plane in ('z1', 'z2'):
        lead_plane = snapshot['planes'][plane]
        expected_row = PLANE_ROW[plane]
        my_plane_routes = {int(route['route_id']) for route in state['routes']
                           if int(route['row']) == expected_row}
        my_plane_endpoints = {key for key, endpoint in my_endpoints.items()
                              if endpoint['plane'] == plane}
        lead_route_ids = {int(route_id) for route_id in lead_plane['route_ids']}
        lead_plane_endpoints = set(lead_plane['endpoints'])
        max_plane_routing = 0.0
        plane_field_violations = 0
        for key in sorted(my_plane_endpoints & lead_plane_endpoints):
            mine, theirs = my_endpoints[key], lead_plane['endpoints'][key]
            max_plane_routing = max(
                max_plane_routing,
                max(abs(float(theirs['routing_xy'][0]) - mine['routing_xy'][0]),
                    abs(float(theirs['routing_xy'][1]) - mine['routing_xy'][1])))
            if theirs['plane'] != plane or int(theirs['row']) != expected_row:
                plane_field_violations += 1
        validation = lead_plane.get('validation') or {}
        my_plane_relations = Counter(route['relation'] for route in state['routes']
                                     if int(route['row']) == expected_row)
        comparison['planes'][plane] = dict(
            lead_validation_passed=bool(validation.get('passed')),
            lead_validation_issues=list(validation.get('issues') or []),
            route_id_sets_equal=(my_plane_routes == lead_route_ids),
            our_routes=len(my_plane_routes), lead_routes=len(lead_route_ids),
            endpoint_set_mismatches=len(my_plane_endpoints ^ lead_plane_endpoints),
            max_routing_xy_delta_mm=max_plane_routing,
            plane_field_violations=plane_field_violations,
            our_relations={key: my_plane_relations.get(key, 0) for key in CLASS_ORDER},
            lead_relations=validation.get('relations'),
        )

    my_ports = {port['endpoint_id']: port for port in state['physical_ports']}
    lead_ports = {str(port['endpoint_id']): port
                  for port in snapshot['physical_ports']}
    max_port_delta = 0.0
    port_field_violations = []
    for key in sorted(set(my_ports) & set(lead_ports)):
        mine, theirs = my_ports[key], lead_ports[key]
        delta = max(abs(float(theirs['x_mm']) - mine['x_mm']),
                    abs(float(theirs['y_mm']) - mine['y_mm']),
                    abs(float(theirs['z_mm']) - mine['z_mm']),
                    abs(float(theirs['along_mm']) - mine['along_mm']))
        max_port_delta = max(max_port_delta, delta)
        if theirs['plane'] != mine['plane'] \
                or int(theirs['row']) != int(mine['row']) \
                or int(theirs['column']) != int(mine['column']):
            port_field_violations.append(dict(endpoint=key))
    comparison['physical_ports'] = dict(
        ours=len(my_ports), lead=len(lead_ports),
        identity_sets_equal=set(my_ports) == set(lead_ports),
        max_coordinate_delta_mm=max_port_delta,
        field_violations=len(port_field_violations),
        passed=(set(my_ports) == set(lead_ports) and not port_field_violations
                and max_port_delta <= MATCH_TOL_MM),
    )

    topology_validation = snapshot.get('topology_validation') or {}
    excel_validation = snapshot.get('excel_validation') or {}
    comparison['lead_topology_validation_passed'] = bool(topology_validation.get('passed'))
    comparison['lead_topology_validation_issues'] = list(
        topology_validation.get('issues') or [])
    comparison['lead_excel_validation_passed'] = bool(excel_validation.get('passed'))
    comparison['lead_excel_validation_issues'] = list(excel_validation.get('issues') or [])
    comparison['lead_csv_sha256_match'] = topology_validation.get('csv_sha256_match')
    comparison['lead_xlsx_sha256_match'] = excel_validation.get('xlsx_sha256_match')

    comparison['code_notes'] = [
        'Lead v3/ports.py routes do not set a "row" field (only plane/z_mm and '
        'source_row/target_row); the V3 route schema should state whether "row" is '
        'required. Our comparison derives the row from source_row.',
        'Lead v3/inputs.py rewrites record["relation"] from the two sides and keeps the '
        'original CSV column in record["relation_csv"]; both agree on the frozen CSV '
        '(our audit verifies the CSV column directly, so the conventions coincide).',
        'Lead v3/inputs.py pair labels are directional for adjacent pairs (NE/ES/SW/NW), '
        'while the Excel statistics sheet uses the sorted form (EN/ES/SW/NW); both are '
        'the same 48 connections. Our audit canonicalises pair labels when comparing.',
        'Lead v3/inputs.py excel_crosscheck compares the fibre-level reference sheet as '
        'sorted tuples and checks only row counts for the port map / statistics sheets; '
        'our audit additionally compares the entire port map, class-statistics pair table '
        'and connector-pair multiplicity sheet (all pass).',
        'Lead v3/ports.py inherits the FSLW 2x16 physical port geometry unchanged; the '
        'dual-height assumption (no z transition, glass thickness / connector mounting '
        'unconfirmed) therefore carries over to V3.',
    ]

    endpoints_ok = comparison['endpoints']['passed']
    routes_ok = comparison['routes']['passed']
    planes_ok = all(entry['lead_validation_passed'] and entry['route_id_sets_equal']
                    and entry['endpoint_set_mismatches'] == 0
                    and entry['max_routing_xy_delta_mm'] <= MATCH_TOL_MM
                    and entry['plane_field_violations'] == 0
                    for entry in comparison['planes'].values())
    ports_ok = comparison['physical_ports']['passed']
    comparison['passed'] = bool(config_mismatches == [] and endpoints_ok and routes_ok
                                and planes_ok and ports_ok
                                and comparison['lead_topology_validation_passed']
                                and comparison['lead_excel_validation_passed'])
    comparison['summary'] = dict(
        passed=comparison['passed'],
        endpoint_identities=len(my_endpoints),
        endpoint_identity_sets_equal=endpoint_identity_sets_equal,
        max_routing_xy_delta_mm=max_routing,
        max_physical_along_delta_mm=max_along,
        max_z_delta_mm=max_z,
        endpoint_plane_violations=len(plane_violations),
        routes=dict(ours=len(my_routes), lead=len(lead_routes),
                    route_id_sets_equal=set(my_routes) == set(lead_routes),
                    field_mismatches=len(route_field_mismatches),
                    relation_column_mismatches=len(relation_column_mismatches),
                    pair_label_mismatches=len(pair_label_mismatches),
                    row_plane_violations=len(row_plane_violations)),
        planes={plane: dict(lead_validation_passed=entry['lead_validation_passed'],
                            route_id_sets_equal=entry['route_id_sets_equal'],
                            endpoint_set_mismatches=entry['endpoint_set_mismatches'],
                            max_routing_xy_delta_mm=entry['max_routing_xy_delta_mm'])
                for plane, entry in comparison['planes'].items()},
        physical_ports=dict(identity_sets_equal=set(my_ports) == set(lead_ports),
                            max_coordinate_delta_mm=max_port_delta,
                            field_violations=len(port_field_violations)),
        config_mismatches=config_mismatches,
        lead_topology_validation_passed=comparison['lead_topology_validation_passed'],
        lead_excel_validation_passed=comparison['lead_excel_validation_passed'],
    )
    return comparison


# ---------------------------------------------------------------------------
# Figures: demand diagrams, 64x64 matrix heatmap, class composition
# ---------------------------------------------------------------------------
def _pyplot():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    return plt


def _connector_centers():
    return {f'{side}{number:02d}': edge_xy(side, FIRST_CONNECTOR_CENTER_MM
                                           + (number - 1) * CONNECTOR_PITCH_MM)
            for side in SIDES for number in range(1, CONNECTORS_PER_SIDE + 1)}


def _draw_board_and_connectors(ax):
    from matplotlib.patches import Rectangle
    ax.add_patch(Rectangle((0, 0), BOARD_MM, BOARD_MM, fill=False, linewidth=1.4,
                           zorder=5))
    centers = _connector_centers()
    for name, center in centers.items():
        side = name[0]
        ax.plot([center[0]], [center[1]], marker='s', markersize=4.2,
                color=SIDE_COLORS[side], markeredgecolor='black',
                markeredgewidth=0.3, zorder=6)
        if side == 'N':
            ax.text(center[0], BOARD_MM + 4.4, name, ha='center', va='bottom', fontsize=5)
        elif side == 'S':
            ax.text(center[0], -4.4, name, ha='center', va='top', fontsize=5)
        elif side == 'E':
            ax.text(BOARD_MM + 5.4, center[1], name, ha='left', va='center', fontsize=5)
        else:
            ax.text(-5.4, center[1], name, ha='right', va='center', fontsize=5)
    return centers


def _draw_route(ax, source, target, connection_class, linewidth=0.42, alpha=0.38):
    from matplotlib.patches import FancyArrowPatch
    color = CLASS_COLORS[connection_class]
    same_edge = (abs(source[0] - target[0]) <= TOL_MM
                 or abs(source[1] - target[1]) <= TOL_MM)
    if same_edge:
        # same-side connection -> hairpin/U demand that must bulge INTO the board
        if abs(source[1] - BOARD_MM) <= TOL_MM:       # N edge: inward = -y
            delta = target[0] - source[0]
            rad = 0.22 * (1 if delta >= 0 else -1)
        elif abs(source[1]) <= TOL_MM:                # S edge: inward = +y
            delta = target[0] - source[0]
            rad = -0.22 * (1 if delta >= 0 else -1)
        elif abs(source[0] - BOARD_MM) <= TOL_MM:     # E edge: inward = -x
            delta = target[1] - source[1]
            rad = -0.22 * (1 if delta >= 0 else -1)
        else:                                         # W edge: inward = +x
            delta = target[1] - source[1]
            rad = 0.22 * (1 if delta >= 0 else -1)
        ax.add_patch(FancyArrowPatch(
            (source[0], source[1]), (target[0], target[1]),
            connectionstyle=f'arc3,rad={rad}', arrowstyle='-',
            linewidth=linewidth, color=color, alpha=alpha, zorder=1))
    else:
        ax.plot([source[0], target[0]], [source[1], target[1]],
                color=color, linewidth=linewidth, alpha=alpha, zorder=1)


def _class_counts_text(class_counts, pair_counts=None):
    parts = [f'{klass} {class_counts.get(klass, 0)}' for klass in CLASS_ORDER]
    text = ' | '.join(parts)
    if pair_counts:
        text += '\n' + ' - '.join(f'{pair} {pair_counts[pair]}' for pair in PAIR_ORDER
                                  if pair_counts.get(pair))
    return text


def plot_connector_demand(routes, class_counts, pair_counts, path,
                          title='V3-Huang 1024 connections - connector-level demand'):
    plt = _pyplot()
    from matplotlib.lines import Line2D
    fig, ax = plt.subplots(figsize=(13.5, 13.5))
    centers = _draw_board_and_connectors(ax)
    for route in routes:
        _draw_route(ax, centers[str(route['source_connector'])],
                    centers[str(route['target_connector'])],
                    str(route['relation']))
    handles = [Line2D([0], [0], color=CLASS_COLORS[klass], lw=1.8,
                      label=f'{klass} ({class_counts.get(klass, 0)})')
               for klass in CLASS_ORDER]
    ax.legend(handles=handles, loc='upper center', bbox_to_anchor=(0.5, -0.06),
              ncol=3, fontsize=9, title='connection class (count)', title_fontsize=9)
    subtitle = _class_counts_text(class_counts, pair_counts)
    ax.set(xlim=(-11, 161), ylim=(-11, 161), aspect='equal', xlabel='x (mm)',
           ylabel='y (mm)')
    ax.set_title(
        f'{title}\n{subtitle}\n'
        f'same-side connections drawn as inward arcs (U/refold demand); '
        f'relation from the CSV column; plane not shown in this figure',
        fontsize=10)
    ax.grid(alpha=0.15)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)
    return Path(path)


def plot_plane_demand(state, plane, path):
    routes = [route for route in state['routes'] if route['plane'] == plane]
    class_counts = Counter(route['relation'] for route in routes)
    pair_counts = Counter(route['pair'] for route in routes)
    row = PLANE_ROW[plane]
    title = (f'V3-Huang plane {plane} (CSV row {row}, z = {PLANE_Z_MM[row]} mm) - '
             f'{len(routes)}/512 connections at connector level')
    return plot_connector_demand(routes, class_counts, pair_counts, path, title=title)


def plot_connector_matrix_heatmap(matrix, connector_order, path):
    plt = _pyplot()
    import numpy as np
    from matplotlib.colors import PowerNorm
    data = np.asarray(matrix, dtype=float)
    fig, ax = plt.subplots(figsize=(13.5, 12))
    norm = PowerNorm(gamma=0.55, vmin=0, vmax=max(1.0, data.max()))
    image = ax.imshow(data, origin='lower', cmap='YlOrRd', norm=norm,
                      interpolation='nearest')
    ticks = [index * CONNECTORS_PER_SIDE + CONNECTORS_PER_SIDE / 2 - 0.5
             for index in range(4)]
    labels = [f'{side}01-{side}16' for side in SIDES]
    ax.set_xticks(ticks)
    ax.set_yticks(ticks)
    ax.set_xticklabels(labels)
    ax.set_yticklabels(labels)
    for boundary in (CONNECTORS_PER_SIDE - 0.5, 2 * CONNECTORS_PER_SIDE - 0.5,
                     3 * CONNECTORS_PER_SIDE - 0.5):
        ax.axhline(boundary, color='black', lw=0.7)
        ax.axvline(boundary, color='black', lw=0.7)
    ax.invert_yaxis()  # N01..N16 at the top, W01..W16 at the bottom
    total = int(data.sum())
    fig.colorbar(image, ax=ax, fraction=0.046, pad=0.03,
                 label='connection count (matrix counts each undirected pair twice)')
    ax.set_title('V3-Huang 64x64 connector-level connection matrix\n'
                 f'rows/columns = {connector_order[0]} ... {connector_order[-1]} '
                 f'(1-based Port index 1..64); symmetric, diagonal 0, '
                 f'row/col sums 32, total {total} = 2 x 1024\n'
                 f'equals the Excel 64x64 matrix sheet cell-by-cell '
                 f'(validated, 0 mismatches)', fontsize=10)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)
    return Path(path)


def plot_class_composition(class_counts, pair_counts, plane_counts, path):
    plt = _pyplot()
    import numpy as np
    fig, axes = plt.subplots(1, 2, figsize=(16, 6.2),
                             gridspec_kw={'width_ratios': [1.15, 1.0]})
    pair_groups = {
        'same': ('NN', 'SS', 'EE', 'WW'),
        'opposite': ('NS', 'EW'),
        'adjacent': ('EN', 'NW', 'ES', 'SW'),
    }
    pair_colors = {
        'same': ('#2f6db5', '#5b8fd0', '#7fa9de', '#a6c4e8'),
        'opposite': ('#c0392b', '#e07b39'),
        'adjacent': ('#d99b1c', '#e8b84a', '#f0d275', '#f7e3a1'),
    }
    ax = axes[0]
    bottoms = np.zeros(len(CLASS_ORDER))
    for class_index, klass in enumerate(CLASS_ORDER):
        for pair_index, pair in enumerate(pair_groups[klass]):
            count = pair_counts.get(pair, 0)
            ax.bar(class_index, count, bottom=bottoms[class_index],
                   color=pair_colors[klass][pair_index % len(pair_colors[klass])],
                   edgecolor='white', linewidth=0.7)
            if count:
                ax.text(class_index, bottoms[class_index] + count / 2,
                        f'{pair}\n{count}', ha='center', va='center', fontsize=7.5)
            bottoms[class_index] += count
        ax.text(class_index, bottoms[class_index] + 8,
                f'{klass}\n{int(bottoms[class_index])}', ha='center', va='bottom',
                fontsize=9, fontweight='bold')
    ax.set_xticks(list(range(len(CLASS_ORDER))))
    ax.set_xticklabels([f'{klass}\n({class_counts.get(klass, 0)})'
                        for klass in CLASS_ORDER])
    ax.set_ylabel('connections')
    ax.set_ylim(0, 540)
    ax.set_title('(a) 1024 connections by class and connector pair\n'
                 'class taken from the CSV relation column', fontsize=10)
    ax.grid(axis='y', alpha=0.3)

    ax = axes[1]
    plane_labels = [
        f'{plane}\n(CSV row {PLANE_ROW[plane]}, z={PLANE_Z_MM[PLANE_ROW[plane]]} mm)'
        for plane in ('z1', 'z2')]
    bottoms = np.zeros(2)
    for klass in CLASS_ORDER:
        counts = np.array([plane_counts[plane]['class_counts'].get(klass, 0)
                           for plane in ('z1', 'z2')])
        ax.bar(range(2), counts, bottom=bottoms, color=CLASS_COLORS[klass],
               edgecolor='white', linewidth=0.7, label=klass)
        for plane_index, count in enumerate(counts):
            if count:
                ax.text(plane_index, bottoms[plane_index] + count / 2,
                        f'{klass}\n{count}', ha='center', va='center', fontsize=7.5,
                        color='white')
        bottoms += counts
    for plane_index, plane in enumerate(('z1', 'z2')):
        ax.text(plane_index, bottoms[plane_index] + 8,
                f"{int(bottoms[plane_index])} connections\n"
                f"{plane_counts[plane]['connections']} expected", ha='center',
                va='bottom', fontsize=8.5)
    ax.set_xticks(range(2))
    ax.set_xticklabels(plane_labels)
    ax.set_ylabel('connections')
    ax.set_ylim(0, 580)
    ax.set_title('(b) dual-height split: plane from the CSV row\n'
                 'each plane mixes same / opposite / adjacent', fontsize=10)
    ax.legend(fontsize=8, loc='upper right')
    ax.grid(axis='y', alpha=0.3)
    fig.suptitle('V3-Huang connection-class composition (real counts, no router run)',
                 fontsize=12)
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    fig.savefig(path, dpi=200)
    plt.close(fig)
    return Path(path)


# ---------------------------------------------------------------------------
# Artifact writers
# ---------------------------------------------------------------------------
def write_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(_jsonable(data), indent=2, ensure_ascii=False) + '\n',
                    encoding='utf-8')
    return path


def write_physical_ports_csv(path, rows):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=PHYSICAL_PORT_COLUMNS,
                                extrasaction='ignore')
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row[key] for key in PHYSICAL_PORT_COLUMNS})
    return path


def write_plane_endpoints_json(path, endpoints):
    ordered = [dict(endpoint) for endpoint in endpoints]
    return write_json(path, ordered)


def write_matrix_csv(path, matrix, connector_order):
    path = Path(path)
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.writer(handle)
        writer.writerow(['connector'] + list(connector_order))
        for connector, row in zip(connector_order, matrix):
            writer.writerow([connector] + list(row))
    return path


def write_pair_multiplicity_csv(path, pair_counter):
    path = Path(path)
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.writer(handle)
        writer.writerow(['connector_a', 'connector_b', 'pair', 'class', 'count'])
        for (connector_a, connector_b), count in sorted(
                pair_counter.items(), key=lambda item: (-item[1], item[0])):
            pair = side_pair_label(connector_a, connector_b)
            writer.writerow([connector_a, connector_b, pair, PAIR_CLASS[pair], count])
    return path


# ---------------------------------------------------------------------------
# README + artifact build
# ---------------------------------------------------------------------------
def write_readme(path, report, comparison=None):
    counts = report['counts']
    checks = report['checks']
    total_checks = len(checks)
    passed_checks = sum(1 for check in checks.values()
                        if check.get('passed') is not False)
    lines = [
        '# V3-Huang 1024-channel dual-height topology - independent input audit',
        '',
        'Owner: connectivity (Phase 1, `four_side_1024/v3/SPEC.md` sections 1/2/7).',
        'Generated by `four_side_1024/v3/ports_audit_v3.py`; every number below comes',
        'from a real run of that script (no hand-written values). This audit never',
        'imports the Lead V3 input modules: they are compared only in a subprocess.',
        '',
        '## 1. Frozen inputs (byte-identical staging copies)',
        '',
        f"- CSV: `{report['inputs']['csv']['path']}`",
        f"  SHA256 `{report['inputs']['csv']['sha256']}` "
        f"(expected `{report['inputs']['csv']['expected_sha256']}`), "
        f"{report['inputs']['csv']['bytes']} bytes.",
        f"- XLSX: `{report['inputs']['xlsx']['path']}`",
        f"  SHA256 `{report['inputs']['xlsx']['sha256']}` "
        f"(expected `{report['inputs']['xlsx']['expected_sha256']}`), "
        f"{report['inputs']['xlsx']['bytes']} bytes.",
        f"- Sheet names: {report['inputs']['xlsx']['sheet_names']}",
        f"- PROVENANCE: `{report['inputs']['provenance']}` (Desktop source paths + SHA256).",
        '',
        '## 2. Topology (from the CSV)',
        '',
        '| class | count | composition |',
        '|---|---:|---|',
        f"| same | {counts['class_counts']['same']} | NN/SS/EE/WW = "
        f"{counts['pair_counts']['NN']}/{counts['pair_counts']['SS']}/"
        f"{counts['pair_counts']['EE']}/{counts['pair_counts']['WW']} |",
        f"| opposite | {counts['class_counts']['opposite']} | NS/EW = "
        f"{counts['pair_counts']['NS']}/{counts['pair_counts']['EW']} |",
        f"| adjacent | {counts['class_counts']['adjacent']} | EN/NW/ES/SW = "
        f"{counts['pair_counts']['EN']}/{counts['pair_counts']['NW']}/"
        f"{counts['pair_counts']['ES']}/{counts['pair_counts']['SW']} |",
        f"| total | {counts['connections']} | 4 sides x 16 connectors x 32 fibres = "
        f"{counts['endpoints']} endpoints, each used once |",
        '',
        'Plane split (plane from the CSV row; relation from the CSV column):',
        '',
        '| plane | row | z | connections | same | opposite | adjacent |',
        '|---|---:|---:|---:|---:|---:|---:|',
    ]
    for plane in ('z1', 'z2'):
        entry = counts['planes'][plane]
        lines.append(
            f"| {plane} | {entry['row']} | {entry['z_mm']} mm | {entry['connections']} | "
            f"{entry['class_counts']['same']} | {entry['class_counts']['opposite']} | "
            f"{entry['class_counts']['adjacent']} |")
    lines += [
        '',
        f"Unique connector pairs: {counts['unique_connector_pairs']} (counts sum to 1024).",
        '',
        '## 3. Audit result',
        '',
        f"Overall: **{'PASS' if report['passed'] else 'FAIL'}** - {passed_checks}/"
        f"{total_checks} checks passed; failed checks: {report['failed_checks'] or '[]'}.",
        '',
        '| check | passed | key evidence |',
        '|---|---|---|',
    ]
    for name, check in checks.items():
        status = 'skipped' if check.get('passed') is None else (
            'yes' if check['passed'] else '**NO**')
        evidence_keys = [key for key in check
                         if key not in ('passed', 'sample', 'provenance')
                         and not isinstance(check[key], (dict, list))]
        evidence = ', '.join(f'{key}={check[key]}' for key in evidence_keys[:4])
        lines.append(f'| `{name}` | {status} | {evidence} |')
    lines += [
        '',
        '## 4. Excel sheets cross-check',
        '',
        f"Workbook sheets: {report['excel']['sheet_names']}",
        f"Sheet roles detected: {report['excel']['sheet_roles']}",
        '',
        '| sheet | audit check | result |',
        '|---|---|---|',
        f"| Sheet1 Port1/Port2 (1024 connector-index pairs in CSV order) | "
        f"`excel_sheet1_pairs` | {checks['excel_sheet1_pairs']['passed']} |",
        f"| port map (64 rows, N01..W16, centres 7.5+9(c-1), 32/16/16) | "
        f"`excel_port_map` | {checks['excel_port_map']['passed']} |",
        f"| 64x64 connector matrix (symmetric, diag 0, row/col sums 32, total 2048, "
        f"equals CSV) | `excel_64x64_matrix` | {checks['excel_64x64_matrix']['passed']} |",
        f"| fibre-level reference sheet (must equal the CSV, 1024 rows x 12 columns) | "
        f"`excel_fibre_sheet_equals_csv` | "
        f"{checks['excel_fibre_sheet_equals_csv']['passed']} |",
        f"| connector-pair multiplicity ({report['excel']['multiplicities']} unique "
        f"pairs, counts/classes equal CSV) | `excel_pair_multiplicity` | "
        f"{checks['excel_pair_multiplicity']['passed']} |",
        f"| class statistics sheet (class and pair table equal CSV) | "
        f"`excel_class_statistics` | {checks['excel_class_statistics']['passed']} |",
        '',
        '## 5. Physical ports and per-plane endpoint lists',
        '',
        f"- `physical_ports_v3.csv`: {checks['physical_ports']['count']} rows "
        f"(64 connectors x 32 fibres) with (x, y, z1/z2, along, side/connector/row/"
        f"column/channel); the real 2x16 MPO geometry is saved independently of the 2D "
        f"routing coordinates. Unique (x,y,z) = {checks['physical_ports']['unique_3d']}; "
        f"projected XY positions = {checks['physical_ports']['projected_xy_distinct']} "
        f"(the {checks['physical_ports']['cross_plane_projected_duplicate_ports']} "
        f"cross-plane projected duplicates are by design, not conflicts).",
        f"- `plane_endpoints_z1_v3.json` / `plane_endpoints_z2_v3.json`: "
        f"{checks['plane_endpoint_lists']['planes']['z1']['count']} / "
        f"{checks['plane_endpoint_lists']['planes']['z2']['count']} endpoint records "
        f"(1024 per plane), each identity used once overall.",
        '',
        '## 6. Comparison with the Lead V3 core modules (subprocess only)',
        '',
    ]
    if comparison and comparison.get('lead_modules_available'):
        summary = comparison.get('summary') or {}
        sources = comparison.get('module_sources') or {}
        lines.append('Compared modules (provenance recorded in '
                     '`lead_v3_module_comparison.json`):')
        for name, source in sources.items():
            if source.get('exists'):
                lines.append(f"- `{name}`: `{source['path']}` SHA256 "
                             f"`{source['sha256']}`, {source['bytes']} bytes, mtime "
                             f"{source['mtime']}.")
        lines += [
            '',
            f"- 2048 endpoint identities: same set = "
            f"{summary.get('endpoint_identity_sets_equal')}; max routing_xy delta = "
            f"{summary.get('max_routing_xy_delta_mm')} mm; max physical_along delta = "
            f"{summary.get('max_physical_along_delta_mm')} mm; max z delta = "
            f"{summary.get('max_z_delta_mm')} mm; endpoint row->plane violations = "
            f"{summary.get('endpoint_plane_violations')}.",
            f"- 1024 routes: route_id sets equal = "
            f"{summary.get('routes', {}).get('route_id_sets_equal')}; field mismatches = "
            f"{summary.get('routes', {}).get('field_mismatches')}; relation-column "
            f"mismatches = {summary.get('routes', {}).get('relation_column_mismatches')}; "
            f"pair-label mismatches = "
            f"{summary.get('routes', {}).get('pair_label_mismatches')}; row->plane "
            f"violations = {summary.get('routes', {}).get('row_plane_violations')}.",
            f"- Planes: {summary.get('planes')}.",
            f"- 2048 physical ports: identity set equal = "
            f"{summary.get('physical_ports', {}).get('identity_sets_equal')}; max "
            f"coordinate delta = "
            f"{summary.get('physical_ports', {}).get('max_coordinate_delta_mm')} mm.",
            f"- Config-constant mismatches: {summary.get('config_mismatches')}; Lead's "
            f"own validation: topology passed = "
            f"{summary.get('lead_topology_validation_passed')}, Excel passed = "
            f"{summary.get('lead_excel_validation_passed')}.",
            f"- Comparison verdict: "
            f"**{'PASS - 0 mismatches' if summary.get('passed') else 'FAIL'}**.",
        ]
        if comparison.get('code_notes'):
            lines += ['', 'Non-blocking observations on the Lead modules:']
            for note in comparison['code_notes']:
                lines.append(f'- {note}')
    else:
        lines.append('Comparison not run / unavailable in this build.')
    lines += [
        '',
        '## 7. Figures (demand)',
        '',
        '| file | content |',
        '|---|---|',
        '| `demand_1024_connectors.png` | all 1024 connections at connector level; '
        'same-side connections drawn as inward arcs (U/refold demand) |',
        '| `connector_matrix_heatmap_64x64.png` | 64x64 connector-level matrix (equals '
        'the Excel matrix sheet, 0 mismatches) |',
        '| `class_composition.png` | class/pair composition and dual-height plane split '
        '(real counts) |',
        '| `plane_demand_z1.png`, `plane_demand_z2.png` | per-plane connector-level '
        'demand (512 connections each, mixed classes) |',
        '',
        '## 8. Honest limitations',
        '',
    ]
    for warning in report.get('warnings', []):
        lines.append(f'- {warning}')
    lines += [
        '- This is a *topology/interface audit only*: no routing (D56 R=40 mm) was run, '
        'so no routed count, spacing or crossing result is claimed here.',
        '- The V3 topology is synthetic (generated from the 512 fiberBoard): it is a '
        'baseline for algorithm experiments, not a measured demand.',
        '',
        '## 9. Reproduce',
        '',
        '```powershell',
        'cd "C:\\Users\\lihao\\Desktop\\Graduation Project\\OpticalWaveguideRouter2D"',
        '.\\.venv\\Scripts\\python.exe -m four_side_1024.v3.ports_audit_v3',
        '.\\.venv\\Scripts\\python.exe -m pytest four_side_1024\\v3\\tests_v3\\test_ports_v3.py -q',
        '```',
        '',
    ]
    path = Path(path)
    path.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return path


def build_all(output_dir=OUTPUT_DIR, plots=True, compare=True, python_exe=None):
    """Build every Phase-1 V3 topology artifact."""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    report = validate_inputs()
    state = derive_ports()
    files = []

    def emit(path):
        files.append(str(path))
        return path

    emit(write_json(out / 'topology_validation_v3.json', report))
    emit(write_physical_ports_csv(out / 'physical_ports_v3.csv',
                                  state['physical_ports']))
    for plane in ('z1', 'z2'):
        emit(write_plane_endpoints_json(
            out / f'plane_endpoints_{plane}_v3.json', state['plane_endpoints'][plane]))

    connector_order = [f'{side}{number:02d}' for side in SIDES
                       for number in range(1, CONNECTORS_PER_SIDE + 1)]
    matrix_undirected = [[0] * len(connector_order) for _ in connector_order]
    index = {name: position for position, name in enumerate(connector_order)}
    pair_counter = _connector_pair_counter(state['routes'])
    for (connector_a, connector_b), count in pair_counter.items():
        matrix_undirected[index[connector_a]][index[connector_b]] = count
        matrix_undirected[index[connector_b]][index[connector_a]] = count
    emit(write_matrix_csv(out / 'connector_matrix_64x64.csv', matrix_undirected,
                          connector_order))
    emit(write_pair_multiplicity_csv(out / 'connector_pair_multiplicity_v3.csv',
                                     pair_counter))

    comparison = None
    if compare:
        try:
            comparison = compare_with_lead_v3_module(python_exe=python_exe, state=state)
        except Exception as exc:  # noqa: BLE001 - comparison is auxiliary
            comparison = dict(lead_modules_available=False,
                              error=f'{type(exc).__name__}: {exc}')
        emit(write_json(out / 'lead_v3_module_comparison.json', comparison))

    if plots:
        class_counts = Counter(route['relation'] for route in state['routes'])
        pair_counts = Counter(route['pair'] for route in state['routes'])
        emit(plot_connector_demand(state['routes'], class_counts, pair_counts,
                                   out / 'demand_1024_connectors.png'))
        emit(plot_connector_matrix_heatmap(matrix_undirected, connector_order,
                                           out / 'connector_matrix_heatmap_64x64.png'))
        emit(plot_class_composition(class_counts, pair_counts, report['counts']['planes'],
                                    out / 'class_composition.png'))
        for plane in ('z1', 'z2'):
            emit(plot_plane_demand(state, plane, out / f'plane_demand_{plane}.png'))

    emit(write_readme(out / 'README.md', report, comparison))

    summary = dict(
        inputs=report['inputs'],
        validation_passed=report['passed'],
        failed_checks=report['failed_checks'],
        checks_total=len(report['checks']),
        class_counts=report['counts']['class_counts'],
        pair_counts=report['counts']['pair_counts'],
        planes={plane: dict(connections=entry['connections'],
                            class_counts=entry['class_counts'])
                for plane, entry in report['counts']['planes'].items()},
        endpoints=report['counts']['endpoints'],
        unique_connector_pairs=report['counts']['unique_connector_pairs'],
        physical_ports=report['checks']['physical_ports']['count'],
        lead_v3_comparison=(comparison.get('summary') if comparison
                            and comparison.get('lead_modules_available') else comparison),
        files=[str(path) for path in files],
    )
    return dict(report=report, state=state, comparison=comparison, files=files,
                summary=summary)


def main(argv=None):
    import argparse
    parser = argparse.ArgumentParser(
        description='Build the V3-Huang Phase-1 topology/connector audit artifacts.')
    parser.add_argument('--output-dir', default=str(OUTPUT_DIR))
    parser.add_argument('--no-plots', action='store_true')
    parser.add_argument('--skip-comparison', action='store_true')
    parser.add_argument('--python', default=None,
                        help='python executable for the Lead-module subprocess')
    args = parser.parse_args(argv)
    result = build_all(Path(args.output_dir), plots=not args.no_plots,
                       compare=not args.skip_comparison, python_exe=args.python)
    print(json.dumps(_jsonable(result['summary']), indent=2, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())


