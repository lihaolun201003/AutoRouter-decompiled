"""Phase-1 acceptance tests for ``four_side_1024/v3/ports_audit_v3.py``.

Covers the V3-Huang input/topology audit (``four_side_1024/v3/SPEC.md`` sections
1/2/7):

* both frozen inputs (CSV + XLSX) and the PROVENANCE SHA256 records;
* 1024 rows / route ids, class counts 448/384/192, pair counts, channel and
  connector identity rules, no same-side self connector, 2048 endpoints used
  exactly once, 64 connectors x 32 fibres;
* dual-height split from the CSV row (relation from the CSV column only):
  z1 = 512 (same 222 / opposite 196 / adjacent 94), z2 = 512 (226/188/98);
* Excel cross-checks: Sheet1 Port1/Port2 == CSV connector pairs in order, port
  map, 64x64 matrix == connector-level counts, fibre-level reference sheet ==
  CSV exactly, connector-pair multiplicity sheet, class statistics sheet;
* physical 2048-port 3D table (2x16 MPO geometry) and the 1024-endpoint lists
  per plane;
* negative controls: tampered CSV and tampered Excel are detected;
* comparison with Lead's V3 core modules (subprocess): 0 mismatches, 0.0 mm
  max deviation, 0 row->plane inconsistencies;
* the full artifact build (JSON/CSV/PNG/README) into a temporary directory.

Run from ``OpticalWaveguideRouter2D``::

    .\\.venv\\Scripts\\python.exe -m pytest four_side_1024\\v3\\tests_v3\\test_ports_v3.py -q
"""
import copy
import csv
import json
import subprocess
import sys
from pathlib import Path

import pytest

from four_side_1024.v3 import ports_audit_v3 as A

INPUTS_AVAILABLE = A.CSV_PATH.exists() and A.XLSX_PATH.exists()
pytestmark = pytest.mark.skipif(not INPUTS_AVAILABLE,
                                reason='V3-Huang staged inputs not found')

REQUIRED_README_PHRASES = [
    'simulation under the two-height interface assumption',
    'no z transition',
    'glass thickness',
    'mechanical',
    'cross-plane',
    'Same-side connections',
]


# ---------------------------------------------------------------------------
# fixtures (the audit is deterministic; load once per session)
# ---------------------------------------------------------------------------
@pytest.fixture(scope='module')
def report():
    return A.validate_inputs()


@pytest.fixture(scope='module')
def state():
    return A.derive_ports()


@pytest.fixture(scope='module')
def comparison():
    return A.compare_with_lead_v3_module()


# ---------------------------------------------------------------------------
# 1. model constants / derivations
# ---------------------------------------------------------------------------
def test_model_constants_and_formulas():
    assert (A.BOARD_MM, A.CONNECTORS_PER_SIDE, A.COLUMNS, A.ROWS) == (150.0, 16, 16, 2)
    assert A.CHANNELS_PER_CONNECTOR == 32
    assert A.FIBER_COLUMN_PITCH_MM == 0.25
    assert A.FIBER_ROW_PITCH_MM == 0.5
    assert A.CONNECTOR_PITCH_MM == 9.0
    assert A.FIRST_CONNECTOR_CENTER_MM == 7.5  # (150 - 15*9)/2
    assert A.PLANE_Z_MM == {1: 0.0, 2: 0.5}
    assert A.PLANE_NAME == {1: 'z1', 2: 'z2'}
    assert A.CLASS_ORDER == ('same', 'opposite', 'adjacent')
    assert A.connector_center_mm('N01') == 7.5
    assert A.connector_center_mm('W16') == 142.5
    assert A.port_along_mm('N01', 1) == pytest.approx(5.625)
    assert A.port_along_mm('N16', 16) == pytest.approx(144.375)
    assert A.port_along_mm('E05', 9) - A.port_along_mm('E05', 8) == pytest.approx(0.25)
    assert A.edge_xy('N', 10.0) == [10.0, 150.0]
    assert A.edge_xy('S', 10.0) == [10.0, 0.0]
    assert A.edge_xy('E', 10.0) == [150.0, 10.0]
    assert A.edge_xy('W', 10.0) == [0.0, 10.0]
    assert A.side_pair_label('N01', 'N06') == 'NN'
    assert A.side_pair_label('E11', 'W03') == 'EW'
    assert A.side_pair_label('N01', 'E01') == 'EN'


def test_sha256_and_provenance(report):
    checks = report['checks']
    assert report['inputs']['csv']['sha256'] == A.CSV_SHA256
    assert report['inputs']['xlsx']['sha256'] == A.XLSX_SHA256
    provenance = checks['sha256_provenance']
    assert provenance['passed'] is True
    for entry in provenance['files']:
        assert entry['matches_stored'] is True
        assert entry['matches_source'] is True
        assert entry['matches_task_expected'] is True
    assert report['inputs']['csv']['expected_sha256'] == A.CSV_SHA256
    assert report['inputs']['xlsx']['expected_sha256'] == A.XLSX_SHA256
    assert checks['csv_header']['observed'] == A.CSV_FIELDS
    assert checks['excel_sheet_count']['observed'] == 7


# ---------------------------------------------------------------------------
# 2. CSV topology checks
# ---------------------------------------------------------------------------
def test_csv_topology(report):
    checks = report['checks']
    assert report['passed'] is True
    assert report['failed_checks'] == []
    assert len(report['checks']) >= 20
    assert checks['connections']['observed'] == 1024
    assert checks['route_ids']['unique'] == 1024
    assert checks['class_counts']['observed'] == {'same': 448, 'opposite': 384,
                                                  'adjacent': 192}
    assert checks['pair_counts']['observed'] == {
        'NN': 112, 'SS': 112, 'EE': 112, 'WW': 112,
        'NS': 192, 'EW': 192, 'EN': 48, 'NW': 48, 'ES': 48, 'SW': 48}
    assert checks['relation_matches_sides']['violations'] == 0
    assert checks['channel_formula']['violations'] == 0
    assert checks['connector_letter']['violations'] == 0
    assert checks['no_same_side_self_connector']['violations'] == 0
    assert checks['endpoints_used_once']['distinct'] == 2048
    assert checks['endpoints_used_once']['usage_min'] == 1
    assert checks['endpoints_used_once']['usage_max'] == 1
    assert checks['endpoints_used_once']['missing'] == 0
    assert checks['endpoint_count']['observed'] == 2048
    assert checks['connector_fibres']['connectors'] == 64
    assert checks['connector_fibres']['channels_ok'] == 64
    assert checks['no_mixed_rows']['violations'] == 0


def test_dual_height_plane_split(report):
    checks = report['checks']
    split = checks['dual_height_split']
    assert split['passed'] is True
    z1, z2 = report['counts']['planes']['z1'], report['counts']['planes']['z2']
    assert z1['connections'] == z2['connections'] == 512
    assert z1['class_counts'] == {'same': 222, 'opposite': 196, 'adjacent': 94}
    assert z2['class_counts'] == {'same': 226, 'opposite': 188, 'adjacent': 98}
    assert z1['row'] == 1 and z1['z_mm'] == 0.0
    assert z2['row'] == 2 and z2['z_mm'] == 0.5
    # relation is NOT inferred from the row: both planes mix all three classes
    for plan in (z1, z2):
        assert all(plan['class_counts'][klass] > 0 for klass in A.CLASS_ORDER)


def test_plane_endpoint_lists(report):
    checks = report['checks']['plane_endpoint_lists']
    assert checks['passed'] is True
    assert checks['planes']['z1']['count'] == 1024
    assert checks['planes']['z2']['count'] == 1024
    assert checks['planes']['z1']['rows'] == [1]
    assert checks['planes']['z2']['rows'] == [2]


# ---------------------------------------------------------------------------
# 3. Excel cross-checks
# ---------------------------------------------------------------------------
def test_excel_cross_checks(report):
    checks = report['checks']
    assert report['excel']['port_pairs'] == 1024
    assert report['excel']['port_map_rows'] == 64
    assert report['excel']['matrix_shape'] == [64, 64]
    assert report['excel']['fibre_rows'] == 1024
    assert report['excel']['multiplicities'] == 549
    assert checks['excel_sheet1_pairs']['passed'] is True
    assert checks['excel_sheet1_pairs']['same_order'] is True
    assert checks['excel_sheet1_pairs']['self_pairs'] == 0
    assert checks['excel_port_map']['passed'] is True
    assert checks['excel_port_map']['names_match'] is True
    assert checks['excel_port_map']['distinct_targets_match'] is True
    matrix_check = checks['excel_64x64_matrix']
    assert matrix_check['passed'] is True
    assert matrix_check['equals_csv_connector_counts'] is True
    assert matrix_check['symmetric'] is True
    assert matrix_check['diagonal_zero'] is True
    assert matrix_check['row_sums'] == [32]
    assert matrix_check['total'] == 2048
    fibre_check = checks['excel_fibre_sheet_equals_csv']
    assert fibre_check['passed'] is True
    assert fibre_check['mismatches'] == 0
    assert fibre_check['rows'] == 1024
    multiplicity = checks['excel_pair_multiplicity']
    assert multiplicity['passed'] is True
    assert multiplicity['unique_pairs'] == 549
    assert multiplicity['csv_unique_pairs'] == 549
    assert multiplicity['total'] == 1024
    stats = checks['excel_class_statistics']
    assert stats['passed'] is True
    assert stats['sheet_counts']['total'] == 1024
    assert report['excel']['class_stats']['same'] == 448


# ---------------------------------------------------------------------------
# 4. physical 2048-port 3D table
# ---------------------------------------------------------------------------
def test_physical_ports(state, report):
    ports = state['physical_ports']
    assert len(ports) == 2048
    check = report['checks']['physical_ports']
    assert check['passed'] is True
    assert check['count'] == 2048
    assert check['z1'] == check['z2'] == 1024
    assert check['unique_3d'] == 2048
    assert check['projected_xy_distinct'] == 1024
    assert check['cross_plane_projected_duplicate_ports'] == 1024
    assert check['max_coordinate_error_mm'] <= A.TOL_MM
    identities = {(port['connector'], port['row'], port['column']) for port in ports}
    assert len(identities) == 2048
    for port in ports:
        assert port['plane'] == A.PLANE_NAME[port['row']]
        assert port['z_mm'] == A.PLANE_Z_MM[port['row']]
        assert port['channel'] == (port['row'] - 1) * 16 + port['column']
        expected_along = A.port_along_mm(port['connector'], port['column'])
        assert port['along_mm'] == pytest.approx(expected_along, abs=A.TOL_MM)
        if port['side'] == 'N':
            assert port['y_mm'] == 150.0
        elif port['side'] == 'S':
            assert port['y_mm'] == 0.0
        elif port['side'] == 'E':
            assert port['x_mm'] == 150.0
        else:
            assert port['x_mm'] == 0.0
    assert report['checks']['port_pitches']['passed'] is True
    assert report['checks']['port_pitches']['column_pitch_observed'] == [0.25]
    assert report['checks']['port_pitches']['row_z_separation_mm'] == 0.5


def test_endpoints_are_2d_projections_of_real_ports(state):
    endpoints = state['endpoints']
    assert len(endpoints) == 2048
    # same XY projection, different z -> the two-height model is real
    assert endpoints['N01:1:1']['routing_xy'] == endpoints['N01:2:1']['routing_xy']
    assert endpoints['N01:1:1']['z_mm'] == 0.0
    assert endpoints['N01:2:1']['z_mm'] == 0.5
    assert endpoints['N01:1:1']['plane'] == 'z1'
    assert endpoints['N01:2:1']['plane'] == 'z2'


# ---------------------------------------------------------------------------
# 5. negative controls
# ---------------------------------------------------------------------------
def _write_tampered_csv(tmp_path, mutate):
    routes, header, _sha = A.read_connectivity_csv(A.CSV_PATH)
    routes = copy.deepcopy(routes)
    mutate(routes)
    path = tmp_path / 'tampered_fiber_pairs.csv'
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=header)
        writer.writeheader()
        for route in routes:
            writer.writerow({field: route[field] for field in header})
    return path


def test_negative_control_duplicate_endpoint(tmp_path):
    def mutate(routes):
        duplicate = routes[0]
        target = routes[1]
        for field in ('source_side', 'source_connector', 'source_row',
                      'source_column', 'source_channel'):
            target[field] = duplicate[field]

    csv_path = _write_tampered_csv(tmp_path, mutate)
    report = A.validate_inputs(csv_path=csv_path, xlsx_path=A.XLSX_PATH)
    assert report['passed'] is False
    assert 'endpoints_used_once' in report['failed_checks']
    assert 'sha256_provenance' in report['failed_checks']


def test_negative_control_channel_formula(tmp_path):
    def mutate(routes):
        routes[0]['target_channel'] = 1 if routes[0]['target_channel'] != 1 else 2

    csv_path = _write_tampered_csv(tmp_path, mutate)
    report = A.validate_inputs(csv_path=csv_path, xlsx_path=A.XLSX_PATH)
    assert report['passed'] is False
    assert 'channel_formula' in report['failed_checks']


def test_negative_control_relation_from_row(tmp_path):
    """A row/relation mismatch must be caught (relation must come from the column)."""

    def mutate(routes):
        # route 0 is a same-side row-1 connection: fake it as opposite without
        # changing the sides -> the CSV relation column no longer matches the sides.
        routes[0]['relation'] = 'opposite'

    csv_path = _write_tampered_csv(tmp_path, mutate)
    report = A.validate_inputs(csv_path=csv_path, xlsx_path=A.XLSX_PATH)
    assert report['passed'] is False
    assert 'relation_matches_sides' in report['failed_checks']
    assert 'class_counts' in report['failed_checks']


def test_negative_control_mixed_rows(tmp_path):
    def mutate(routes):
        # same-side connection: move only the target to row 2 but keep columns valid
        route = routes[0]
        column = int(route['target_column'])
        route['target_row'] = 2
        route['target_channel'] = 16 + column

    csv_path = _write_tampered_csv(tmp_path, mutate)
    report = A.validate_inputs(csv_path=csv_path, xlsx_path=A.XLSX_PATH)
    assert report['passed'] is False
    assert 'no_mixed_rows' in report['failed_checks']
    assert 'plane_endpoint_lists' in report['failed_checks']
    assert 'endpoints_used_once' in report['failed_checks']


def test_negative_control_tampered_matrix(tmp_path):
    import openpyxl
    workbook = openpyxl.load_workbook(A.XLSX_PATH, data_only=False)
    matrix_sheet = [sheet for sheet in workbook.worksheets
                    if sheet.max_row == 65 and sheet.max_column == 65][0]
    matrix_sheet.cell(row=2, column=2).value = 99  # N01 x N01 diagonal
    tampered_path = tmp_path / 'tampered_connectors.xlsx'
    workbook.save(tampered_path)
    workbook.close()
    report = A.validate_inputs(csv_path=A.CSV_PATH, xlsx_path=tampered_path)
    assert report['passed'] is False
    assert 'excel_64x64_matrix' in report['failed_checks']
    assert 'sha256_provenance' in report['failed_checks']


def test_negative_control_tampered_fibre_reference_sheet(tmp_path):
    import openpyxl
    workbook = openpyxl.load_workbook(A.XLSX_PATH, data_only=False)
    fibre_sheet = [sheet for sheet in workbook.worksheets
                   if str(sheet.cell(row=1, column=1).value).strip() == 'route_id'][0]
    fibre_sheet.cell(row=2, column=5).value = 99  # source_column of route 1
    tampered_path = tmp_path / 'tampered_fibre.xlsx'
    workbook.save(tampered_path)
    workbook.close()
    report = A.validate_inputs(csv_path=A.CSV_PATH, xlsx_path=tampered_path)
    assert report['passed'] is False
    assert 'excel_fibre_sheet_equals_csv' in report['failed_checks']


# ---------------------------------------------------------------------------
# 6. comparison with Lead's V3 core modules
# ---------------------------------------------------------------------------
def test_comparison_with_lead_v3_modules(comparison):
    assert comparison['lead_modules_available'] is True
    summary = comparison['summary']
    assert summary['passed'] is True
    assert summary['endpoint_identities'] == 2048
    assert summary['endpoint_identity_sets_equal'] is True
    assert summary['max_routing_xy_delta_mm'] == 0.0
    assert summary['max_physical_along_delta_mm'] == 0.0
    assert summary['max_z_delta_mm'] == 0.0
    assert summary['endpoint_plane_violations'] == 0
    routes = summary['routes']
    assert routes['route_id_sets_equal'] is True
    assert routes['field_mismatches'] == 0
    assert routes['relation_column_mismatches'] == 0
    assert routes['pair_label_mismatches'] == 0
    assert routes['row_plane_violations'] == 0
    planes = summary['planes']
    for plane in ('z1', 'z2'):
        assert planes[plane]['lead_validation_passed'] is True
        assert planes[plane]['route_id_sets_equal'] is True
        assert planes[plane]['endpoint_set_mismatches'] == 0
        assert planes[plane]['max_routing_xy_delta_mm'] == 0.0
    ports = summary['physical_ports']
    assert ports['identity_sets_equal'] is True
    assert ports['max_coordinate_delta_mm'] == 0.0
    assert ports['field_violations'] == 0
    assert summary['config_mismatches'] == []
    assert summary['lead_topology_validation_passed'] is True
    assert summary['lead_excel_validation_passed'] is True
    # provenance of the compared modules is recorded
    for name in ('inputs', 'ports', 'config'):
        source = comparison['module_sources'][name]
        assert source['exists'] is True
        assert len(source['sha256']) == 64


# ---------------------------------------------------------------------------
# 7. independence of the audit from Lead's V3 input modules
# ---------------------------------------------------------------------------
def test_audit_module_does_not_import_lead_v3_input_modules():
    code = ('import sys; '
            'from four_side_1024.v3 import ports_audit_v3 as A; '
            'A.derive_ports(); '
            "print(int('four_side_1024.v3.inputs' in sys.modules), "
            "int('four_side_1024.v3.ports' in sys.modules))")
    completed = subprocess.run([sys.executable, '-c', code], cwd=str(A.CODE_DIR),
                               capture_output=True, text=True, timeout=300)
    assert completed.returncode == 0, completed.stderr
    assert completed.stdout.strip() == '0 0'


# ---------------------------------------------------------------------------
# 8. artifact build
# ---------------------------------------------------------------------------
def test_build_all_artifacts(tmp_path):
    result = A.build_all(tmp_path, plots=True, compare=True)
    summary = result['summary']
    assert summary['validation_passed'] is True
    assert summary['lead_v3_comparison']['passed'] is True
    expected_files = [
        'topology_validation_v3.json', 'physical_ports_v3.csv',
        'plane_endpoints_z1_v3.json', 'plane_endpoints_z2_v3.json',
        'connector_matrix_64x64.csv', 'connector_pair_multiplicity_v3.csv',
        'lead_v3_module_comparison.json', 'README.md',
        'demand_1024_connectors.png', 'connector_matrix_heatmap_64x64.png',
        'class_composition.png', 'plane_demand_z1.png', 'plane_demand_z2.png',
    ]
    for name in expected_files:
        path = tmp_path / name
        assert path.exists() and path.stat().st_size > 0, name
    report = json.loads((tmp_path / 'topology_validation_v3.json').read_text(
        encoding='utf-8'))
    assert report['passed'] is True
    assert report['counts']['class_counts'] == {'same': 448, 'opposite': 384,
                                                'adjacent': 192}
    with (tmp_path / 'physical_ports_v3.csv').open(encoding='utf-8', newline='') as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 2048
    for column in A.PHYSICAL_PORT_COLUMNS:
        assert column in rows[0]
    assert {row['plane'] for row in rows} == {'z1', 'z2'}
    matrix_rows = list(csv.reader((tmp_path / 'connector_matrix_64x64.csv').open(
        encoding='utf-8')))
    assert len(matrix_rows) == 65 and len(matrix_rows[0]) == 65
    multiplicity_rows = list(csv.reader(
        (tmp_path / 'connector_pair_multiplicity_v3.csv').open(encoding='utf-8')))
    assert len(multiplicity_rows) == 550  # 549 unique pairs + header
    with (tmp_path / 'plane_endpoints_z1_v3.json').open(encoding='utf-8') as handle:
        plane_z1 = json.load(handle)
    assert len(plane_z1) == 1024
    assert {entry['plane'] for entry in plane_z1} == {'z1'}
    readme = (tmp_path / 'README.md').read_text(encoding='utf-8')
    for phrase in REQUIRED_README_PHRASES:
        assert phrase in readme, phrase
    comparison = json.loads((tmp_path / 'lead_v3_module_comparison.json').read_text(
        encoding='utf-8'))
    assert comparison['summary']['passed'] is True
    assert comparison['summary']['max_routing_xy_delta_mm'] == 0.0
